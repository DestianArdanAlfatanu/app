"""Backend tests for Sistem LPK Jepang MVP."""
import os
import io
import base64
import pytest
import requests
from datetime import date

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL")
if not BASE_URL:
    # fall back to reading frontend/.env
    with open("/app/frontend/.env") as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL="):
                BASE_URL = line.split("=", 1)[1].strip()
BASE_URL = BASE_URL.rstrip("/")
API = f"{BASE_URL}/api"

CREDS = {
    "owner": ("owner@lpk.id", "owner123"),
    "admin": ("admin@lpk.id", "password123"),
    "finance": ("finance@lpk.id", "password123"),
    "hr": ("hr@lpk.id", "password123"),
    "guru": ("guru@lpk.id", "password123"),
    "marketing": ("marketing@lpk.id", "password123"),
    "staff": ("staff@lpk.id", "password123"),
}


def _login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=30)
    return r


def _headers(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="session")
def tokens():
    out = {}
    for role, (email, pwd) in CREDS.items():
        r = _login(email, pwd)
        assert r.status_code == 200, f"login {role} failed: {r.status_code} {r.text}"
        out[role] = r.json()["token"]
    return out


# ---------------- AUTH ----------------
class TestAuth:
    def test_login_owner(self):
        r = _login(*CREDS["owner"])
        assert r.status_code == 200
        d = r.json()
        assert "token" in d and "user" in d
        assert d["user"]["email"] == "owner@lpk.id"
        assert d["user"]["role"] == "owner"

    def test_login_wrong_password(self):
        r = _login("owner@lpk.id", "wrongpwd")
        assert r.status_code == 401
        assert "salah" in r.json().get("detail", "").lower()

    def test_me(self, tokens):
        r = requests.get(f"{API}/auth/me", headers=_headers(tokens["owner"]))
        assert r.status_code == 200
        assert r.json()["role"] == "owner"

    def test_guru_cannot_post_payment(self, tokens):
        r = requests.post(f"{API}/payments",
                          headers=_headers(tokens["guru"]),
                          json={"student_id": "x", "nominal": 100000, "tanggal": date.today().isoformat(),
                                "metode": "tunai", "account_id": "x"})
        assert r.status_code == 403

    def test_finance_cannot_post_employee(self, tokens):
        r = requests.post(f"{API}/employees",
                          headers=_headers(tokens["finance"]),
                          json={"name": "TEST_x", "role": "guru"})
        assert r.status_code == 403

    def test_nonowner_cannot_create_user(self, tokens):
        r = requests.post(f"{API}/users",
                          headers=_headers(tokens["admin"]),
                          json={"name": "TEST_u", "email": "test_nou@x.id", "password": "abcdef", "role": "staff"})
        assert r.status_code == 403


# ---------------- STUDENTS ----------------
class TestStudents:
    @pytest.fixture(scope="class")
    def student_id(self, tokens):
        r = requests.post(f"{API}/students",
                          headers=_headers(tokens["admin"]),
                          json={"nama_lengkap": "TEST_Budi Santoso", "jenis_kelamin": "L",
                                "tanggal_lahir": "2000-05-10", "no_hp": "0811",
                                "alamat": {"jalan": "Jl. Test", "kota": "Jakarta"}})
        assert r.status_code == 200, r.text
        sid = r.json()["id"]
        return sid

    def test_list_enriched(self, tokens):
        r = requests.get(f"{API}/students", headers=_headers(tokens["admin"]))
        assert r.status_code == 200
        rows = r.json()
        assert isinstance(rows, list)
        if rows:
            row = rows[0]
            assert "pembayaran" in row
            assert set(["total", "bayar", "sisa"]).issubset(row["pembayaran"].keys())

    def test_created_default_feeplan(self, tokens, student_id):
        r = requests.get(f"{API}/students/{student_id}", headers=_headers(tokens["admin"]))
        assert r.status_code == 200
        prof = r.json()
        total = prof["pembayaran"]["total"]
        assert total == 18500000, f"expected 18,500,000, got {total}"
        assert prof["status_history"], "status_history missing"

    def test_status_change(self, tokens, student_id):
        r = requests.put(f"{API}/students/{student_id}/status",
                         headers=_headers(tokens["admin"]),
                         json={"status": "seleksi", "catatan": "TEST"})
        assert r.status_code == 200, r.text
        r2 = requests.get(f"{API}/students/{student_id}", headers=_headers(tokens["admin"]))
        d = r2.json()
        assert d["status"] == "seleksi"
        assert any(h["status"] == "seleksi" for h in d["status_history"])

    def test_selection(self, tokens, student_id):
        r = requests.post(f"{API}/students/{student_id}/selections",
                          headers=_headers(tokens["admin"]),
                          json={"tanggal": date.today().isoformat(), "jenis": "wawancara",
                                "hasil": "lulus", "catatan": "ok"})
        assert r.status_code == 200, r.text

    def test_documents_checklist(self, tokens, student_id):
        r = requests.get(f"{API}/students/{student_id}/documents", headers=_headers(tokens["admin"]))
        assert r.status_code == 200
        docs = r.json()
        assert len(docs) == 13, f"expected 13 doc checklist items, got {len(docs)}"

    def test_doc_status_toggle(self, tokens, student_id):
        r = requests.put(f"{API}/students/{student_id}/documents/status",
                         headers=_headers(tokens["admin"]),
                         json={"jenis": "KTP", "kategori": "identitas", "status": "tersedia"})
        assert r.status_code == 200, r.text

    def test_doc_upload_and_download(self, tokens, student_id):
        # 1x1 PNG
        png_b64 = ("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR4nGNgAAIAAAUAAeI"
                   "mDPMAAAAASUVORK5CYII=")
        png = base64.b64decode(png_b64)
        files = {"file": ("test.png", io.BytesIO(png), "image/png")}
        data = {"jenis": "KTP", "kategori": "identitas"}
        r = requests.post(f"{API}/students/{student_id}/documents",
                          headers=_headers(tokens["admin"]), files=files, data=data)
        if r.status_code != 200:
            pytest.skip(f"Object storage upload not available: {r.status_code} {r.text[:200]}")
        doc = r.json()
        file_id = doc.get("file_id")
        if not file_id:
            pytest.skip("No file_id in response")
        r2 = requests.get(f"{API}/files/{file_id}?auth={tokens['admin']}")
        assert r2.status_code == 200
        assert len(r2.content) > 0


# ---------------- PAYMENTS ----------------
class TestPayments:
    @pytest.fixture(scope="class")
    def ctx(self, tokens):
        # pick an existing student and account
        s = requests.get(f"{API}/students", headers=_headers(tokens["admin"])).json()
        a = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"])).json()
        assert s and a, "need seeded students & accounts"
        return {"student_id": s[0]["id"], "account_id": a[0]["id"], "saldo_before": a[0]["saldo"]}

    def test_create_payment(self, tokens, ctx):
        r = requests.post(f"{API}/payments", headers=_headers(tokens["finance"]),
                          json={"student_id": ctx["student_id"], "nominal": 500000,
                                "tanggal": date.today().isoformat(), "metode": "tunai",
                                "account_id": ctx["account_id"], "keterangan": "TEST payment"})
        assert r.status_code == 200, r.text
        p = r.json()
        assert p["no_kwitansi"].startswith("KW-")
        ctx["payment_id"] = p["id"]
        # verify saldo went up
        a = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"])).json()
        after = next(x["saldo"] for x in a if x["id"] == ctx["account_id"])
        assert after == ctx["saldo_before"] + 500000, f"saldo not updated: {after} vs {ctx['saldo_before']}+500000"
        # verify transaction created
        tx = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"])).json()
        assert any(t.get("ref_type") == "payment" and t.get("ref_id") == p["id"] for t in tx)

    def test_receipt(self, tokens, ctx):
        pid = ctx.get("payment_id")
        assert pid
        r = requests.get(f"{API}/payments/{pid}/receipt", headers=_headers(tokens["finance"]))
        assert r.status_code == 200
        d = r.json()
        for k in ("total_tagihan", "total_dibayar", "sisa"):
            assert k in d, f"missing {k}"

    def test_update_without_alasan(self, tokens, ctx):
        pid = ctx.get("payment_id")
        r = requests.put(f"{API}/payments/{pid}", headers=_headers(tokens["finance"]),
                         json={"student_id": ctx["student_id"], "nominal": 600000,
                               "tanggal": date.today().isoformat(), "account_id": ctx["account_id"],
                               "metode": "tunai"})
        assert r.status_code == 400, r.text

    def test_update_with_alasan(self, tokens, ctx):
        pid = ctx.get("payment_id")
        r = requests.put(f"{API}/payments/{pid}", headers=_headers(tokens["finance"]),
                         json={"student_id": ctx["student_id"], "nominal": 600000,
                               "tanggal": date.today().isoformat(), "account_id": ctx["account_id"],
                               "metode": "tunai", "alasan": "TEST koreksi"})
        assert r.status_code == 200, r.text

    def test_arrears(self, tokens):
        r = requests.get(f"{API}/payments-arrears?filter=terlambat", headers=_headers(tokens["finance"]))
        assert r.status_code == 200
        assert isinstance(r.json(), list)


# ---------------- FINANCE ----------------
class TestFinance:
    def test_summary(self, tokens):
        r = requests.get(f"{API}/finance/summary?period=bulan", headers=_headers(tokens["finance"]))
        assert r.status_code == 200
        d = r.json()
        for k in ("pemasukan", "pengeluaran", "saldo_kas", "piutang", "accounts"):
            assert k in d

    def test_expense_reduces_saldo(self, tokens):
        # Use last account to avoid race with parallel payment tests on first account
        acc = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"])).json()
        aid = acc[-1]["id"]
        before = acc[-1]["saldo"]
        r = requests.post(f"{API}/finance/transactions", headers=_headers(tokens["finance"]),
                          json={"jenis": "pengeluaran", "kategori": "Operasional Lainnya", "nominal": 100000,
                                "deskripsi": "TEST expense", "tanggal": date.today().isoformat(),
                                "account_id": aid})
        assert r.status_code == 200, r.text
        acc2 = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"])).json()
        after = next(x["saldo"] for x in acc2 if x["id"] == aid)
        assert after == before - 100000, f"expected {before - 100000}, got {after}"

    def test_tx_update_needs_alasan(self, tokens):
        tx = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"])).json()
        # find a manual tx (not from payment)
        manual = next((t for t in tx if t.get("ref_type") != "payment"), None)
        assert manual, "no manual transaction found"
        tid = manual["id"]
        r = requests.put(f"{API}/finance/transactions/{tid}", headers=_headers(tokens["finance"]),
                         json={"jenis": manual["jenis"], "kategori": manual["kategori"],
                               "nominal": 12345, "tanggal": manual["tanggal"],
                               "account_id": manual["account_id"]})
        assert r.status_code == 400, r.text

    def test_reconciliation_selisih(self, tokens):
        acc = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"])).json()
        aid = acc[0]["id"]
        sistem = acc[0]["saldo"]
        # no alasan with selisih != 0 should 400
        r = requests.post(f"{API}/finance/reconciliations", headers=_headers(tokens["finance"]),
                          json={"account_id": aid, "tanggal": date.today().isoformat(),
                                "saldo_aktual": sistem + 5000})
        assert r.status_code == 400
        r2 = requests.post(f"{API}/finance/reconciliations", headers=_headers(tokens["finance"]),
                           json={"account_id": aid, "tanggal": date.today().isoformat(),
                                 "saldo_aktual": sistem + 5000, "alasan": "TEST selisih kas"})
        assert r2.status_code == 200
        assert r2.json().get("selisih") == 5000

    def test_cashflow(self, tokens):
        r = requests.get(f"{API}/finance/cashflow", headers=_headers(tokens["finance"]))
        assert r.status_code == 200
        assert len(r.json()) == 6


# ---------------- ACADEMICS ----------------
class TestAcademics:
    def test_classes_enriched(self, tokens):
        r = requests.get(f"{API}/classes", headers=_headers(tokens["admin"]))
        assert r.status_code == 200
        rows = r.json()
        assert rows
        assert "guru_nama" in rows[0] and "jumlah_siswa" in rows[0]

    def test_attendance_upsert(self, tokens):
        cls = requests.get(f"{API}/classes", headers=_headers(tokens["admin"])).json()
        cid = cls[0]["id"]
        det = requests.get(f"{API}/classes/{cid}", headers=_headers(tokens["admin"])).json()
        students = det.get("students") or det.get("siswa") or []
        if not students:
            pytest.skip("no students in class")
        sid = students[0]["id"]
        today = date.today().isoformat()
        payload = {"class_id": cid, "tanggal": today,
                   "records": [{"student_id": sid, "status": "hadir"}]}
        r1 = requests.post(f"{API}/attendance", headers=_headers(tokens["admin"]), json=payload)
        assert r1.status_code == 200, r1.text
        r2 = requests.post(f"{API}/attendance", headers=_headers(tokens["admin"]), json=payload)
        assert r2.status_code == 200, "second upsert should not fail"
        rec = requests.get(f"{API}/attendance/recap?class_id={cid}", headers=_headers(tokens["admin"]))
        assert rec.status_code == 200
        d = rec.json()
        # recap is a dict keyed by student_id
        assert isinstance(d, dict)
        if d:
            first = next(iter(d.values()))
            assert "persentase" in first

    def test_grade_computed(self, tokens):
        cls = requests.get(f"{API}/classes", headers=_headers(tokens["admin"])).json()
        cid = cls[0]["id"]
        det = requests.get(f"{API}/classes/{cid}", headers=_headers(tokens["admin"])).json()
        students = det.get("students") or []
        if not students:
            pytest.skip("no students")
        sid = students[0]["id"]
        r = requests.post(f"{API}/grades", headers=_headers(tokens["admin"]),
                          json={"student_id": sid, "class_id": cid, "periode": "TEST_2026-Q1",
                                "komponen": {"tulis": 80, "lisan": 90}})
        assert r.status_code == 200, r.text
        assert abs(r.json()["nilai_akhir"] - 85) < 0.01

    def test_guru_attendance_own_class_only(self, tokens):
        cls = requests.get(f"{API}/classes", headers=_headers(tokens["admin"])).json()
        # find class not taught by guru user
        guru_me = requests.get(f"{API}/auth/me", headers=_headers(tokens["guru"])).json()
        emp_id = guru_me.get("employee_id")
        other = next((c for c in cls if c.get("guru_id") != emp_id), None)
        if not other:
            pytest.skip("no other class")
        r = requests.post(f"{API}/attendance", headers=_headers(tokens["guru"]),
                          json={"class_id": other["id"], "tanggal": date.today().isoformat(), "records": []})
        assert r.status_code == 403


# ---------------- JOB ORDERS ----------------
class TestJobs:
    def test_list_and_candidates(self, tokens):
        r = requests.get(f"{API}/job-orders", headers=_headers(tokens["admin"]))
        assert r.status_code == 200
        rows = r.json()
        if not rows:
            pytest.skip("no job orders")
        jid = rows[0]["id"]
        c = requests.get(f"{API}/job-orders/{jid}/candidates", headers=_headers(tokens["admin"]))
        assert c.status_code == 200
        d = c.json()
        cands = d.get("candidates", [])
        assert isinstance(cands, list)
        if cands:
            first = cands[0]
            assert "checks" in first and "memenuhi" in first


# ---------------- DASHBOARD / MISC ----------------
class TestDashboard:
    @pytest.mark.parametrize("p", ["hari", "minggu", "bulan", "tahun"])
    def test_dashboard(self, tokens, p):
        r = requests.get(f"{API}/dashboard?period={p}", headers=_headers(tokens["owner"]))
        assert r.status_code == 200
        d = r.json()
        assert "siswa" in d and "operasional" in d
        assert "keuangan" in d  # owner has access

    def test_dashboard_guru_no_finance(self, tokens):
        r = requests.get(f"{API}/dashboard?period=bulan", headers=_headers(tokens["guru"]))
        assert r.status_code == 200
        d = r.json()
        assert d.get("keuangan") in (None, {}, False) or "keuangan" not in d

    def test_notifications(self, tokens):
        r = requests.get(f"{API}/notifications", headers=_headers(tokens["admin"]))
        assert r.status_code == 200

    def test_search(self, tokens):
        r = requests.get(f"{API}/search?q=Budi", headers=_headers(tokens["admin"]))
        assert r.status_code == 200


# ---------------- SEARCH RBAC (D-01) ----------------
class TestSearchRBAC:
    Q = "a"

    def _search(self, tokens, role):
        return requests.get(f"{API}/search?q={self.Q}", headers=_headers(tokens[role]), timeout=30)

    @pytest.mark.parametrize("role", ["owner", "admin", "finance", "hr"])
    def test_fin_roles_see_payment(self, tokens, role):
        r = self._search(tokens, role)
        assert r.status_code == 200
        rows = r.json()["students"]
        assert rows, "need seeded students for search test"
        for row in rows:
            assert "pembayaran" in row, f"payment missing for {role}"
            assert "total" in row["pembayaran"] and "bayar" in row["pembayaran"]

    @pytest.mark.parametrize("role", ["guru", "marketing", "staff"])
    def test_nonfin_roles_no_payment(self, tokens, role):
        r = self._search(tokens, role)
        assert r.status_code == 200
        for row in r.json()["students"]:
            assert "pembayaran" not in row, f"payment leaked to {role}"
            assert "total" not in row and "bayar" not in row and "sisa" not in row and "tagihan" not in row
        assert '"pembayaran"' not in r.text

    def test_guru_scoped_to_own_classes(self, tokens):
        me = requests.get(f"{API}/auth/me", headers=_headers(tokens["guru"]), timeout=30).json()
        cls = requests.get(f"{API}/classes", headers=_headers(tokens["guru"]), timeout=30).json()
        own = {c["id"] for c in cls}
        g = self._search(tokens, "guru").json()["students"]
        for row in g:
            d = requests.get(f"{API}/students/{row['id']}", headers=_headers(tokens["owner"]), timeout=30).json()
            assert d.get("class_id") in own, f"guru sees out-of-class student {row['id']}"

    def test_detail_masks_payment(self, tokens):
        oid = self._search(tokens, "owner").json()["students"][0]["id"]
        for role in ["guru", "marketing", "staff"]:
            r = requests.get(f"{API}/students/{oid}", headers=_headers(tokens[role]), timeout=30)
            assert r.status_code in (200, 403), role
            if r.status_code == 200:
                d = r.json()
                assert "pembayaran" not in d and d.get("payments", []) == [], role
        r = requests.get(f"{API}/students/{oid}", headers=_headers(tokens["owner"]), timeout=30)
        assert r.status_code == 200 and "pembayaran" in r.json()

    def test_guru_out_of_class_forbidden(self, tokens):
        o = self._search(tokens, "owner").json()["students"]
        gids = {x["id"] for x in self._search(tokens, "guru").json()["students"]}
        outside = [x for x in o if x["id"] not in gids]
        if not outside:
            pytest.skip("no out-of-class student in seed data")
        r = requests.get(f"{API}/students/{outside[0]['id']}", headers=_headers(tokens["guru"]), timeout=30)
        assert r.status_code == 403


# ---------------- NOTIFICATIONS (D1) ----------------
class TestNotifications:
    def _list(self, tokens, role):
        return requests.get(f"{API}/notifications", headers=_headers(tokens[role]), timeout=30)

    def test_list_and_unread(self, tokens):
        r = self._list(tokens, "owner")
        assert r.status_code == 200
        rows = r.json()
        assert isinstance(rows, list)
        for n in rows:
            assert {"id", "tipe", "judul", "pesan", "read", "link"} <= set(n.keys())
        c = requests.get(f"{API}/notifications/unread-count", headers=_headers(tokens["owner"]), timeout=30)
        assert c.status_code == 200 and c.json()["unread"] == sum(1 for n in rows if not n["read"])

    def test_no_duplicate_on_refresh(self, tokens):
        a = self._list(tokens, "owner").json()
        b = self._list(tokens, "owner").json()
        assert [x["id"] for x in a] == [x["id"] for x in b]
        assert len({x["dedupe_key"] for x in a if x.get("dedupe_key")}) == len([x for x in a if x.get("dedupe_key")])

    def test_role_recipient(self, tokens):
        g = self._list(tokens, "guru").json()
        assert all("pembayaran" != n["tipe"] for n in g), "guru must not receive payment notifications"
        o = self._list(tokens, "owner").json()
        assert isinstance(o, list)

    def test_mark_read(self, tokens):
        rows = self._list(tokens, "owner").json()
        if not rows:
            pytest.skip("no notifications to mark")
        nid = rows[0]["id"]
        assert requests.post(f"{API}/notifications/{nid}/read", headers=_headers(tokens["owner"]), timeout=30).status_code == 200
        rows2 = self._list(tokens, "owner").json()
        assert next(x for x in rows2 if x["id"] == nid)["read"] is True

    def test_wrong_role_cannot_read_others(self, tokens):
        rows = self._list(tokens, "owner").json()
        pay = [x for x in rows if x["tipe"] == "pembayaran"]
        if not pay:
            pytest.skip("no payment notification available")
        r = requests.post(f"{API}/notifications/{pay[0]['id']}/read", headers=_headers(tokens["guru"]), timeout=30)
        assert r.status_code == 404


# ---------------- SEARCH MULTI-ENTITY (D3) ----------------
class TestSearchMulti:
    @pytest.mark.parametrize("role", ["owner", "admin", "finance", "hr", "guru", "marketing", "staff"])
    def test_sections_present(self, tokens, role):
        r = requests.get(f"{API}/search?q=a", headers=_headers(tokens[role]), timeout=30)
        assert r.status_code == 200
        d = r.json()
        assert {"students", "employees", "job_orders"} <= set(d.keys())

    def test_employee_no_salary_leak(self, tokens):
        for role in ["owner", "admin", "finance", "hr", "guru", "marketing", "staff"]:
            r = requests.get(f"{API}/search?q=suci", headers=_headers(tokens[role]), timeout=30)
            for e in r.json().get("employees", []):
                assert set(e.keys()) == {"entity_type", "entity_id", "title", "subtitle", "route"}, e.keys()
            assert "gaji" not in r.text and "tunjangan" not in r.text and "nik" not in r.text.lower()

    def test_job_orders(self, tokens):
        r = requests.get(f"{API}/search?q=manufacturing", headers=_headers(tokens["marketing"]), timeout=30).json()
        assert isinstance(r.get("job_orders"), list)

    def test_limits_and_empty(self, tokens):
        r = requests.get(f"{API}/search?q=a", headers=_headers(tokens["owner"]), timeout=30).json()
        assert len(r["students"]) <= 8 and len(r["employees"]) <= 8 and len(r["job_orders"]) <= 8
        r = requests.get(f"{API}/search?q=zzz-tidak-ada-zzz", headers=_headers(tokens["owner"]), timeout=30).json()
        assert r["students"] == [] and r["employees"] == [] and r["job_orders"] == []
        r = requests.get(f"{API}/search?q=a%25b", headers=_headers(tokens["owner"]), timeout=30)
        assert r.status_code == 200

    def test_guru_scope_and_masking_kept(self, tokens):
        g = requests.get(f"{API}/search?q=a", headers=_headers(tokens["guru"]), timeout=30).json()
        assert all("pembayaran" not in s for s in g["students"])
        assert '"pembayaran"' not in requests.get(f"{API}/search?q=a", headers=_headers(tokens["guru"]), timeout=30).text


# ---------------- REPORTS PAGINATION + INFO-1 (D5) ----------------
class TestReportsPaging:
    def test_students_pagination(self, tokens):
        h = _headers(tokens["admin"])
        full = requests.get(f"{API}/reports/students", headers=h, timeout=60).json()
        assert full["total"] == len(full["rows"]) and full["page"] == 1
        p1 = requests.get(f"{API}/reports/students?page=1&limit=2", headers=h, timeout=60).json()
        assert p1["total"] == full["total"] and len(p1["rows"]) == min(2, full["total"])
        assert p1["per_status"] == full["per_status"]
        if full["total"] > 2:
            p2 = requests.get(f"{API}/reports/students?page=2&limit=2", headers=h, timeout=60).json()
            assert [r["nama"] for r in p2["rows"]] != [r["nama"] for r in p1["rows"]]
            assert set(p1.keys()) >= {"rows", "total", "page", "limit"}

    def test_training_pagination(self, tokens):
        h = _headers(tokens["admin"])
        full = requests.get(f"{API}/reports/training", headers=h, timeout=60).json()
        assert full["total"] == len(full["rows"])
        p1 = requests.get(f"{API}/reports/training?page=1&limit=1", headers=h, timeout=60).json()
        assert p1["total"] == full["total"] and len(p1["rows"]) == 1
        row = p1["rows"][0]
        assert {"hadir", "persentase", "nilai", "ujian"} <= set(row.keys())

    def test_info1_create_masks_payment(self, tokens):
        import uuid
        body = {"nama_lengkap": "TEST Info1", "nik": f"99{uuid.uuid4().hex[:14]}", "jenis_kelamin": "L"}
        r = requests.post(f"{API}/students", headers=_headers(tokens["marketing"]), json=body, timeout=30)
        assert r.status_code == 200, r.text
        assert "pembayaran" not in r.json()
        sid = r.json()["id"]
        r2 = requests.post(f"{API}/students", headers=_headers(tokens["admin"]),
                           json={"nama_lengkap": "TEST Info1b", "nik": f"98{uuid.uuid4().hex[:14]}", "jenis_kelamin": "P"}, timeout=30)
        assert "pembayaran" in r2.json()
        sid2 = r2.json()["id"]
        assert requests.delete(f"{API}/students/{sid}", headers=_headers(tokens["admin"]), timeout=30).status_code == 200
        assert requests.delete(f"{API}/students/{sid2}", headers=_headers(tokens["admin"]), timeout=30).status_code == 200


# ---------------- DASHBOARD PENDING (D4) ----------------
class TestDashboardPending:
    def test_pending_shape_and_roles(self, tokens):
        d = requests.get(f"{API}/dashboard?period=bulan", headers=_headers(tokens["owner"]), timeout=30).json()
        assert "pending" in d
        assert set(d["pending"].keys()) == {"expense", "leave", "payroll"}
        assert d["pending"]["expense"]["count"] >= 0 and d["pending"]["payroll"]["total"] >= 0
        d2 = requests.get(f"{API}/dashboard?period=bulan", headers=_headers(tokens["guru"]), timeout=30).json()
        assert d2.get("pending", {}) == {}
        assert "keuangan" not in d2
        d3 = requests.get(f"{API}/dashboard?period=bulan", headers=_headers(tokens["finance"]), timeout=30).json()
        assert set(d3["pending"].keys()) == {"expense"}
        d4 = requests.get(f"{API}/dashboard?period=bulan", headers=_headers(tokens["hr"]), timeout=30).json()
        assert set(d4["pending"].keys()) == {"leave", "payroll"}

    def test_financial_kpi_unchanged(self, tokens):
        d = requests.get(f"{API}/dashboard?period=bulan", headers=_headers(tokens["owner"]), timeout=30).json()
        assert {"pemasukan", "pengeluaran", "saldo_kas", "piutang"} <= set(d["keuangan"].keys())
        assert {"total", "aktif", "per_status"} <= set(d["siswa"].keys())


# ---------------- EXPORT (D2) ----------------
class TestExport:
    def _dl(self, tokens, role, tab, fmt, extra=""):
        return requests.get(f"{API}/reports/{tab}/export?format={fmt}{extra}", headers=_headers(tokens[role]), timeout=60)

    def test_xlsx_valid(self, tokens):
        r = self._dl(tokens, "owner", "siswa", "xlsx")
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers.get("Content-Type", "")
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        assert ws["A1"].value == "Nama"
        nums = [ws.cell(row=2, column=8).value, ws.cell(row=2, column=9).value]
        assert all(isinstance(v, (int, float)) for v in nums), nums

    def test_pdf_valid(self, tokens):
        r = self._dl(tokens, "owner", "keuangan", "pdf", "&dari=2026-01-01&sampai=2026-12-31")
        assert r.status_code == 200, r.text
        assert r.headers.get("Content-Type", "").startswith("application/pdf")
        assert r.content[:5] == b"%PDF-"

    def test_export_rbac(self, tokens):
        assert self._dl(tokens, "guru", "siswa", "xlsx").status_code == 403
        assert self._dl(tokens, "guru", "keuangan", "pdf").status_code == 403
        assert self._dl(tokens, "guru", "pelatihan", "xlsx").status_code == 200
        assert self._dl(tokens, "marketing", "siswa", "xlsx").status_code == 200
        assert self._dl(tokens, "marketing", "keuangan", "xlsx").status_code == 403

    def test_export_audit(self, tokens):
        self._dl(tokens, "owner", "sdm", "xlsx")
        rows = requests.get(f"{API}/audit-logs?entity=export&limit=5", headers=_headers(tokens["owner"]), timeout=30).json()
        assert any(a["action"] == "export" and a.get("after", {}).get("tab") == "sdm" for a in rows)

    def test_audit_logs(self, tokens):
        r = requests.get(f"{API}/audit-logs?entity=payment", headers=_headers(tokens["owner"]))
        assert r.status_code == 200

    @pytest.mark.parametrize("kind", ["students", "finance", "hr", "training"])
    def test_reports(self, tokens, kind):
        r = requests.get(f"{API}/reports/{kind}", headers=_headers(tokens["owner"]))
        assert r.status_code == 200


# ---------------- HR ATTENDANCE ----------------
class TestHrAttendance:
    TEST_DATE = "2020-01-05"  # tanggal uji terisolasi, di luar periode operasional

    @pytest.fixture(scope="class")
    def emp_ids(self, tokens):
        r = requests.get(f"{API}/employees", headers=_headers(tokens["hr"]), timeout=30)
        assert r.status_code == 200
        ids = [e["id"] for e in r.json() if e.get("aktif", True)][:2]
        assert len(ids) == 2, "butuh minimal 2 employee aktif untuk uji absensi"
        return ids

    def _batch(self, emp_ids, status="hadir"):
        return {"tanggal": self.TEST_DATE, "records": [
            {"employee_id": eid, "status": status, "jam_masuk": "08:00",
             "jam_pulang": "17:00", "keterangan": "TEST"} for eid in emp_ids]}

    def _post(self, tokens, role, body):
        return requests.post(f"{API}/hr/attendance", headers=_headers(tokens[role]), json=body, timeout=30)

    def test_owner_post_and_idempotent(self, tokens, emp_ids):
        assert self._post(tokens, "owner", self._batch(emp_ids)).status_code == 200
        assert self._post(tokens, "owner", self._batch(emp_ids)).status_code == 200
        r = requests.get(f"{API}/hr/attendance?dari={self.TEST_DATE}&sampai={self.TEST_DATE}",
                         headers=_headers(tokens["owner"]), timeout=30)
        assert r.status_code == 200
        got = {(x["employee_id"], x["tanggal"]) for x in r.json()}
        assert len(got) == 2, f"duplikat terdeteksi: {r.json()}"

    def test_hr_post(self, tokens, emp_ids):
        assert self._post(tokens, "hr", self._batch(emp_ids)).status_code == 200

    def test_admin_post_forbidden(self, tokens, emp_ids):
        assert self._post(tokens, "admin", self._batch(emp_ids)).status_code == 403

    @pytest.mark.parametrize("role", ["guru", "finance", "marketing", "staff"])
    def test_write_roles_forbidden(self, tokens, emp_ids, role):
        assert self._post(tokens, role, self._batch(emp_ids)).status_code == 403

    @pytest.mark.parametrize("role", ["owner", "hr", "admin"])
    def test_read_allowed(self, tokens, role):
        r = requests.get(f"{API}/hr/attendance?dari={self.TEST_DATE}&sampai={self.TEST_DATE}",
                         headers=_headers(tokens[role]), timeout=30)
        assert r.status_code == 200

    def test_invalid_employee(self, tokens):
        r = self._post(tokens, "owner", {"tanggal": self.TEST_DATE, "records": [
            {"employee_id": "no-such-id", "status": "hadir"}]})
        assert r.status_code == 404

    def test_invalid_status(self, tokens, emp_ids):
        r = self._post(tokens, "owner", {"tanggal": self.TEST_DATE, "records": [
            {"employee_id": emp_ids[0], "status": "liburan"}]})
        assert r.status_code == 400

    def test_future_date(self, tokens, emp_ids):
        future = "2999-01-01"
        r = self._post(tokens, "owner", {"tanggal": future, "records": [
            {"employee_id": emp_ids[0], "status": "hadir"}]})
        assert r.status_code == 400

    def test_invalid_time(self, tokens, emp_ids):
        r = self._post(tokens, "owner", {"tanggal": self.TEST_DATE, "records": [
            {"employee_id": emp_ids[0], "status": "hadir", "jam_masuk": "25:99"}]})
        assert r.status_code == 400

    def test_recap(self, tokens, emp_ids):
        r = requests.get(f"{API}/hr/attendance/recap?dari={self.TEST_DATE}&sampai={self.TEST_DATE}",
                         headers=_headers(tokens["hr"]), timeout=30)
        assert r.status_code == 200
        by_id = {x["employee_id"]: x for x in r.json()}
        for eid in emp_ids:
            row = by_id[eid]
            assert row["total"] == 1 and row["hadir"] == 1 and row["persentase"] == 100.0


# ---------------- LEAVE / CUTI ----------------
class TestLeave:
    D1, D2 = "2020-03-01", "2020-03-05"  # rentang uji utama (E1)
    D3, D4 = "2020-03-10", "2020-03-11"  # uji durasi / reject (E2)

    @pytest.fixture(scope="class")
    def emp_ids(self, tokens):
        r = requests.get(f"{API}/employees", headers=_headers(tokens["hr"]), timeout=30)
        assert r.status_code == 200
        ids = [e["id"] for e in r.json() if e.get("aktif", True)][:2]
        assert len(ids) == 2, "butuh minimal 2 employee aktif untuk uji cuti"
        return ids

    def _post(self, tokens, role, body):
        return requests.post(f"{API}/leaves", headers=_headers(tokens[role]), json=body, timeout=30)

    def _decide(self, tokens, role, lid, setuju, alasan=""):
        return requests.put(f"{API}/leaves/{lid}/decide", headers=_headers(tokens[role]),
                            json={"setuju": setuju, "alasan": alasan}, timeout=30)

    def _main_id(self, tokens):
        r = requests.get(f"{API}/leaves?status=menunggu", headers=_headers(tokens["owner"]), timeout=30)
        for x in r.json():
            if x["dari"] == self.D1 and x["sampai"] == self.D2:
                return x["id"]
        return None

    def test_create_valid_menunggu(self, tokens, emp_ids):
        r = self._post(tokens, "owner", {"employee_id": emp_ids[0], "jenis": "Cuti",
                                        "dari": self.D1, "sampai": self.D2, "alasan": "TEST"})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["status"] == "menunggu" and d["durasi_hari"] == 5

    def test_duration(self, tokens, emp_ids):
        r = self._post(tokens, "hr", {"employee_id": emp_ids[1], "jenis": "Cuti",
                                     "dari": "2020-03-06", "sampai": "2020-03-06", "alasan": "TEST"})
        assert r.json()["durasi_hari"] == 1
        r = self._post(tokens, "hr", {"employee_id": emp_ids[1], "jenis": "Cuti",
                                     "dari": self.D3, "sampai": self.D4, "alasan": "TEST"})
        assert r.json()["durasi_hari"] == 2

    def test_invalid_employee(self, tokens):
        r = self._post(tokens, "owner", {"employee_id": "no-such-id", "jenis": "Cuti",
                                        "dari": self.D1, "sampai": self.D2, "alasan": "x"})
        assert r.status_code == 404

    def test_invalid_date(self, tokens, emp_ids):
        r = self._post(tokens, "owner", {"employee_id": emp_ids[0], "jenis": "Cuti",
                                        "dari": self.D2, "sampai": self.D1, "alasan": "x"})
        assert r.status_code == 400

    def test_overlap_pending(self, tokens, emp_ids):
        r = self._post(tokens, "owner", {"employee_id": emp_ids[0], "jenis": "Cuti",
                                        "dari": "2020-03-03", "sampai": "2020-03-07", "alasan": "x"})
        assert r.status_code == 409

    @pytest.mark.parametrize("role", ["owner", "hr", "admin"])
    def test_read_allowed(self, tokens, role):
        assert requests.get(f"{API}/leaves", headers=_headers(tokens[role]), timeout=30).status_code == 200

    @pytest.mark.parametrize("role", ["guru", "finance", "marketing", "staff"])
    def test_read_forbidden(self, tokens, role):
        assert requests.get(f"{API}/leaves", headers=_headers(tokens[role]), timeout=30).status_code == 403

    def test_create_rbac(self, tokens, emp_ids):
        ok = {"tanggal": "2020-04-10", "x": 0}
        for role, code in [("owner", 200), ("hr", 200), ("admin", 200)]:
            d = f"2020-04-{10 + ok['x']}"
            r = self._post(tokens, role, {"employee_id": emp_ids[0], "jenis": "Cuti", "dari": d, "sampai": d, "alasan": "TEST"})
            assert r.status_code == code, f"{role}: {r.status_code} {r.text}"
            ok["x"] += 1
        for role in ["guru", "finance", "marketing", "staff"]:
            r = self._post(tokens, role, {"employee_id": emp_ids[0], "jenis": "Cuti",
                                         "dari": "2020-04-20", "sampai": "2020-04-20", "alasan": "x"})
            assert r.status_code == 403, role

    def test_approve_marks_attendance(self, tokens, emp_ids):
        lid = self._main_id(tokens)
        assert lid, "leave utama tidak ditemukan"
        r = self._decide(tokens, "hr", lid, True)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["status"] == "disetujui" and d["attendance_marked"] is True
        assert d["approver_id"] and d["decided_at"]
        a = requests.get(f"{API}/hr/attendance?dari={self.D1}&sampai={self.D2}&employee_id={emp_ids[0]}",
                         headers=_headers(tokens["hr"]), timeout=30).json()
        assert len(a) == 5 and all(x["status"] == "cuti" for x in a)
        c = requests.get(f"{API}/hr/attendance/recap?dari={self.D1}&sampai={self.D2}",
                         headers=_headers(tokens["hr"]), timeout=30).json()
        row = next(x for x in c if x["employee_id"] == emp_ids[0])
        assert row["cuti"] == 5

    def test_overlap_approved(self, tokens, emp_ids):
        r = self._post(tokens, "owner", {"employee_id": emp_ids[0], "jenis": "Cuti",
                                        "dari": "2020-03-04", "sampai": "2020-03-06", "alasan": "x"})
        assert r.status_code == 409

    def test_double_decision(self, tokens):
        r = requests.get(f"{API}/leaves?status=disetujui", headers=_headers(tokens["owner"]), timeout=30)
        lid = next(x["id"] for x in r.json() if x["dari"] == self.D1)
        assert self._decide(tokens, "hr", lid, True).status_code == 400
        assert self._decide(tokens, "hr", lid, False, "x").status_code == 400

    def test_reject_flow(self, tokens, emp_ids):
        r = self._post(tokens, "owner", {"employee_id": emp_ids[1], "jenis": "Cuti",
                                        "dari": "2020-03-12", "sampai": "2020-03-13", "alasan": "TEST"})
        assert r.status_code == 200, r.text
        lid = r.json()["id"]
        assert self._decide(tokens, "hr", lid, False, "").status_code == 400
        r = self._decide(tokens, "hr", lid, False, "Personil kurang")
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["status"] == "ditolak" and d["reject_reason"] == "Personil kurang"
        a = requests.get(f"{API}/hr/attendance?dari=2020-03-12&sampai=2020-03-13&employee_id={emp_ids[1]}",
                         headers=_headers(tokens["hr"]), timeout=30).json()
        assert all(x["status"] != "cuti" for x in a)
        r = self._post(tokens, "owner", {"employee_id": emp_ids[1], "jenis": "Cuti",
                                        "dari": "2020-03-12", "sampai": "2020-03-13", "alasan": "TEST ulang"})
        assert r.status_code == 200, "overlap dengan yang ditolak harus lolos"

    def test_decide_rbac(self, tokens, emp_ids):
        r = self._post(tokens, "owner", {"employee_id": emp_ids[0], "jenis": "Cuti",
                                        "dari": "2020-05-01", "sampai": "2020-05-02", "alasan": "TEST"})
        lid = r.json()["id"]
        for role in ["admin", "guru", "finance", "marketing", "staff"]:
            assert self._decide(tokens, role, lid, True).status_code == 403, role
        assert self._decide(tokens, "owner", lid, True).status_code == 200


# ---------------- PAYROLL & HONOR ----------------
class TestPayroll:
    P1, P2, P3 = "2020-07", "2020-08", "2020-09"  # bulanan, alfa, honor
    P4, P5 = "2020-10", "2020-11"  # approve/pay, snapshot

    @pytest.fixture(scope="class")
    def staff(self, tokens):
        hdr = _headers(tokens["hr"])
        mk = lambda body: requests.post(f"{API}/employees", headers=hdr, json=body, timeout=30).json()
        kar = mk({"nama": "TEST_Gaji", "tipe": "karyawan", "jabatan": "Staff", "tanggal_masuk": "2020-01-01",
                  "status_kerja": "tetap", "gaji_pokok": 4000000, "tunjangan": 500000, "aktif": True})
        gur = mk({"nama": "TEST_Honor", "tipe": "guru", "jabatan": "Pengajar", "tanggal_masuk": "2020-01-01",
                  "status_kerja": "freelance", "gaji_pokok": 0, "tunjangan": 0,
                  "honor_per_pertemuan": 150000, "aktif": True})
        yield kar, gur
        for e in (kar, gur):
            requests.delete(f"{API}/employees/{e['id']}", headers=hdr, timeout=30)

    def _calc(self, tokens, role, body):
        return requests.post(f"{API}/payrolls/calculate", headers=_headers(tokens[role]), json=body, timeout=30)

    def test_invalid_employee(self, tokens):
        r = self._calc(tokens, "hr", {"periode": self.P1, "items": [{"employee_id": "no-such"}]})
        assert r.status_code == 404

    def test_invalid_periode(self, tokens, staff):
        r = self._calc(tokens, "hr", {"periode": "2020-13", "employee_ids": [staff[0]["id"]]})
        assert r.status_code == 400

    def test_monthly_calc(self, tokens, staff):
        kar = staff[0]
        r = self._calc(tokens, "hr", {"periode": self.P1, "employee_ids": [kar["id"]]})
        assert r.status_code == 200, r.text
        p = r.json()["rows"][0]
        assert p["status"] == "draft" and p["snapshot"]["model"] == "bulanan"
        assert p["snapshot"]["gaji_pokok"] == kar["gaji_pokok"]
        assert p["bruto"] == kar["gaji_pokok"] + kar["tunjangan"]
        assert p["bersih"] == p["bruto"] and p["no_slip"].startswith("SLIP-")

    def test_duplicate_skipped(self, tokens, staff):
        r = self._calc(tokens, "hr", {"periode": self.P1, "employee_ids": [staff[0]["id"]]})
        assert r.status_code == 200
        assert r.json()["rows"] == [] and staff[0]["id"] in r.json()["skipped"]
        allp = requests.get(f"{API}/payrolls?periode={self.P1}", headers=_headers(tokens["hr"]), timeout=30).json()
        assert sum(1 for x in allp if x["employee_id"] == staff[0]["id"]) == 1

    def test_alfa_deduction(self, tokens, staff):
        kar = staff[0]
        for t in ("2020-08-03", "2020-08-04"):
            r = requests.post(f"{API}/hr/attendance", headers=_headers(tokens["hr"]), timeout=30,
                              json={"tanggal": t, "records": [{"employee_id": kar["id"], "status": "alfa"}]})
            assert r.status_code == 200
        r = self._calc(tokens, "hr", {"periode": self.P2, "employee_ids": [kar["id"]]})
        p = r.json()["rows"][0]
        dim = 31
        pot = round(kar["gaji_pokok"] / dim * 2)
        assert p["komponen"]["attendance"]["alfa"] == 2 and p["komponen"]["potongan_alfa"] == pot
        assert p["bersih"] == p["bruto"] - pot

    def test_honor_b1(self, tokens, staff):
        gur = staff[1]
        r = self._calc(tokens, "hr", {"periode": self.P3, "items": [{"employee_id": gur["id"], "honor_pertemuan": 10}]})
        assert r.status_code == 200, r.text
        p = r.json()["rows"][0]
        assert p["snapshot"]["model"] == "honor"
        assert p["komponen"]["honor_total"] == gur["honor_per_pertemuan"] * 10
        assert p["bersih"] == p["komponen"]["honor_total"] + p["komponen"]["tunjangan_hitung"]

    def test_honor_requires_meetings(self, tokens, staff):
        r = self._calc(tokens, "hr", {"periode": "2020-12", "items": [{"employee_id": staff[1]["id"]}]})
        assert r.status_code == 400

    def test_snapshot_frozen(self, tokens, staff):
        kar = staff[0]
        r = self._calc(tokens, "hr", {"periode": self.P5, "employee_ids": [kar["id"]]})
        old_gaji = r.json()["rows"][0]["snapshot"]["gaji_pokok"]
        full = requests.get(f"{API}/employees", headers=_headers(tokens["hr"]), timeout=30).json()
        emp = next(e for e in full if e["id"] == kar["id"])
        emp["gaji_pokok"] = old_gaji + 1000000
        emp.pop("kelas", None)
        ur = requests.put(f"{API}/employees/{kar['id']}", headers=_headers(tokens["hr"]), json=emp, timeout=30)
        assert ur.status_code == 200
        r2 = self._calc(tokens, "hr", {"periode": "2020-12", "employee_ids": [kar["id"]]})
        assert r2.json()["rows"][0]["snapshot"]["gaji_pokok"] == old_gaji + 1000000
        back = requests.get(f"{API}/payrolls?periode={self.P5}", headers=_headers(tokens["hr"]), timeout=30).json()
        assert back[0]["snapshot"]["gaji_pokok"] == old_gaji, "payroll lama ikut berubah!"
        emp["gaji_pokok"] = old_gaji
        assert requests.put(f"{API}/employees/{kar['id']}", headers=_headers(tokens["hr"]), json=emp, timeout=30).status_code == 200

    def test_edit_draft_needs_alasan(self, tokens, staff):
        allp = requests.get(f"{API}/payrolls?periode={self.P5}", headers=_headers(tokens["hr"]), timeout=30).json()
        pid = allp[0]["id"]
        assert requests.put(f"{API}/payrolls/{pid}", headers=_headers(tokens["hr"]), timeout=30,
                            json={"bonus": 100000}).status_code == 400
        r = requests.put(f"{API}/payrolls/{pid}", headers=_headers(tokens["hr"]), timeout=30,
                         json={"bonus": 100000, "potongan": [{"jenis": "Kasbon", "nominal": 50000}], "alasan": "TEST"})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["komponen"]["bonus"] == 100000 and d["bersih"] == d["bruto"] - 50000 - d["komponen"]["potongan_alfa"]

    def test_approve_flow(self, tokens, staff):
        kar = staff[0]
        r = self._calc(tokens, "hr", {"periode": self.P4, "employee_ids": [kar["id"]]})
        pid = r.json()["rows"][0]["id"]
        assert requests.post(f"{API}/payrolls/{pid}/approve", headers=_headers(tokens["admin"]), timeout=30).status_code == 403
        assert requests.post(f"{API}/payrolls/{pid}/pay", headers=_headers(tokens["finance"]), timeout=30,
                             json={"account_id": "x"}).status_code == 400
        assert requests.post(f"{API}/payrolls/{pid}/approve", headers=_headers(tokens["owner"]), timeout=30).status_code == 200
        d = requests.get(f"{API}/payrolls/{pid}", headers=_headers(tokens["owner"]), timeout=30).json()
        assert d["status"] == "disetujui" and d["approved_by"]
        assert requests.post(f"{API}/payrolls/{pid}/approve", headers=_headers(tokens["owner"]), timeout=30).status_code == 400
        assert requests.put(f"{API}/payrolls/{pid}", headers=_headers(tokens["hr"]), timeout=30,
                            json={"bonus": 1, "alasan": "x"}).status_code == 400

    def test_pay_idempotent(self, tokens, staff):
        allp = requests.get(f"{API}/payrolls?periode={self.P4}", headers=_headers(tokens["owner"]), timeout=30).json()
        pid = allp[0]["id"]
        acc = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"]), timeout=30).json()[0]
        assert requests.post(f"{API}/payrolls/{pid}/pay", headers=_headers(tokens["admin"]), timeout=30,
                             json={"account_id": acc["id"]}).status_code == 403
        r = requests.post(f"{API}/payrolls/{pid}/pay", headers=_headers(tokens["finance"]), timeout=30,
                          json={"account_id": acc["id"]})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["status"] == "dibayar" and d["transaction_id"]
        txs = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30).json()
        mine = [t for t in txs if t.get("ref_type") == "payroll" and t.get("ref_id") == pid]
        assert len(mine) == 1 and mine[0]["nominal"] == d["bersih"] and mine[0]["kategori"] == "Gaji"
        r2 = requests.post(f"{API}/payrolls/{pid}/pay", headers=_headers(tokens["finance"]), timeout=30,
                           json={"account_id": acc["id"]})
        assert r2.status_code == 200 and r2.json().get("already_paid") is True
        txs2 = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30).json()
        assert sum(1 for t in txs2 if t.get("ref_type") == "payroll" and t.get("ref_id") == pid) == 1

    @pytest.mark.parametrize("role", ["owner", "hr", "admin", "finance"])
    def test_read_allowed(self, tokens, role):
        assert requests.get(f"{API}/payrolls?periode={self.P1}", headers=_headers(tokens[role]), timeout=30).status_code == 200

    @pytest.mark.parametrize("role", ["guru", "marketing", "staff"])
    def test_read_forbidden(self, tokens, role):
        assert requests.get(f"{API}/payrolls?periode={self.P1}", headers=_headers(tokens[role]), timeout=30).status_code == 403

    def test_calc_forbidden_roles(self, tokens, staff):
        for role in ["admin", "finance", "guru", "marketing", "staff"]:
            r = self._calc(tokens, role, {"periode": self.P1, "employee_ids": [staff[0]["id"]]})
            assert r.status_code == 403, role

    def test_audit_trail(self, tokens):
        rows = requests.get(f"{API}/audit-logs?entity=payroll&limit=50", headers=_headers(tokens["owner"]), timeout=30).json()
        acts = {r["action"] for r in rows}
        assert {"create", "update", "approve", "pay"} <= acts


# ---------------- PAYROLL CORRECTIVE (F1/F2/F3/F5) ----------------
class TestPayrollCorrective:
    P1, P2, P3 = "2021-01", "2021-02", "2021-03"

    @pytest.fixture(scope="class")
    def staff(self, tokens):
        hdr = _headers(tokens["hr"])
        mk = lambda body: requests.post(f"{API}/employees", headers=hdr, json=body, timeout=30).json()
        joiner = mk({"nama": "TEST_Join", "tipe": "karyawan", "jabatan": "Staff", "tanggal_masuk": "2021-01-11",
                     "status_kerja": "tetap", "gaji_pokok": 3000000, "tunjangan": 300000, "aktif": True})
        exiter = mk({"nama": "TEST_Exit", "tipe": "karyawan", "jabatan": "Staff", "tanggal_masuk": "2020-01-01",
                     "kontrak_berakhir": "2021-01-20", "status_kerja": "kontrak",
                     "gaji_pokok": 3000000, "tunjangan": 300000, "aktif": True})
        gsal = mk({"nama": "TEST_GuruGaji", "tipe": "guru", "jabatan": "Pengajar", "tanggal_masuk": "2020-01-01",
                   "status_kerja": "tetap", "gaji_pokok": 3000000, "tunjangan": 0,
                   "honor_per_pertemuan": 0, "aktif": True})
        yield joiner, exiter, gsal
        for e in (joiner, exiter, gsal):
            requests.delete(f"{API}/employees/{e['id']}", headers=hdr, timeout=30)

    def _calc(self, tokens, role, body):
        return requests.post(f"{API}/payrolls/calculate", headers=_headers(tokens[role]), json=body, timeout=30)

    def _att(self, tokens, tanggal, eid, status="alfa"):
        return requests.post(f"{API}/hr/attendance", headers=_headers(tokens["hr"]), timeout=30,
                             json={"tanggal": tanggal, "records": [{"employee_id": eid, "status": status}]})

    def test_f1_alfa_join_window(self, tokens, staff):
        joiner = staff[0]
        assert self._att(tokens, "2021-01-10", joiner["id"]).status_code == 200
        assert self._att(tokens, "2021-01-11", joiner["id"]).status_code == 200
        r = self._calc(tokens, "hr", {"periode": self.P1, "employee_ids": [joiner["id"]]})
        assert r.status_code == 200, r.text
        p = r.json()["rows"][0]
        assert p["komponen"]["attendance"]["alfa"] == 1, f"alfa luar window ikut dihitung: {p['komponen']['attendance']}"
        assert p["komponen"]["hari_aktif"] == 21
        assert p["komponen"]["potongan_alfa"] == round(3000000 / 31 * 1)
        assert p["komponen"]["tunjangan_hitung"] == round(300000 / 31 * 21)
        assert p["komponen"]["base"] == round(3000000 / 31 * 21)

    def test_f1_alfa_exit_window(self, tokens, staff):
        exiter = staff[1]
        assert self._att(tokens, "2021-01-20", exiter["id"]).status_code == 200
        assert self._att(tokens, "2021-01-21", exiter["id"]).status_code == 200
        r = self._calc(tokens, "hr", {"periode": self.P1, "employee_ids": [exiter["id"]]})
        p = r.json()["rows"][0]
        assert p["komponen"]["attendance"]["alfa"] == 1
        assert p["komponen"]["hari_aktif"] == 20

    def test_f5_monthly_guru_category(self, tokens, staff):
        gsal = staff[2]
        r = self._calc(tokens, "hr", {"periode": self.P1, "employee_ids": [gsal["id"]]})
        pid = r.json()["rows"][0]["id"]
        assert requests.post(f"{API}/payrolls/{pid}/approve", headers=_headers(tokens["owner"]), timeout=30).status_code == 200
        acc = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"]), timeout=30).json()[0]
        r = requests.post(f"{API}/payrolls/{pid}/pay", headers=_headers(tokens["finance"]), timeout=30,
                          json={"account_id": acc["id"]})
        assert r.status_code == 200, r.text
        txs = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30).json()
        mine = [t for t in txs if t.get("ref_type") == "payroll" and t.get("ref_id") == pid]
        assert len(mine) == 1 and mine[0]["kategori"] == "Gaji"

    def _zero_neg_ids(self, tokens, staff):
        gsal = staff[2]
        ids = []
        for tag, periode, extra in (("nol", self.P2, 0), ("neg", "2021-04", 100000)):
            prev = self._calc(tokens, "hr", {"periode": periode, "employee_ids": [gsal["id"]], "preview": True})
            bruto = prev.json()["rows"][0]["bruto"]
            r = self._calc(tokens, "hr", {"periode": periode, "items": [
                {"employee_id": gsal["id"], "potongan": [{"jenis": "TEST", "nominal": bruto + extra,
                                                          "keterangan": "QA f3"}]}]})
            assert r.status_code == 200 and r.json()["rows"], f"{tag}: {r.status_code} {r.text}"
            ids.append((tag, r.json()["rows"][0]["id"]))
        return ids

    def test_f3_net_nonpositif_rejected(self, tokens, staff):
        for tag, pid in self._zero_neg_ids(tokens, staff):
            assert requests.post(f"{API}/payrolls/{pid}/approve", headers=_headers(tokens["owner"]), timeout=30).status_code == 200
            r = requests.post(f"{API}/payrolls/{pid}/pay", headers=_headers(tokens["finance"]), timeout=30,
                              json={"account_id": requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"]), timeout=30).json()[0]["id"]})
            assert r.status_code == 400, f"{tag}: {r.status_code} {r.text}"
            d = requests.get(f"{API}/payrolls/{pid}", headers=_headers(tokens["owner"]), timeout=30).json()
            assert d["status"] == "disetujui" and not d.get("transaction_id"), tag
            txs = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30).json()
            assert sum(1 for t in txs if t.get("ref_type") == "payroll" and t.get("ref_id") == pid) == 0, tag

    def test_f2_concurrent_pay(self, tokens, staff):
        import concurrent.futures
        gsal = staff[2]
        r = self._calc(tokens, "hr", {"periode": self.P3, "employee_ids": [gsal["id"]]})
        pid = r.json()["rows"][0]["id"]
        assert requests.post(f"{API}/payrolls/{pid}/approve", headers=_headers(tokens["owner"]), timeout=30).status_code == 200
        acc = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"]), timeout=30).json()[0]
        hdr = _headers(tokens["finance"])
        def one(_):
            return requests.post(f"{API}/payrolls/{pid}/pay", headers=hdr, json={"account_id": acc["id"]}, timeout=60).status_code
        with concurrent.futures.ThreadPoolExecutor(max_workers=6) as ex:
            codes = list(ex.map(one, range(6)))
        assert all(c == 200 for c in codes), codes
        txs = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30).json()
        assert sum(1 for t in txs if t.get("ref_type") == "payroll" and t.get("ref_id") == pid) == 1
        d = requests.get(f"{API}/payrolls/{pid}", headers=_headers(tokens["owner"]), timeout=30).json()
        assert d["status"] == "dibayar" and d["transaction_id"]


# ---------------- EXPENSE APPROVAL ----------------
class TestExpense:
    @pytest.fixture(scope="class")
    def acc(self, tokens):
        a = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"]), timeout=30).json()
        assert a, "need seeded accounts"
        return a[0]["id"]

    def _mk(self, tokens, role, nominal, **kw):
        body = {"judul": "TEST expense", "kategori": "ATK", "nominal": nominal,
                "tanggal": date.today().isoformat(), "deskripsi": "TEST", "account_id": kw.pop("account_id", ""),
                **kw}
        return requests.post(f"{API}/expenses", headers=_headers(tokens[role]), json=body, timeout=30)

    def _lifecycle(self, tokens, nominal, acc):
        r = self._mk(tokens, "staff", nominal, account_id=acc)
        assert r.status_code == 200, r.text
        eid = r.json()["id"]
        assert r.json()["status"] == "draft"
        assert requests.post(f"{API}/expenses/{eid}/submit", headers=_headers(tokens["staff"]), timeout=30).status_code == 200
        return eid

    def test_full_lifecycle_small(self, tokens, acc):
        eid = self._lifecycle(tokens, 500000, acc)
        assert requests.post(f"{API}/expenses/{eid}/decide", headers=_headers(tokens["finance"]), timeout=30,
                             json={"setuju": True}).status_code == 200
        r = requests.post(f"{API}/expenses/{eid}/pay", headers=_headers(tokens["finance"]), timeout=30,
                          json={"account_id": acc})
        assert r.status_code == 200, r.text
        assert r.json()["status"] == "dibayar"
        txs = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30).json()
        mine = [t for t in txs if t.get("ref_type") == "expense" and t.get("ref_id") == eid]
        assert len(mine) == 1 and mine[0]["nominal"] == 500000 and mine[0]["kategori"] == "ATK"

    def test_threshold_boundary(self, tokens, acc):
        e1 = self._lifecycle(tokens, 999999, acc)
        assert requests.post(f"{API}/expenses/{e1}/decide", headers=_headers(tokens["finance"]), timeout=30,
                             json={"setuju": True}).status_code == 200
        e2 = self._lifecycle(tokens, 1000000, acc)
        assert requests.post(f"{API}/expenses/{e2}/decide", headers=_headers(tokens["finance"]), timeout=30,
                             json={"setuju": True}).status_code == 403
        assert requests.post(f"{API}/expenses/{e2}/decide", headers=_headers(tokens["owner"]), timeout=30,
                             json={"setuju": True}).status_code == 200
        e3 = self._lifecycle(tokens, 1000001, acc)
        assert requests.post(f"{API}/expenses/{e3}/decide", headers=_headers(tokens["finance"]), timeout=30,
                             json={"setuju": True}).status_code == 403

    def test_reject_flow(self, tokens, acc):
        eid = self._lifecycle(tokens, 300000, acc)
        r = requests.post(f"{API}/expenses/{eid}/decide", headers=_headers(tokens["finance"]), timeout=30,
                          json={"setuju": False, "alasan": ""})
        assert r.status_code == 400
        r = requests.post(f"{API}/expenses/{eid}/decide", headers=_headers(tokens["finance"]), timeout=30,
                          json={"setuju": False, "alasan": "TEST tolak"})
        assert r.status_code == 200 and r.json()["status"] == "ditolak"
        assert requests.post(f"{API}/expenses/{eid}/pay", headers=_headers(tokens["finance"]), timeout=30,
                             json={"account_id": acc}).status_code == 400

    def test_invalid_transitions(self, tokens, acc):
        r = self._mk(tokens, "staff", 100000, account_id=acc)
        eid = r.json()["id"]
        assert requests.post(f"{API}/expenses/{eid}/pay", headers=_headers(tokens["finance"]), timeout=30,
                             json={"account_id": acc}).status_code == 400
        assert requests.post(f"{API}/expenses/{eid}/decide", headers=_headers(tokens["finance"]), timeout=30,
                             json={"setuju": True}).status_code == 400

    def test_double_pay(self, tokens, acc):
        eid = self._lifecycle(tokens, 200000, acc)
        requests.post(f"{API}/expenses/{eid}/decide", headers=_headers(tokens["finance"]), timeout=30,
                      json={"setuju": True})
        hdr = _headers(tokens["finance"])
        assert requests.post(f"{API}/expenses/{eid}/pay", headers=hdr, json={"account_id": acc}, timeout=30).status_code == 200
        r2 = requests.post(f"{API}/expenses/{eid}/pay", headers=hdr, json={"account_id": acc}, timeout=30)
        assert r2.status_code == 200 and r2.json().get("already_paid") is True
        txs = requests.get(f"{API}/finance/transactions", headers=hdr, timeout=30).json()
        assert sum(1 for t in txs if t.get("ref_type") == "expense" and t.get("ref_id") == eid) == 1

    def test_concurrent_pay(self, tokens, acc):
        import concurrent.futures
        eid = self._lifecycle(tokens, 150000, acc)
        requests.post(f"{API}/expenses/{eid}/decide", headers=_headers(tokens["finance"]), timeout=30,
                      json={"setuju": True})
        hdr = _headers(tokens["finance"])
        def one(_):
            return requests.post(f"{API}/expenses/{eid}/pay", headers=hdr, json={"account_id": acc}, timeout=60).status_code
        with concurrent.futures.ThreadPoolExecutor(max_workers=6) as ex:
            codes = list(ex.map(one, range(6)))
        assert all(c == 200 for c in codes), codes
        txs = requests.get(f"{API}/finance/transactions", headers=hdr, timeout=30).json()
        assert sum(1 for t in txs if t.get("ref_type") == "expense" and t.get("ref_id") == eid) == 1

    def test_expense_rbac(self, tokens, acc):
        r = self._mk(tokens, "staff", 100000, account_id=acc)
        eid = r.json()["id"]
        for role in ["hr", "guru"]:
            assert requests.get(f"{API}/expenses", headers=_headers(tokens[role]), timeout=30).status_code == 403, role
            assert self._mk(tokens, role, 100000, account_id=acc).status_code == 403, role
        for role in ["owner", "admin", "finance", "staff", "marketing"]:
            assert requests.get(f"{API}/expenses", headers=_headers(tokens[role]), timeout=30).status_code == 200, role
            if role in ("owner", "finance", "staff", "marketing"):
                assert self._mk(tokens, role, 50000, account_id=acc).status_code == 200, role
        assert requests.post(f"{API}/expenses/{eid}/decide", headers=_headers(tokens["admin"]), timeout=30,
                             json={"setuju": True}).status_code == 403
        assert requests.post(f"{API}/expenses/{eid}/pay", headers=_headers(tokens["admin"]), timeout=30,
                             json={"account_id": acc}).status_code == 403
        assert self._mk(tokens, "admin", 100000, account_id=acc).status_code == 403
        assert self._mk(tokens, "guru", 100000, account_id=acc).status_code == 403

    def test_draft_not_in_ledger(self, tokens, acc):
        before = requests.get(f"{API}/finance/summary?period=tahun", headers=_headers(tokens["finance"]), timeout=30).json()["pengeluaran"]
        self._mk(tokens, "staff", 777000, account_id=acc)
        after = requests.get(f"{API}/finance/summary?period=tahun", headers=_headers(tokens["finance"]), timeout=30).json()["pengeluaran"]
        assert after == before

    def test_category_validation(self, tokens, acc):
        r = self._mk(tokens, "staff", 100000, account_id=acc, kategori="Ngawur")
        assert r.status_code == 400
        r = requests.post(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30,
                          json={"jenis": "pengeluaran", "kategori": "Ngawur", "nominal": 1000,
                                "tanggal": date.today().isoformat(), "account_id": acc})
        assert r.status_code == 400

    def test_linked_tx_guards(self, tokens, acc):
        eid = self._lifecycle(tokens, 120000, acc)
        requests.post(f"{API}/expenses/{eid}/decide", headers=_headers(tokens["finance"]), timeout=30, json={"setuju": True})
        requests.post(f"{API}/expenses/{eid}/pay", headers=_headers(tokens["finance"]), timeout=30, json={"account_id": acc})
        txs = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30).json()
        tx = next(t for t in txs if t.get("ref_type") == "expense" and t.get("ref_id") == eid)
        hdr = _headers(tokens["finance"])
        assert requests.put(f"{API}/finance/transactions/{tx['id']}", headers=hdr, timeout=30,
                            json={"jenis": "pengeluaran", "kategori": "ATK", "nominal": 1,
                                  "tanggal": date.today().isoformat(), "account_id": acc, "alasan": "x"}).status_code == 400
        assert requests.delete(f"{API}/finance/transactions/{tx['id']}", headers=_headers(tokens["owner"]), timeout=30).status_code == 400


# ---------------- PAYMENT IDEMPOTENCY ----------------
class TestPaymentIdem:
    @pytest.fixture(scope="class")
    def ctx(self, tokens):
        s = requests.get(f"{API}/students", headers=_headers(tokens["admin"]), timeout=30).json()
        a = requests.get(f"{API}/finance/accounts", headers=_headers(tokens["finance"]), timeout=30).json()
        return {"student_id": s[0]["id"], "account_id": a[0]["id"]}

    def _pay(self, tokens, ctx, key, nominal=100000):
        return requests.post(f"{API}/payments", headers=_headers(tokens["finance"]), timeout=30,
                             json={"student_id": ctx["student_id"], "nominal": nominal,
                                   "tanggal": date.today().isoformat(), "metode": "cash",
                                   "account_id": ctx["account_id"], "idem_key": key})

    def test_same_key_once(self, tokens, ctx):
        import uuid
        key = f"TEST-{uuid.uuid4()}"
        r1 = self._pay(tokens, ctx, key)
        assert r1.status_code == 200, r1.text
        r2 = self._pay(tokens, ctx, key)
        assert r2.status_code == 200 and r2.json().get("duplicate") is True
        assert r2.json()["id"] == r1.json()["id"]
        txs = requests.get(f"{API}/finance/transactions", headers=_headers(tokens["finance"]), timeout=30).json()
        assert sum(1 for t in txs if t.get("ref_type") == "payment" and t.get("ref_id") == r1.json()["id"]) == 1

    def test_concurrent_same_key(self, tokens, ctx):
        import concurrent.futures, uuid
        key = f"TEST-{uuid.uuid4()}"
        def one(_):
            r = self._pay(tokens, ctx, key)
            return (r.status_code, r.json().get("id"))
        with concurrent.futures.ThreadPoolExecutor(max_workers=6) as ex:
            out = list(ex.map(one, range(6)))
        assert all(c == 200 for c, _ in out), out
        ids = {i for _, i in out}
        assert len(ids) == 1, ids
    def test_different_keys_separate(self, tokens, ctx):
        import uuid
        ids = set()
        for _ in range(2):
            r = self._pay(tokens, ctx, f"TEST-{uuid.uuid4()}")
            assert r.status_code == 200 and not r.json().get("duplicate")
            ids.add(r.json()["id"])
        assert len(ids) == 2


# ---------------- WHATSAPP E1 ----------------
WA2 = "http://127.0.0.1:8001/api"


class TestWA:
    @pytest.fixture(scope="class")
    def t2(self, tokens):
        out = {}
        for role in ["owner", "admin", "finance", "hr", "guru", "marketing", "staff"]:
            r = requests.post(f"{WA2}/auth/login",
                              json={"email": f"{role}@lpk.id" if role != "owner" else "owner@lpk.id",
                                    "password": "owner123" if role == "owner" else "password123"}, timeout=15)
            assert r.status_code == 200, role
            out[role] = r.json()["token"]
        return out

    @pytest.fixture(scope="function")
    def wabackend(self, monkeypatch):
        import os
        import sys
        monkeypatch.setenv("MONGO_URL", "mongodb://127.0.0.1:27017")
        monkeypatch.setenv("DB_NAME", "lpk")
        backend = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        if backend not in sys.path:
            sys.path.insert(0, backend)
        import routers.whatsapp as wa
        return wa

    def _h2(self, t2, role):
        return {"Authorization": f"Bearer {t2[role]}"}

    def test_status_and_no_secrets(self, t2):
        r = requests.get(f"{WA2}/wa/status", headers=self._h2(t2, "owner"), timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert d["enabled"] is True and d["dry_run"] is True
        assert "TOKEN" not in r.text and "test-secret" not in r.text and "WA_TOKEN" not in r.text
        assert requests.get(f"{WA2}/wa/status", headers=self._h2(t2, "admin"), timeout=15).status_code == 200
        assert requests.get(f"{WA2}/wa/status", headers=self._h2(t2, "guru"), timeout=15).status_code == 403

    def test_connection_and_rbac(self, t2):
        r = requests.post(f"{WA2}/wa/test-connection", headers=self._h2(t2, "owner"), timeout=15)
        assert r.status_code == 200 and r.json()["mode"] == "dry_run"
        assert requests.post(f"{WA2}/wa/test-connection", headers=self._h2(t2, "finance"), timeout=15).status_code == 403

    def test_templates_crud_rbac(self, t2):
        import uuid
        key = f"TEST_tpl_{uuid.uuid4().hex[:8]}"
        hdr_o, hdr_f = self._h2(t2, "owner"), self._h2(t2, "finance")
        assert requests.get(f"{WA2}/wa/templates", headers=hdr_o, timeout=15).status_code == 200
        assert requests.get(f"{WA2}/wa/templates", headers=hdr_f, timeout=15).status_code == 200
        assert requests.get(f"{WA2}/wa/templates", headers=self._h2(t2, "guru"), timeout=15).status_code == 403
        assert requests.post(f"{WA2}/wa/templates", headers=hdr_f, timeout=15,
                             json={"key": key, "body": "x"}).status_code == 403
        r = requests.post(f"{WA2}/wa/templates", headers=hdr_o, timeout=15,
                          json={"key": key, "body": "Halo {{nama}}, TEST."})
        assert r.status_code == 400  # key tidak dikenal
        # version increment pada template existing
        cur = next(t for t in requests.get(f"{WA2}/wa/templates", headers=hdr_o, timeout=15).json() if t["key"] == "payment_due")
        v0 = cur["version"]
        r = requests.put(f"{WA2}/wa/templates/payment_due", headers=hdr_o, timeout=15,
                         json={"key": "payment_due", "body": cur["body"] + " "})
        assert r.status_code == 200 and r.json()["version"] == v0 + 1
        requests.put(f"{WA2}/wa/templates/payment_due", headers=hdr_o, timeout=15,
                     json={"key": "payment_due", "body": cur["body"]})

    def test_normalize_phone(self, wabackend):
        normalize_phone = wabackend.normalize_phone
        assert normalize_phone("0812 3456 7890") == "6281234567890"
        assert normalize_phone("+62-812-3456-7890") == "6281234567890"
        assert normalize_phone("6281234567890") == "6281234567890"
        assert normalize_phone("123") is None
        assert normalize_phone("") is None
        assert normalize_phone("0800abc") is None

    def test_consent_and_manual_dryrun(self, t2, tokens):
        hdr_o = self._h2(t2, "owner")
        s = requests.get(f"{API}/students", headers=_headers(tokens["owner"]), timeout=30).json()[0]
        sid = s["id"]
        orig_hp, orig_opt = s.get("no_hp", ""), bool(s.get("wa_student_opt_in"))
        try:
            # nomor invalid -> 400, tanpa record baru
            before = len(requests.get(f"{WA2}/wa/messages?student_id={sid}", headers=hdr_o, timeout=15).json())
            r = requests.post(f"{WA2}/wa/messages", headers=self._h2(t2, "finance"), timeout=30,
                              json={"student_id": sid, "recipient": "siswa", "template_key": "payment_due",
                                    "variables": {"nama": "X", "jumlah": "Rp1", "jatuh_tempo": "2026-01-01", "cara_bayar": "Cash"}})
            if not s.get("no_hp"):
                assert r.status_code == 400 and "valid" in r.json().get("detail", "")
            # nomor valid tapi tanpa consent -> 400 Consent
            requests.put(f"{API}/students/{sid}", headers=hdr_o, timeout=30,
                         json={**{k: v for k, v in s.items() if k not in ("_id", "id")},
                               "no_hp": "081234567890", "wa_student_opt_in": False})
            r = requests.post(f"{WA2}/wa/messages", headers=self._h2(t2, "finance"), timeout=30,
                              json={"student_id": sid, "recipient": "siswa", "template_key": "payment_due",
                                    "variables": {"nama": "X", "jumlah": "Rp1", "jatuh_tempo": "2026-01-01", "cara_bayar": "Cash"}})
            assert r.status_code == 400 and "Consent" in r.json().get("detail", "")
            after = len(requests.get(f"{WA2}/wa/messages?student_id={sid}", headers=hdr_o, timeout=15).json())
            assert after == before
            # beri consent + nomor
            requests.put(f"{API}/students/{sid}", headers=hdr_o, timeout=30,
                         json={**{k: v for k, v in s.items() if k not in ("_id", "id")},
                               "no_hp": "081234567890", "wa_student_opt_in": True})
            r = requests.post(f"{WA2}/wa/messages", headers=self._h2(t2, "finance"), timeout=30,
                              json={"student_id": sid, "recipient": "siswa", "template_key": "payment_due",
                                    "variables": {"nama": "X", "jumlah": "Rp1", "jatuh_tempo": "2026-01-01", "cara_bayar": "Cash"}})
            assert r.status_code == 200, r.text
            d = r.json()
            assert d["status"] == "sent" and d["dry_run"] is True and d["phone"] != "6281234567890"
            assert "6281234567890" not in requests.get(f"{WA2}/wa/messages?student_id={sid}", headers=hdr_o, timeout=15).text
        finally:
            requests.put(f"{API}/students/{sid}", headers=hdr_o, timeout=30,
                         json={**{k: v for k, v in s.items() if k not in ("_id", "id")},
                               "no_hp": orig_hp, "wa_student_opt_in": orig_opt})

    def test_manual_rbac(self, t2, tokens):
        hdr_o = self._h2(t2, "owner")
        s = requests.get(f"{API}/students", headers=_headers(tokens["owner"]), timeout=30).json()[0]
        body = {"student_id": s["id"], "recipient": "siswa", "template_key": "payment_due", "variables": {}}
        for role in ["admin", "hr", "guru", "marketing", "staff"]:
            assert requests.post(f"{WA2}/wa/messages", headers=self._h2(t2, role), timeout=15, json=body).status_code == 403, role
        for role in ["guru", "marketing", "staff"]:
            assert requests.get(f"{WA2}/wa/messages", headers=self._h2(t2, role), timeout=15).status_code == 403, role

    def test_retry_only_failed(self, t2):
        hdr_f = self._h2(t2, "finance")
        rows = requests.get(f"{WA2}/wa/messages?status=sent", headers=self._h2(t2, "owner"), timeout=15).json()
        if rows:
            r = requests.post(f"{WA2}/wa/messages/{rows[0]['id']}/retry", headers=hdr_f, timeout=15)
            assert r.status_code == 400
        r = requests.post(f"{WA2}/wa/messages/ tidak-ada /retry".replace(" ", ""), headers=hdr_f, timeout=15)
        assert r.status_code == 404

    def test_webhook(self, t2):
        import hmac, hashlib, json as _json, time
        hdr_o = self._h2(t2, "owner")
        r = requests.get(f"{WA2}/wa/webhook?hub.mode=subscribe&hub.verify_token=test-verify&hub.challenge=CHAL123", timeout=15)
        assert r.status_code == 200 and r.text == "CHAL123"
        r = requests.get(f"{WA2}/wa/webhook?hub.mode=subscribe&hub.verify_token=salah&hub.challenge=x", timeout=15)
        assert r.status_code == 403
        # buat pesan terkirim dengan provider id via fake adapter in-process? gunakan dry-run + update manual tidak bisa;
        # uji unknown-event path + invalid signature
        payload = {"entry": [{"changes": [{"value": {"statuses": [{"id": "wamid.TEST123", "status": "delivered",
                    "timestamp": str(int(time.time())), "recipient_id": "6281"}]}}]}]}
        raw = _json.dumps(payload).encode()
        sig = "sha256=" + hmac.new(b"test-secret-123", raw, hashlib.sha256).hexdigest()
        r = requests.post(f"{WA2}/wa/webhook", data=raw, headers={"X-Hub-Signature-256": sig}, timeout=15)
        assert r.status_code == 200
        r = requests.post(f"{WA2}/wa/webhook", data=raw, headers={"X-Hub-Signature-256": "sha256=dead"}, timeout=15)
        assert r.status_code == 403
        # duplicate
        r = requests.post(f"{WA2}/wa/webhook", data=raw, headers={"X-Hub-Signature-256": sig}, timeout=15)
        assert r.status_code == 200

    def test_service_fake_fail_retry_idempotent(self, t2, tokens, wabackend, monkeypatch):
        import asyncio
        import uuid
        wa = wabackend
        hdr_o = self._h2(t2, "owner")
        s = requests.get(f"{API}/students", headers=_headers(tokens["owner"]), timeout=30).json()[0]
        sid, orig = s["id"], dict(s)
        suffix = f"TEST-{uuid.uuid4().hex[:8]}"
        actor = {"id": "t", "name": "t", "role": "owner"}
        try:
            requests.put(f"{API}/students/{sid}", headers=hdr_o, timeout=30,
                         json={**{k: v for k, v in s.items() if k not in ("_id", "id")},
                               "no_hp": "081234567890", "wa_student_opt_in": True})
            monkeypatch.setenv("WA_PROVIDER", "fake_timeout")
            monkeypatch.setenv("WA_ENABLED", "true")
            monkeypatch.setenv("WA_DRY_RUN", "false")
            loop = asyncio.new_event_loop()
            m1 = loop.run_until_complete(wa.queue_message(
                "payment_due", sid, "siswa", suffix,
                {"nama": "T", "jumlah": "Rp1", "jatuh_tempo": "2026-01-01", "cara_bayar": "C"}, created_by=actor))
            assert m1["status"] == "queued"
            m1 = loop.run_until_complete(wa.attempt_send(m1["id"], actor))
            assert m1["status"] == "failed" and m1["retryable"] is True and m1["attempts"] == 1
            # duplicate kirim event sama -> record sama
            m2 = loop.run_until_complete(wa.queue_message(
                "payment_due", sid, "siswa", suffix,
                {"nama": "T", "jumlah": "Rp1", "jatuh_tempo": "2026-01-01", "cara_bayar": "C"}, created_by=actor))
            assert m2.get("duplicate") is True and m2["id"] == m1["id"]
            # retry dengan provider sukses
            monkeypatch.setenv("WA_PROVIDER", "fake_success")
            r = loop.run_until_complete(wa.wa_retry(m1["id"], actor))
            assert r["status"] == "sent" and r.get("retry_of") == m1["id"]
            assert r["provider_msg_id"].startswith("fake-")
            # adapter mapping murni (kontrak: fail = rejection non-retryable, timeout = retryable)
            assert wa.FakeAdapter("success").send("6281", "x", None, "id")["ok"] is True
            fail_res = wa.FakeAdapter("fail").send("6281", "x", None, "id")
            assert fail_res["ok"] is False and fail_res["retryable"] is False
            assert wa.FakeAdapter("timeout").send("6281", "x", None, "id") == {
                "ok": False, "retryable": True, "error": "Timeout provider (fake)"}
            assert wa.CloudApiAdapter().verify_signature(b"abc", "fake") is False
        finally:
            requests.put(f"{API}/students/{sid}", headers=hdr_o, timeout=30,
                         json={**{k: v for k, v in orig.items() if k not in ("_id", "id")}})
        # cleanup pesan uji via pymongo sinkron
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        d.whatsapp_messages.delete_many({"student_id": sid, "idempotency_key": {"$regex": suffix[:13]}})

    def test_webhook_progression(self, t2, tokens):
        import asyncio, hmac, hashlib, json as _json, time
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        mid = f"TEST-wamid-{int(time.time())}"
        s = requests.get(f"{API}/students", headers=_headers(tokens["owner"]), timeout=30).json()[0]
        d.whatsapp_messages.insert_one(
            {"id": f"TEST-msg-{mid}", "idempotency_key": f"TEST-hook-{mid}", "template_key": "payment_due",
             "template_version": 1, "body_snapshot": "x", "recipient_role": "siswa", "student_id": s["id"],
             "student_nama": s["nama_lengkap"], "recipient_name": "t", "phone": "6281234567890",
             "status": "sent", "provider": "fake_success", "provider_msg_id": mid, "fail_reason": None,
             "retryable": False, "attempts": 1, "next_retry_at": None, "sent_at": "2026-01-01",
             "retry_of": None, "created_by": "t", "created_at": "2026-01-01", "updated_at": "2026-01-01",
             "dry_run": False})
        try:
            def post(status, ts):
                payload = {"entry": [{"changes": [{"value": {"statuses": [
                    {"id": mid, "status": status, "timestamp": str(ts), "recipient_id": "6281"}]}}]}]}
                raw = _json.dumps(payload).encode()
                sig = "sha256=" + hmac.new(b"test-secret-123", raw, hashlib.sha256).hexdigest()
                return requests.post(f"{WA2}/wa/webhook", data=raw,
                                     headers={"X-Hub-Signature-256": sig}, timeout=15)
            assert post("delivered", 100).status_code == 200
            assert d.whatsapp_messages.find_one({"id": f"TEST-msg-{mid}"})["status"] == "delivered"
            assert post("delivered", 100).status_code == 200  # duplicate no-op
            assert post("read", 200).status_code == 200
            assert d.whatsapp_messages.find_one({"id": f"TEST-msg-{mid}"})["status"] == "read"
            assert post("delivered", 50).status_code == 200  # out-of-order: tidak turun
            assert d.whatsapp_messages.find_one({"id": f"TEST-msg-{mid}"})["status"] == "read"
        finally:
            d.whatsapp_messages.delete_many({"id": f"TEST-msg-{mid}"})
            d.whatsapp_events.delete_many({"message_id": mid})

    def test_audit_no_secrets(self, t2, tokens):
        rows = requests.get(f"{API}/audit-logs?entity=wa_message&limit=20", headers=_headers(tokens["owner"]), timeout=30).json()
        blob = str(rows).lower()
        assert "test-secret" not in blob and "wa_token" not in blob
        rows = requests.get(f"{API}/audit-logs?entity=wa_template&limit=20", headers=_headers(tokens["owner"]), timeout=30).json()
        assert isinstance(rows, list)


# ---------------- STUDENT PORTAL ----------------
def _portal_login(email, password):
    return requests.post(f"{API}/student/auth/login", json={"email": email, "password": password}, timeout=30)


@pytest.fixture(scope="module")
def portal_pair(tokens, worker_id):
    """Two TEST students + student accounts. Unique per xdist worker. Yields dict; cleans up afterwards."""
    import uuid
    tag_suffix = uuid.uuid4().hex[:6]
    admin_h, owner_h = _headers(tokens["admin"]), _headers(tokens["owner"])
    def _email(tag):
        return f"test.portal.{tag.lower()}.{tag_suffix}@lpk.id"
    def _name(tag):
        return f"TEST Portal {tag} {tag_suffix}"
    for tag in ("PA", "PB"):
        for u in requests.get(f"{API}/users", headers=owner_h, timeout=30).json():
            if u["email"] == _email(tag):
                requests.delete(f"{API}/users/{u['id']}", headers=owner_h, timeout=30)
        for s in requests.get(f"{API}/students", headers=admin_h, timeout=30).json():
            if s["nama_lengkap"] == _name(tag):
                requests.delete(f"{API}/students/{s['id']}", headers=owner_h, timeout=30)
    created = {"sids": [], "uids": []}
    try:
        pair = {}
        for tag in ("PA", "PB"):
            r = requests.post(f"{API}/students", headers=admin_h,
                              json={"nama_lengkap": _name(tag), "jenis_kelamin": "L",
                                    "no_hp": "0811", "alamat": {"jalan": "Jl. Test", "kota": "Jakarta"}})
            assert r.status_code == 200, r.text
            sid = r.json()["id"]
            created["sids"].append(sid)
            email = _email(tag)
            r = requests.post(f"{API}/users", headers=owner_h,
                              json={"name": _name(tag), "email": email,
                                    "password": "portal123", "role": "student", "student_id": sid})
            assert r.status_code == 200, r.text
            created["uids"].append(r.json()["id"])
            r = _portal_login(email, "portal123")
            assert r.status_code == 200, r.text
            pair[tag] = {"sid": sid, "email": email, "token": r.json()["token"]}
        yield pair
    finally:
        for uid in created["uids"]:
            requests.delete(f"{API}/users/{uid}", headers=owner_h, timeout=30)
        for sid in created["sids"]:
            requests.delete(f"{API}/students/{sid}", headers=owner_h, timeout=30)


def _ph(token):
    return {"Authorization": f"Bearer {token}"}


class TestStudentPortalAuth:
    def test_login_valid(self, portal_pair):
        assert portal_pair["PA"]["token"]
        r = requests.get(f"{API}/student/auth/me", headers=_ph(portal_pair["PA"]["token"]), timeout=30)
        assert r.status_code == 200
        d = r.json()
        assert d["role"] == "student" and d["student_id"] == portal_pair["PA"]["sid"]
        assert d["must_change_password"] is True
        assert "password_hash" not in d

    def test_login_invalid(self, portal_pair):
        r = _portal_login(portal_pair["PA"]["email"], "salah123")
        assert r.status_code == 401

    def test_login_staff_email_rejected(self, tokens):
        r = _portal_login("admin@lpk.id", "password123")
        assert r.status_code == 401

    def test_inactive_account(self, portal_pair, tokens):
        owner_h = _headers(tokens["owner"])
        users = requests.get(f"{API}/users", headers=owner_h, timeout=30).json()
        uid = next(u["id"] for u in users if u["email"] == portal_pair["PB"]["email"])
        base = next(u for u in users if u["id"] == uid)
        try:
            body = {"name": base["name"], "email": base["email"], "role": "student",
                    "employee_id": None, "aktif": False}
            assert requests.put(f"{API}/users/{uid}", headers=owner_h, json=body, timeout=30).status_code == 200
            assert _portal_login(portal_pair["PB"]["email"], "portal123").status_code == 403
        finally:
            body["aktif"] = True
            requests.put(f"{API}/users/{uid}", headers=owner_h, json=body, timeout=30)

    def test_must_change_gate_and_change_password(self, portal_pair):
        tok = portal_pair["PA"]["token"]
        assert requests.get(f"{API}/student/dashboard", headers=_ph(tok), timeout=30).status_code == 403
        assert requests.post(f"{API}/student/auth/change-password", headers=_ph(tok),
                             json={"old_password": "salah", "new_password": "baru1234"}, timeout=30).status_code == 401
        assert requests.post(f"{API}/student/auth/change-password", headers=_ph(tok),
                             json={"old_password": "portal123", "new_password": "abc"}, timeout=30).status_code == 400
        r = requests.post(f"{API}/student/auth/change-password", headers=_ph(tok),
                          json={"old_password": "portal123", "new_password": "baru1234"}, timeout=30)
        assert r.status_code == 200, r.text
        assert requests.get(f"{API}/student/dashboard", headers=_ph(tok), timeout=30).status_code == 200
        r = _portal_login(portal_pair["PA"]["email"], "baru1234")
        assert r.status_code == 200
        assert r.json()["user"]["must_change_password"] is False
        portal_pair["PA"]["token"] = r.json()["token"]
        portal_pair["PA"]["password"] = "baru1234"

    def test_logout(self, portal_pair):
        tok = portal_pair["PA"].get("token")
        assert requests.post(f"{API}/student/auth/logout", headers=_ph(tok), timeout=30).status_code == 200


class TestStudentPortalRBAC:
    @pytest.mark.parametrize("path", ["/dashboard", "/profile", "/class", "/attendance",
                                      "/grades", "/exams", "/payments", "/documents",
                                      "/notifications", "/whatsapp-consent"])
    def test_staff_cannot_use_portal(self, tokens, path):
        for role in ("owner", "admin", "finance", "hr", "guru", "marketing", "staff"):
            r = requests.get(f"{API}/student{path}", headers=_headers(tokens[role]), timeout=30)
            assert r.status_code == 403, f"{role} {path}: {r.status_code}"

    @pytest.mark.parametrize("path", ["/students", "/finance/summary", "/audit-logs",
                                      "/payrolls", "/expenses", "/dashboard", "/hr/attendance"])
    def test_student_cannot_use_internal(self, portal_pair, path):
        tok = portal_pair["PA"].get("token")
        r = requests.get(f"{API}{path}", headers=_ph(tok), timeout=30)
        assert r.status_code in (401, 403), f"{path}: {r.status_code}"

    def test_unauthenticated(self, portal_pair):
        assert requests.get(f"{API}/student/dashboard", timeout=30).status_code == 401


class TestStudentPortalIDOR:
    def test_own_profile(self, portal_pair):
        tok = portal_pair["PA"].get("token")
        r = requests.get(f"{API}/student/profile", headers=_ph(tok), timeout=30)
        assert r.status_code == 200
        assert r.json()["id"] == portal_pair["PA"]["sid"]

    def test_cross_student_payment_detail(self, portal_pair, tokens):
        finance_h = _headers(tokens["finance"])
        acc = requests.get(f"{API}/finance/accounts", headers=finance_h, timeout=30).json()[0]["id"]
        sid_b = portal_pair["PB"]["sid"]
        r = requests.post(f"{API}/payments", headers=finance_h,
                          json={"student_id": sid_b, "nominal": 100000, "tanggal": date.today().isoformat(),
                                "account_id": acc, "catatan": "TEST portal"})
        assert r.status_code == 200, r.text
        pid_b, tx_id = r.json()["id"], None
        try:
            tx = requests.get(f"{API}/finance/transactions", headers=finance_h, timeout=30).json()
            tx_id = next(t["id"] for t in tx if t.get("ref_id") == pid_b)
            tok_a = portal_pair["PA"].get("token")
            assert requests.get(f"{API}/student/payments/{pid_b}", headers=_ph(tok_a), timeout=30).status_code == 404
            rows = requests.get(f"{API}/student/payments", headers=_ph(tok_a), timeout=30).json()["rows"]
            assert all(p["id"] != pid_b for p in rows)
        finally:
            requests.delete(f"{API}/payments/{pid_b}", headers=_headers(tokens["owner"]), timeout=30)

    def test_cross_student_notifications(self, portal_pair, tokens):
        owner_h = _headers(tokens["owner"])
        all_n = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        assert all_n, "need seeded notifications"
        tok_a = portal_pair["PA"].get("token")
        mine = requests.get(f"{API}/student/notifications", headers=_ph(tok_a), timeout=30).json()["rows"]
        assert isinstance(mine, list)
        foreign = next((n for n in all_n if "TEST Portal" not in n.get("judul", "")), None)
        if foreign:
            assert requests.post(f"{API}/student/notifications/{foreign['id']}/read",
                                 headers=_ph(tok_a), timeout=30).status_code in (403, 404)

    def test_cross_student_documents(self, portal_pair):
        tok_a = portal_pair["PA"].get("token")
        rows = requests.get(f"{API}/student/documents", headers=_ph(tok_a), timeout=30).json()["rows"]
        assert isinstance(rows, list)
        assert requests.get(f"{API}/student/documents/bogus-id/download", headers=_ph(tok_a), timeout=30).status_code == 404

    def test_cross_student_consent(self, portal_pair):
        tok_a = portal_pair["PA"].get("token")
        before = requests.get(f"{API}/student/whatsapp-consent", headers=_ph(tok_a), timeout=30).json()
        assert before["wa_guardian_opt_in"] in (True, False)
        r = requests.put(f"{API}/student/whatsapp-consent", headers=_ph(tok_a),
                         json={"wa_student_opt_in": True, "wa_guardian_opt_in": True, "no_hp": "0899"},
                         timeout=30)
        assert r.status_code == 200
        after = requests.get(f"{API}/student/whatsapp-consent", headers=_ph(tok_a), timeout=30).json()
        assert after["wa_student_opt_in"] is True
        assert after["wa_guardian_opt_in"] == before["wa_guardian_opt_in"]
        requests.put(f"{API}/student/whatsapp-consent", headers=_ph(tok_a),
                     json={"wa_student_opt_in": before.get("wa_student_opt_in", False)}, timeout=30)


class TestStudentPortalData:
    def test_dashboard_shape(self, portal_pair):
        tok = portal_pair["PA"].get("token")
        d = requests.get(f"{API}/student/dashboard", headers=_ph(tok), timeout=30).json()
        assert set(["profile", "academic", "finance", "documents", "notifications", "whatsapp"]) <= set(d.keys())
        assert d["profile"]["nama"]
        blob = str(d)
        assert "gaji_pokok" not in blob and "account_id" not in blob and "petugas" not in blob

    def test_academic_smoke(self, portal_pair):
        tok = portal_pair["PA"].get("token")
        for p in ("class", "attendance", "grades", "exams"):
            assert requests.get(f"{API}/student/{p}", headers=_ph(tok), timeout=30).status_code == 200, p

    def test_finance_privacy(self, portal_pair):
        tok = portal_pair["PA"].get("token")
        d = requests.get(f"{API}/student/payments", headers=_ph(tok), timeout=30).json()
        assert set(["fee_plan", "total", "bayar", "sisa", "rows"]) <= set(d.keys())
        assert "transactions" not in d and "accounts" not in d


class TestStudentPortalProvisioning:
    def test_admin_create_and_duplicates(self, portal_pair, tokens):
        import uuid
        tag = uuid.uuid4().hex[:6]
        admin_h, owner_h = _headers(tokens["admin"]), _headers(tokens["owner"])
        s = requests.post(f"{API}/students", headers=admin_h,
                          json={"nama_lengkap": f"TEST Portal PC {tag}", "jenis_kelamin": "L"}).json()
        sid = s["id"]
        try:
            base = {"name": f"TEST Portal PC {tag}", "email": f"test.portal.pc.{tag}@lpk.id",
                    "password": "portal123", "role": "student", "student_id": sid}
            r = requests.post(f"{API}/users", headers=admin_h, json=base, timeout=30)
            assert r.status_code == 200, r.text
            assert r.json()["must_change_password"] is True
            uid = r.json()["id"]
            try:
                r = requests.post(f"{API}/users", headers=admin_h,
                                  json={**base, "email": f"test.portal.pc2.{tag}@lpk.id"}, timeout=30)
                assert r.status_code == 409, r.text
                r = requests.post(f"{API}/users", headers=admin_h,
                                  json={"name": "X", "email": base["email"],
                                        "password": "portal123", "role": "student", "student_id": sid}, timeout=30)
                assert r.status_code == 409, r.text
                s2 = requests.post(f"{API}/students", headers=admin_h,
                                   json={"nama_lengkap": f"TEST Portal PC2 {tag}", "jenis_kelamin": "L"}).json()
                try:
                    r = requests.post(f"{API}/users", headers=admin_h,
                                      json={"name": "X", "email": base["email"],
                                            "password": "portal123", "role": "student",
                                            "student_id": s2["id"]}, timeout=30)
                    assert r.status_code == 400, r.text
                finally:
                    requests.delete(f"{API}/students/{s2['id']}", headers=owner_h, timeout=30)
                r = requests.post(f"{API}/users", headers=admin_h,
                                  json={**base, "email": f"test.portal.pc3.{tag}@lpk.id", "student_id": "nope"}, timeout=30)
                assert r.status_code == 400, r.text
                for role in ("finance", "guru", "marketing", "staff", "hr"):
                    r = requests.post(f"{API}/users", headers=_headers(tokens[role]),
                                      json={**base, "email": f"test.portal.{role}.{tag}@lpk.id"}, timeout=30)
                    assert r.status_code == 403, role
            finally:
                requests.delete(f"{API}/users/{uid}", headers=owner_h, timeout=30)
        finally:
            requests.delete(f"{API}/students/{sid}", headers=owner_h, timeout=30)


class TestStudentProvisioningRace:
    """RV-DEF-01 corrective: concurrent duplicate provisioning of the same
    (email, student_id) must deterministically yield 1x200 + 4x409, exactly
    one account, no 400/500 — regardless of thread interleaving."""

    def _race_once(self, tokens, run_tag):
        import concurrent.futures
        import uuid
        tag = uuid.uuid4().hex[:6]
        owner_h = _headers(tokens["owner"])
        s = requests.post(f"{API}/students", headers=owner_h,
                          json={"nama_lengkap": f"TEST RV-DEF-01 {run_tag} {tag}",
                                "jenis_kelamin": "L"}, timeout=30).json()
        sid = s["id"]
        uid = None
        try:
            body = {"name": f"TEST RV-DEF-01 {run_tag} {tag}",
                    "email": f"test.rvdef01.{run_tag.lower()}.{tag}@lpk.id",
                    "password": "Rahasia123", "role": "student", "student_id": sid}

            def hit(_):
                r = requests.post(f"{API}/users", headers=owner_h, json=body, timeout=60)
                return r.status_code

            with concurrent.futures.ThreadPoolExecutor(max_workers=5) as ex:
                codes = sorted(ex.map(hit, range(5)))
            assert codes.count(200) == 1, codes
            assert codes.count(409) == 4, codes
            assert not any(c in (400, 500) for c in codes), codes
            from pymongo import MongoClient
            d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
            users = list(d.users.find({"student_id": sid}, {"_id": 0, "id": 1}))
            assert len(users) == 1, users
            uid = users[0]["id"]
        finally:
            if uid:
                requests.delete(f"{API}/users/{uid}", headers=owner_h, timeout=30)
            requests.delete(f"{API}/students/{sid}", headers=owner_h, timeout=30)

    def test_race_run1(self, tokens):
        self._race_once(tokens, "R1")

    def test_race_run2(self, tokens):
        self._race_once(tokens, "R2")

    def test_race_run3(self, tokens):
        self._race_once(tokens, "R3")

    def test_nonstudent_duplicate_email_stays_400(self, tokens):
        import uuid
        tag = uuid.uuid4().hex[:6]
        owner_h = _headers(tokens["owner"])
        email = f"test.rvdef01.staff.{tag}@lpk.id"
        uid = None
        try:
            r = requests.post(f"{API}/users", headers=owner_h,
                              json={"name": f"TEST RVDEF01 S {tag}", "email": email,
                                    "password": "Rahasia123", "role": "staff"}, timeout=30)
            assert r.status_code == 200, r.text
            uid = r.json()["id"]
            r = requests.post(f"{API}/users", headers=owner_h,
                              json={"name": f"TEST RVDEF01 S2 {tag}", "email": email,
                                    "password": "Rahasia123", "role": "staff"}, timeout=30)
            assert r.status_code == 400, r.text
            assert "terdaftar" in r.json().get("detail", "")
        finally:
            if uid:
                requests.delete(f"{API}/users/{uid}", headers=owner_h, timeout=30)


# ---------------- COLLECTION / PENAGIHAN TUNGGAKAN (P1.1) ----------------
class TestCollection:
    @pytest.fixture(scope="class")
    def col_student(self, tokens):
        import uuid
        from datetime import date, timedelta
        tag = uuid.uuid4().hex[:6]
        admin_h, owner_h, fin_h = _headers(tokens["admin"]), _headers(tokens["owner"]), _headers(tokens["finance"])
        r = requests.post(f"{API}/students", headers=admin_h,
                          json={"nama_lengkap": f"TEST Collection {tag}", "jenis_kelamin": "L",
                                "no_hp": "0812000111", "wa_student_opt_in": True}, timeout=30)
        assert r.status_code == 200, r.text
        sid = r.json()["id"]
        past = (date.today() - timedelta(days=3)).isoformat()
        r = requests.put(f"{API}/students/{sid}/fee-plan", headers=fin_h,
                         json={"items": [{"nama": "Uang Pangkal", "nominal": 5000000}], "jatuh_tempo": past}, timeout=30)
        assert r.status_code == 200, r.text
        r = requests.put(f"{API}/students/{sid}/status", headers=admin_h,
                         json={"status": "pelatihan", "catatan": "test collection"}, timeout=30)
        assert r.status_code == 200, r.text
        yield {"sid": sid, "tag": tag}
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        d.collection_activities.delete_many({"student_id": sid})
        d.whatsapp_messages.delete_many({"student_id": sid})
        requests.delete(f"{API}/students/{sid}", headers=owner_h, timeout=30)

    def _idem(self, prefix):
        import uuid
        return f"TEST-{prefix}-{uuid.uuid4().hex[:8]}"

    def test_manual_create_finance(self, tokens, col_student):
        from datetime import date, timedelta
        fu = (date.today() + timedelta(days=2)).isoformat()
        r = requests.post(f"{API}/collections/activities", headers=_headers(tokens["finance"]),
                          json={"student_id": col_student["sid"], "channel": "phone", "outcome": "promised_payment",
                                "note": "Janji transfer Jumat", "next_follow_up_at": fu,
                                "next_follow_up_note": "Tagih lagi", "idem_key": self._idem("m1")}, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["actor_role"] == "finance" and d["actor_name"]
        assert d["sisa_snapshot"]["sisa"] == 5000000
        assert d["next_follow_up_at"] == fu
        assert d["wa_message_id"] is None

    def test_invalid_student_404(self, tokens):
        r = requests.post(f"{API}/collections/activities", headers=_headers(tokens["finance"]),
                          json={"student_id": "tidak-ada", "channel": "phone", "outcome": "contacted",
                                "note": "x", "idem_key": self._idem("bad")}, timeout=30)
        assert r.status_code == 404, r.text

    def test_roles_forbidden(self, tokens, col_student):
        body = {"student_id": col_student["sid"], "channel": "phone", "outcome": "contacted",
                "note": "x", "idem_key": self._idem("f1")}
        for role in ("guru", "marketing", "staff", "hr"):
            r = requests.post(f"{API}/collections/activities", headers=_headers(tokens[role]), json=body, timeout=30)
            assert r.status_code == 403, role
        for role in ("guru", "marketing", "staff", "hr"):
            assert requests.get(f"{API}/collections/overview", headers=_headers(tokens[role]), timeout=30).status_code == 403, role
            assert requests.get(f"{API}/collections/students/{col_student['sid']}/activities",
                                headers=_headers(tokens[role]), timeout=30).status_code == 403, role

    def test_admin_read_only(self, tokens, col_student):
        assert requests.get(f"{API}/collections/overview", headers=_headers(tokens["admin"]), timeout=30).status_code == 200
        r = requests.post(f"{API}/collections/activities", headers=_headers(tokens["admin"]),
                          json={"student_id": col_student["sid"], "channel": "phone", "outcome": "contacted",
                                "note": "x", "idem_key": self._idem("adm")}, timeout=30)
        assert r.status_code == 403, r.text

    def test_financial_integrity(self, tokens, col_student):
        fin_h = _headers(tokens["finance"])
        before_tx = len(requests.get(f"{API}/finance/transactions", headers=fin_h, timeout=30).json())
        before_sum = requests.get(f"{API}/finance/summary?period=bulan", headers=fin_h, timeout=30).json()
        before_ar = next(s for s in requests.get(f"{API}/payments-arrears", headers=fin_h, timeout=30).json()
                         if s["id"] == col_student["sid"])
        r = requests.post(f"{API}/collections/activities", headers=fin_h,
                          json={"student_id": col_student["sid"], "channel": "in_person", "outcome": "contacted",
                                "note": "Ketemu di LPK", "idem_key": self._idem("fi")}, timeout=30)
        assert r.status_code == 200, r.text
        after_tx = len(requests.get(f"{API}/finance/transactions", headers=fin_h, timeout=30).json())
        after_sum = requests.get(f"{API}/finance/summary?period=bulan", headers=fin_h, timeout=30).json()
        after_ar = next(s for s in requests.get(f"{API}/payments-arrears", headers=fin_h, timeout=30).json()
                        if s["id"] == col_student["sid"])
        assert after_tx == before_tx
        assert after_sum["saldo_kas"] == before_sum["saldo_kas"]
        assert after_ar["sisa"] == before_ar["sisa"] == 5000000

    def test_update_needs_alasan(self, tokens, col_student):
        fin_h = _headers(tokens["finance"])
        r = requests.post(f"{API}/collections/activities", headers=fin_h,
                          json={"student_id": col_student["sid"], "channel": "other", "outcome": "no_response",
                                "note": "Belum diangkat", "idem_key": self._idem("up")}, timeout=30)
        aid = r.json()["id"]
        r = requests.put(f"{API}/collections/activities/{aid}", headers=fin_h,
                         json={"outcome": "promised_payment"}, timeout=30)
        assert r.status_code == 400, r.text
        r = requests.put(f"{API}/collections/activities/{aid}", headers=fin_h,
                         json={"outcome": "promised_payment", "alasan": "Siswa menghubungi balik"}, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json()["outcome"] == "promised_payment"

    def test_duplicate_idem_manual(self, tokens, col_student):
        fin_h = _headers(tokens["finance"])
        body = {"student_id": col_student["sid"], "channel": "phone", "outcome": "contacted",
                "note": "Double klik", "idem_key": self._idem("dup")}
        r1 = requests.post(f"{API}/collections/activities", headers=fin_h, json=body, timeout=30)
        r2 = requests.post(f"{API}/collections/activities", headers=fin_h, json=body, timeout=30)
        assert r1.status_code == 200 and r2.status_code == 200
        assert r2.json().get("duplicate") is True and r2.json()["id"] == r1.json()["id"]

    def test_send_wa_dry_run(self, tokens, col_student):
        t2 = requests.post("http://127.0.0.1:8001/api/auth/login",
                           json={"email": "finance@lpk.id", "password": "password123"}, timeout=15).json()["token"]
        h2 = {"Authorization": f"Bearer {t2}"}
        r = requests.post("http://127.0.0.1:8001/api/collections/activities/send-wa", headers=h2,
                          json={"student_id": col_student["sid"], "recipient": "siswa",
                                "outcome": "contacted", "note": "Reminder via WA",
                                "idem_key": self._idem("wa")}, timeout=60)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["wa_sent"] is True and d["wa_status"] == "sent" and d["wa_dry_run"] is True
        assert d["channel"] == "whatsapp" and d["wa_message_id"]
        r2 = requests.post("http://127.0.0.1:8001/api/collections/activities/send-wa", headers=h2,
                           json={"student_id": col_student["sid"], "recipient": "siswa",
                                 "outcome": "contacted", "note": "double",
                                 "idem_key": d["idem_key"]}, timeout=60)
        assert r2.status_code == 200 and r2.json().get("duplicate") is True

    def test_send_wa_disabled_honest(self, tokens, col_student):
        fin_h = _headers(tokens["finance"])
        before = len(requests.get(f"{API}/collections/students/{col_student['sid']}/activities",
                                  headers=fin_h, timeout=30).json())
        r = requests.post(f"{API}/collections/activities/send-wa", headers=fin_h,
                          json={"student_id": col_student["sid"], "recipient": "siswa",
                                "outcome": "contacted", "note": "x", "idem_key": self._idem("dis")}, timeout=30)
        assert r.status_code == 400, r.text
        after = len(requests.get(f"{API}/collections/students/{col_student['sid']}/activities",
                                 headers=fin_h, timeout=30).json())
        assert after == before

    def test_student_role_blocked(self, tokens, portal_pair):
        import requests as _rq
        tok = portal_pair["PA"]["token"]
        assert _rq.get(f"{API}/collections/overview", headers={"Authorization": f"Bearer {tok}"},
                       timeout=30).status_code == 403

    def test_followup_notification_no_wa(self, tokens, col_student):
        from datetime import date
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        fin_h, owner_h = _headers(tokens["finance"]), _headers(tokens["owner"])
        wa_before = d.whatsapp_messages.count_documents({"student_id": col_student["sid"]})
        r = requests.post(f"{API}/collections/activities", headers=fin_h,
                          json={"student_id": col_student["sid"], "channel": "phone",
                                "outcome": "promised_payment", "note": "Follow up hari ini",
                                "next_follow_up_at": date.today().isoformat(),
                                "idem_key": self._idem("nt")}, timeout=30)
        assert r.status_code == 200, r.text
        notifs = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        col = [n for n in notifs if n.get("tipe") == "collection"]
        assert any("Follow-up penagihan" in n.get("judul", "") for n in col), [n.get("judul") for n in col]
        assert all("/pembayaran" in (n.get("link") or "") for n in col)
        wa_after = d.whatsapp_messages.count_documents({"student_id": col_student["sid"]})
        assert wa_after == wa_before


# ---------------- COLLECTION NOTIFICATION QA-COL-01 ----------------
class TestCollectionNotification:
    @pytest.fixture(scope="class")
    def ncol_student(self, tokens):
        import uuid
        from datetime import date, timedelta
        tag = uuid.uuid4().hex[:6]
        admin_h, owner_h, fin_h = _headers(tokens["admin"]), _headers(tokens["owner"]), _headers(tokens["finance"])
        r = requests.post(f"{API}/students", headers=admin_h,
                          json={"nama_lengkap": f"TEST ColNotif {tag}", "jenis_kelamin": "L",
                                "no_hp": "0812000222", "wa_student_opt_in": True}, timeout=30)
        assert r.status_code == 200, r.text
        sid = r.json()["id"]
        past = (date.today() - timedelta(days=3)).isoformat()
        r = requests.put(f"{API}/students/{sid}/fee-plan", headers=fin_h,
                         json={"items": [{"nama": "Uang Pangkal", "nominal": 5000000}], "jatuh_tempo": past}, timeout=30)
        assert r.status_code == 200, r.text
        r = requests.put(f"{API}/students/{sid}/status", headers=admin_h,
                         json={"status": "pelatihan", "catatan": "test colnotif"}, timeout=30)
        assert r.status_code == 200, r.text
        yield {"sid": sid, "tag": tag}
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        d.collection_activities.delete_many({"student_id": sid})
        d.whatsapp_messages.delete_many({"student_id": sid})
        requests.delete(f"{API}/students/{sid}", headers=owner_h, timeout=30)

    def _idem(self, prefix):
        import uuid
        return f"TEST-CN-{prefix}-{uuid.uuid4().hex[:8]}"

    def _act(self, tokens, sid, outcome, note, fu=None):
        r = requests.post(f"{API}/collections/activities", headers=_headers(tokens["finance"]),
                          json={"student_id": sid, "channel": "phone", "outcome": outcome, "note": note,
                                "next_follow_up_at": fu, "idem_key": self._idem(note[:4])}, timeout=30)
        assert r.status_code == 200, r.text
        return r.json()

    def _col_notifs(self, tokens):
        rows = requests.get(f"{API}/notifications", headers=_headers(tokens["owner"]), timeout=30).json()
        return [n for n in rows if n.get("tipe") == "collection"]

    def test_a_close_paid_removes_stale(self, tokens, ncol_student):
        from datetime import date, timedelta
        sid = ncol_student["sid"]
        past = (date.today() - timedelta(days=2)).isoformat()
        self._act(tokens, sid, "no_response", "A-old", past)
        assert any(ncol_student["tag"] in n.get("judul", "") for n in self._col_notifs(tokens))
        self._act(tokens, sid, "paid_after_contact", "A-paid")
        assert not any(ncol_student["tag"] in n.get("judul", "") for n in self._col_notifs(tokens)), \
            [n.get("judul") for n in self._col_notifs(tokens)]

    def test_b_latest_without_fu_suppresses(self, tokens, ncol_student):
        from datetime import date, timedelta
        sid = ncol_student["sid"]
        past = (date.today() - timedelta(days=2)).isoformat()
        self._act(tokens, sid, "no_response", "B-old", past)
        assert any(ncol_student["tag"] in n.get("judul", "") for n in self._col_notifs(tokens))
        self._act(tokens, sid, "contacted", "B-new")
        assert not any(ncol_student["tag"] in n.get("judul", "") for n in self._col_notifs(tokens))

    def test_c_latest_active_fu_wins(self, tokens, ncol_student):
        from datetime import date, timedelta
        sid = ncol_student["sid"]
        past = (date.today() - timedelta(days=2)).isoformat()
        old = self._act(tokens, sid, "no_response", "C-old", past)
        new = self._act(tokens, sid, "promised_payment", "C-new", date.today().isoformat())
        mine = [n for n in self._col_notifs(tokens) if ncol_student["tag"] in n.get("judul", "")]
        assert len(mine) == 1, [n.get("judul") for n in mine]
        assert mine[0]["dedupe_key"] == f"col-fu:{new['id']}"
        assert f"col-fu:{old['id']}" not in [n.get("dedupe_key") for n in self._col_notifs(tokens)]

    def test_d_no_duplicates_on_refresh(self, tokens, ncol_student):
        first = [n for n in self._col_notifs(tokens) if ncol_student["tag"] in n.get("judul", "")]
        second = [n for n in self._col_notifs(tokens) if ncol_student["tag"] in n.get("judul", "")]
        assert [n.get("dedupe_key") for n in first] == [n.get("dedupe_key") for n in second]
        assert len({n.get("dedupe_key") for n in first}) == len(first)


# ---------------- CANDIDATE FOLLOW-UP (P1.2) ----------------
class TestCandidateFollowup:
    @pytest.fixture(scope="class")
    def cf_setup(self, tokens):
        import uuid
        tag = uuid.uuid4().hex[:6]
        admin_h = _headers(tokens["admin"])
        mk = {}
        def _mk(name, extra):
            r = requests.post(f"{API}/students", headers=admin_h,
                              json={"nama_lengkap": name, "jenis_kelamin": "L", **extra}, timeout=30)
            assert r.status_code == 200, r.text
            return r.json()["id"]
        mk["c1"] = _mk(f"TEST Calon {tag}", {"no_hp": "0812333444", "wa_student_opt_in": True,
                                             "sumber_prospek": "sosmed", "pemilik_lead": "Marketing A"})
        mk["c2"] = _mk(f"TEST Calon NC {tag}", {"no_hp": "0812333555"})
        mk["c3"] = _mk(f"TEST Calon NP {tag}", {})
        mk["n1"] = _mk(f"TEST NonCalon {tag}", {"no_hp": "0812333666", "wa_student_opt_in": True})
        r = requests.put(f"{API}/students/{mk['n1']}/status", headers=admin_h,
                         json={"status": "pelatihan", "catatan": "test p12"}, timeout=30)
        assert r.status_code == 200, r.text
        mk["tag"] = tag
        yield mk
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        for sid in (mk["c1"], mk["c2"], mk["c3"], mk["n1"]):
            d.candidate_followups.delete_many({"student_id": sid})
            d.whatsapp_messages.delete_many({"student_id": sid})
            requests.delete(f"{API}/students/{sid}", headers=_headers(tokens["owner"]), timeout=30)

    def _idem(self, prefix):
        import uuid
        return f"TEST-CF-{prefix}-{uuid.uuid4().hex[:8]}"

    def _mk(self, tokens):
        return _headers(tokens["marketing"])

    def test_01_create_manual(self, tokens, cf_setup):
        r = requests.post(f"{API}/candidate-followups/activities", headers=self._mk(tokens),
                          json={"student_id": cf_setup["c1"], "channel": "phone", "outcome": "minat",
                                "note": "Tertarik program", "next_follow_up_at": "2099-01-01",
                                "idem_key": self._idem("m1")}, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["actor_role"] == "marketing" and d["actor_name"] and d["created_at"]
        assert d["prospect_snapshot"]["status"] == "calon_siswa"
        assert "sisa" not in d and "nominal" not in str(d)

    def test_02_invalid_student(self, tokens):
        r = requests.post(f"{API}/candidate-followups/activities", headers=self._mk(tokens),
                          json={"student_id": "tidak-ada", "channel": "phone", "outcome": "minat",
                                "note": "x", "idem_key": self._idem("bad")}, timeout=30)
        assert r.status_code == 404, r.text

    def test_03_non_candidate_rejected(self, tokens, cf_setup):
        r = requests.post(f"{API}/candidate-followups/activities", headers=self._mk(tokens),
                          json={"student_id": cf_setup["n1"], "channel": "phone", "outcome": "minat",
                                "note": "x", "idem_key": self._idem("nc")}, timeout=30)
        assert r.status_code == 400, r.text

    def test_04_enums(self, tokens, cf_setup):
        for body in ({"channel": "fax", "outcome": "minat"}, {"channel": "phone", "outcome": "kaya"}):
            r = requests.post(f"{API}/candidate-followups/activities", headers=self._mk(tokens),
                              json={"student_id": cf_setup["c1"], **body, "note": "x",
                                    "idem_key": self._idem("en")}, timeout=30)
            assert r.status_code == 400, r.text

    def test_05_history(self, tokens, cf_setup):
        rows = requests.get(f"{API}/candidate-followups/students/{cf_setup['c1']}/activities",
                            headers=self._mk(tokens), timeout=30).json()
        assert any(r["outcome"] == "minat" and r["note"] == "Tertarik program" for r in rows)

    def test_06_correction_reason(self, tokens, cf_setup):
        mkh = self._mk(tokens)
        r = requests.post(f"{API}/candidate-followups/activities", headers=mkh,
                          json={"student_id": cf_setup["c1"], "channel": "sosmed", "outcome": "terhubungi",
                                "note": "Chat IG", "idem_key": self._idem("cr")}, timeout=30)
        aid = r.json()["id"]
        assert requests.put(f"{API}/candidate-followups/activities/{aid}", headers=mkh,
                            json={"outcome": "minat"}, timeout=30).status_code == 400
        r = requests.put(f"{API}/candidate-followups/activities/{aid}", headers=mkh,
                         json={"outcome": "minat", "alasan": "Koreksi hasil"}, timeout=30)
        assert r.status_code == 200, r.text
        au = requests.get(f"{API}/audit-logs?entity=candidate_followup&limit=50",
                          headers=_headers(tokens["owner"]), timeout=30).json()
        upd = [x for x in au if x.get("entity_id") == aid and x["action"] == "update"]
        assert upd and upd[0]["before"].get("outcome") == "terhubungi"

    def test_08_no_delete(self, tokens, cf_setup):
        r = requests.delete(f"{API}/candidate-followups/activities/xxx", headers=self._mk(tokens), timeout=30)
        assert r.status_code in (404, 405), r.text

    def test_09_idempotent_manual(self, tokens, cf_setup):
        body = {"student_id": cf_setup["c1"], "channel": "kunjungan", "outcome": "terhubungi",
                "note": "Datang ke LPK", "idem_key": self._idem("dup")}
        r1 = requests.post(f"{API}/candidate-followups/activities", headers=self._mk(tokens), json=body, timeout=30)
        r2 = requests.post(f"{API}/candidate-followups/activities", headers=self._mk(tokens), json=body, timeout=30)
        assert r1.status_code == 200 and r2.json().get("duplicate") is True
        assert r2.json()["id"] == r1.json()["id"]

    def test_10_concurrent_idem(self, tokens, cf_setup):
        import concurrent.futures
        body = {"student_id": cf_setup["c1"], "channel": "phone", "outcome": "terhubungi",
                "note": "race", "idem_key": self._idem("race")}
        h = self._mk(tokens)

        def hit(_):
            r = requests.post(f"{API}/candidate-followups/activities", headers=h, json=body, timeout=60)
            return (r.status_code, r.json().get("duplicate", False))
        with concurrent.futures.ThreadPoolExecutor(max_workers=5) as ex:
            res = list(ex.map(hit, range(5)))
        assert all(c == 200 for c, _ in res)
        assert sum(1 for _, d_ in res if d_) == 4
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        assert d.candidate_followups.count_documents({"idem_key": body["idem_key"]}) == 1

    def test_11_wa_dry_run(self, tokens, cf_setup):
        t2 = requests.post("http://127.0.0.1:8001/api/auth/login",
                           json={"email": "marketing@lpk.id", "password": "password123"}, timeout=15).json()["token"]
        h2 = {"Authorization": f"Bearer {t2}"}
        r = requests.post("http://127.0.0.1:8001/api/candidate-followups/activities/send-wa", headers=h2,
                          json={"student_id": cf_setup["c1"], "outcome": "terhubungi", "note": "Sapa WA",
                                "tanggal_follow_up": "Senin", "idem_key": self._idem("wa")}, timeout=60)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["wa_sent"] is True and d["wa_status"] == "sent" and d["wa_dry_run"] is True
        assert d["channel"] == "whatsapp" and d["template_key"] == "followup_calon" and d["wa_message_id"]
        r2 = requests.post("http://127.0.0.1:8001/api/candidate-followups/activities/send-wa", headers=h2,
                           json={"student_id": cf_setup["c1"], "outcome": "terhubungi", "note": "x",
                                 "idem_key": d["idem_key"]}, timeout=60)
        assert r2.json().get("duplicate") is True

    def test_12_wa_disabled(self, tokens, cf_setup):
        mkh = self._mk(tokens)
        n0 = len(requests.get(f"{API}/candidate-followups/students/{cf_setup['c1']}/activities",
                              headers=mkh, timeout=30).json())
        r = requests.post(f"{API}/candidate-followups/activities/send-wa", headers=mkh,
                          json={"student_id": cf_setup["c1"], "outcome": "terhubungi", "note": "x",
                                "idem_key": self._idem("dis")}, timeout=30)
        assert r.status_code == 400, r.text
        n1 = len(requests.get(f"{API}/candidate-followups/students/{cf_setup['c1']}/activities",
                              headers=mkh, timeout=30).json())
        assert n1 == n0

    def test_13_no_consent(self, tokens, cf_setup):
        t2 = requests.post("http://127.0.0.1:8001/api/auth/login",
                           json={"email": "marketing@lpk.id", "password": "password123"}, timeout=15).json()["token"]
        r = requests.post("http://127.0.0.1:8001/api/candidate-followups/activities/send-wa",
                          headers={"Authorization": f"Bearer {t2}"},
                          json={"student_id": cf_setup["c2"], "outcome": "terhubungi", "note": "x",
                                "idem_key": self._idem("nc")}, timeout=30)
        assert r.status_code == 400, r.text

    def test_14_invalid_number(self, tokens, cf_setup):
        t2 = requests.post("http://127.0.0.1:8001/api/auth/login",
                           json={"email": "marketing@lpk.id", "password": "password123"}, timeout=15).json()["token"]
        s = requests.post("http://127.0.0.1:8001/api/students", headers={"Authorization": f"Bearer {t2}"},
                          json={"nama_lengkap": f"TEST Calon NP2 {cf_setup['tag']}", "jenis_kelamin": "L",
                                "wa_student_opt_in": True}, timeout=30).json()
        try:
            r = requests.post("http://127.0.0.1:8001/api/candidate-followups/activities/send-wa",
                              headers={"Authorization": f"Bearer {t2}"},
                              json={"student_id": s["id"], "outcome": "terhubungi", "note": "x",
                                    "idem_key": self._idem("nn")}, timeout=30)
            assert r.status_code == 400, r.text
        finally:
            t0 = _headers(tokens["owner"])
            from pymongo import MongoClient
            MongoClient("mongodb://127.0.0.1:27017")["lpk"].candidate_followups.delete_many({"student_id": s["id"]})
            MongoClient("mongodb://127.0.0.1:27017")["lpk"].whatsapp_messages.delete_many({"student_id": s["id"]})
            requests.delete(f"{API}/students/{s['id']}", headers=t0, timeout=30)

    def test_16_17_notif_due_overdue(self, tokens, cf_setup):
        from datetime import date, timedelta
        mkh, owner_h = self._mk(tokens), _headers(tokens["owner"])
        past = (date.today() - timedelta(days=1)).isoformat()
        r = requests.post(f"{API}/candidate-followups/activities", headers=mkh,
                          json={"student_id": cf_setup["c2"], "channel": "phone", "outcome": "terhubungi",
                                "note": "Notif test", "next_follow_up_at": past,
                                "idem_key": self._idem("nt")}, timeout=30)
        assert r.status_code == 200, r.text
        rows = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        mine = [n for n in rows if n.get("tipe") == "candidate" and cf_setup["tag"] in n.get("judul", "")]
        assert len(mine) == 1 and mine[0]["level"] == "danger", [n.get("judul") for n in mine]
        assert "/followup" in (mine[0].get("link") or "")

    def test_18_latest_wins(self, tokens, cf_setup):
        from datetime import date, timedelta
        mkh, owner_h = self._mk(tokens), _headers(tokens["owner"])
        past = (date.today() - timedelta(days=2)).isoformat()
        old = requests.post(f"{API}/candidate-followups/activities", headers=mkh,
                            json={"student_id": cf_setup["c2"], "channel": "phone", "outcome": "terhubungi",
                                  "note": "old", "next_follow_up_at": past,
                                  "idem_key": self._idem("lw")}, timeout=30).json()
        new = requests.post(f"{API}/candidate-followups/activities", headers=mkh,
                            json={"student_id": cf_setup["c2"], "channel": "phone", "outcome": "janji_datang",
                                  "note": "new", "next_follow_up_at": date.today().isoformat(),
                                  "idem_key": self._idem("lw2")}, timeout=30).json()
        rows = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        keys = [n.get("dedupe_key") for n in rows if n.get("tipe") == "candidate"]
        assert f"candidate-fu:{new['id']}" in keys
        assert f"candidate-fu:{old['id']}" not in keys

    def test_20_refresh_dedupe(self, tokens, cf_setup):
        owner_h = _headers(tokens["owner"])
        a = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        b = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        ka = [n.get("dedupe_key") for n in a if n.get("tipe") == "candidate"]
        kb = [n.get("dedupe_key") for n in b if n.get("tipe") == "candidate"]
        assert ka == kb and len(set(ka)) == len(ka)

    def test_19_closed_suppresses(self, tokens, cf_setup):
        mkh, owner_h = self._mk(tokens), _headers(tokens["owner"])
        r = requests.post(f"{API}/candidate-followups/activities", headers=mkh,
                          json={"student_id": cf_setup["c2"], "channel": "phone", "outcome": "mendaftar",
                                "note": "Daftar!", "idem_key": self._idem("cl")}, timeout=30)
        assert r.status_code == 200, r.text
        rows = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        assert not any(cf_setup["tag"] in n.get("judul", "") and n.get("tipe") == "candidate" for n in rows)

    def test_21_rbac_matrix(self, tokens, cf_setup):
        exp = {"owner": (200, 200), "admin": (200, 200), "marketing": (200, 200), "staff": (200, 200),
               "hr": (403, 403), "finance": (403, 403), "guru": (403, 403)}
        for role, (ev, ec) in exp.items():
            v = requests.get(f"{API}/candidate-followups/overview", headers=_headers(tokens[role]), timeout=15).status_code
            c = requests.post(f"{API}/candidate-followups/activities", headers=_headers(tokens[role]),
                              json={"student_id": cf_setup["c1"], "channel": "phone", "outcome": "minat",
                                    "note": "x", "idem_key": self._idem(role)}, timeout=15).status_code
            assert (v, c) == (ev, ec), role

    def test_22_idor(self, tokens, cf_setup):
        guru_h = _headers(tokens["guru"])
        assert requests.get(f"{API}/candidate-followups/students/{cf_setup['c1']}/activities",
                            headers=guru_h, timeout=30).status_code == 403
        assert requests.put(f"{API}/candidate-followups/activities/xxx", headers=guru_h,
                            json={"outcome": "minat", "alasan": "x"}, timeout=30).status_code in (403, 404)

    def test_23_student_denied(self, tokens, portal_pair):
        tok = portal_pair["PA"]["token"]
        assert requests.get(f"{API}/candidate-followups/overview",
                            headers={"Authorization": f"Bearer {tok}"}, timeout=30).status_code == 403

    def test_24_financial_unchanged(self, tokens, cf_setup):
        mkh = self._mk(tokens)
        fin_h = _headers(tokens["finance"])

        def snap():
            from pymongo import MongoClient
            d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
            s = requests.get(f"{API}/finance/summary?period=bulan", headers=fin_h, timeout=30).json()
            return (d.transactions.count_documents({}), d.payments.count_documents({}), s["saldo_kas"])
        s0 = snap()
        requests.post(f"{API}/candidate-followups/activities", headers=mkh,
                      json={"student_id": cf_setup["c1"], "channel": "sosmed", "outcome": "minat",
                            "note": "fin check", "idem_key": self._idem("fi")}, timeout=30)
        requests.get(f"{API}/notifications", headers=_headers(tokens["owner"]), timeout=30)
        assert snap() == s0

    def test_26_source_field(self, tokens, cf_setup):
        admin_h = _headers(tokens["admin"])
        s = requests.get(f"{API}/students/{cf_setup['c1']}", headers=admin_h, timeout=30).json()
        assert s["sumber_prospek"] == "sosmed" and s["pemilik_lead"] == "Marketing A"
        r = requests.post(f"{API}/students", headers=admin_h,
                          json={"nama_lengkap": "TEST Bad Sumber", "jenis_kelamin": "L",
                                "sumber_prospek": "alien"}, timeout=30)
        assert r.status_code == 400, r.text

    def test_27_summary_map(self, tokens, cf_setup):
        m = requests.get(f"{API}/candidate-followups/summary-map", headers=self._mk(tokens), timeout=30).json()
        e = m.get(cf_setup["c1"])
        assert e and e["count"] >= 1 and e["last_outcome"] and e["last_contacted"]

    def test_28_overview_matches(self, tokens, cf_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        ov = requests.get(f"{API}/candidate-followups/overview", headers=self._mk(tokens), timeout=30).json()
        assert ov["total_calon"] == d.students.count_documents({"status": "calon_siswa"})
        assert any(x["id"] == cf_setup["c1"] for x in ov["never_contacted"]) is False
        assert set(ov.keys()) >= {"total_calon", "never_contacted", "due_today", "overdue", "closed", "recent"}

    def test_25_conversion_preserves(self, tokens, cf_setup):
        admin_h, mkh = _headers(tokens["admin"]), self._mk(tokens)
        before = requests.get(f"{API}/candidate-followups/students/{cf_setup['c1']}/activities",
                              headers=mkh, timeout=30).json()
        assert len(before) >= 1
        r = requests.put(f"{API}/students/{cf_setup['c1']}/status", headers=admin_h,
                         json={"status": "pendaftaran", "catatan": "mendaftar"}, timeout=30)
        assert r.status_code == 200, r.text
        after = requests.get(f"{API}/candidate-followups/students/{cf_setup['c1']}/activities",
                             headers=mkh, timeout=30).json()
        assert len(after) == len(before)
        r = requests.post(f"{API}/candidate-followups/activities", headers=mkh,
                          json={"student_id": cf_setup["c1"], "channel": "phone", "outcome": "minat",
                                "note": "x", "idem_key": self._idem("post")}, timeout=30)
        assert r.status_code == 400, r.text


# ---------------- DEPARTURE READINESS (P1.3) ----------------
class TestDeparture:
    @pytest.fixture(scope="class")
    def dep_setup(self, tokens):
        import uuid
        from datetime import date, timedelta
        tag = uuid.uuid4().hex[:6]
        admin_h, owner_h = _headers(tokens["admin"]), _headers(tokens["owner"])
        mk = {}
        for key, nm in (("d1", f"TEST Dep {tag}"), ("d2", f"TEST Dep Exp {tag}")):
            r = requests.post(f"{API}/students", headers=admin_h,
                              json={"nama_lengkap": nm, "jenis_kelamin": "L",
                                    "no_hp": "0812444001", "wa_student_opt_in": True}, timeout=30)
            assert r.status_code == 200, r.text
            mk[key] = r.json()["id"]
            r = requests.put(f"{API}/students/{mk[key]}/status", headers=admin_h,
                             json={"status": "pemberkasan", "catatan": "test p13"}, timeout=30)
            assert r.status_code == 200, r.text
        mk["tag"] = tag
        mk["future"] = (date.today() + timedelta(days=365)).isoformat()
        mk["past"] = (date.today() - timedelta(days=10)).isoformat()
        yield mk
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        for sid in (mk["d1"], mk["d2"]):
            for pid in [p["id"] for p in d.departure_profiles.find({"student_id": sid}, {"_id": 0, "id": 1})]:
                d.departure_checklist.delete_many({"departure_profile_id": pid})
            d.departure_profiles.delete_many({"student_id": sid})
            d.documents.delete_many({"student_id": sid})
            requests.delete(f"{API}/students/{sid}", headers=owner_h, timeout=30)

    def _idem(self, prefix):
        import uuid
        return f"TEST-DEP-{prefix}-{uuid.uuid4().hex[:8]}"

    def _mkdoc(self, tokens, sid, jenis, exp):
        r = requests.post(f"{API}/students/{sid}/documents", headers=_headers(tokens["admin"]),
                          data={"jenis": jenis, "kategori": "jepang", "tanggal_kadaluarsa": exp}, timeout=30)
        assert r.status_code == 200, r.text
        return r.json()["id"]

    def _mkprofile(self, tokens, sid, **kw):
        r = requests.post(f"{API}/departures", headers=_headers(tokens["owner"]),
                          json={"student_id": sid, **kw}, timeout=30)
        assert r.status_code == 200, r.text
        return r.json()

    def _items(self, tokens, pid):
        return requests.get(f"{API}/departures/{pid}/checklist",
                            headers=_headers(tokens["owner"]), timeout=30).json()

    def test_01_create_and_checklist(self, tokens, dep_setup):
        from datetime import date, timedelta
        target = (date.today() + timedelta(days=60)).isoformat()
        p = self._mkprofile(tokens, dep_setup["d1"], target_departure_date=target,
                            destination="Osaka", pic_name="PIC A")
        assert p["readiness"]["checklist_total"] == 6
        codes = {i["requirement_code"] for i in self._items(tokens, p["id"])}
        assert codes == {"passport", "coe", "visa", "medical", "ticket", "contract"}
        assert p["readiness"]["readiness_status"] == "BLOCKED"

    def test_02_duplicate_profile(self, tokens, dep_setup):
        r = requests.post(f"{API}/departures", headers=_headers(tokens["owner"]),
                          json={"student_id": dep_setup["d1"]}, timeout=30)
        assert r.status_code == 409, r.text

    def test_02b_concurrent_duplicate(self, tokens, dep_setup):
        import concurrent.futures
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        before = d.departure_profiles.count_documents({"student_id": dep_setup["d2"]})
        h = _headers(tokens["owner"])

        def hit(_):
            return requests.post(f"{API}/departures", headers=h,
                                 json={"student_id": dep_setup["d2"]}, timeout=60).status_code
        with concurrent.futures.ThreadPoolExecutor(max_workers=5) as ex:
            codes = sorted(ex.map(hit, range(5)))
        assert codes.count(200) == 1 and codes.count(409) == 4, codes
        assert d.departure_profiles.count_documents({"student_id": dep_setup["d2"]}) - before == 1

    def test_03_invalid_ids(self, tokens):
        assert requests.post(f"{API}/departures", headers=_headers(tokens["owner"]),
                             json={"student_id": "nope"}, timeout=30).status_code == 404
        assert requests.get(f"{API}/departures/nope", headers=_headers(tokens["owner"]),
                            timeout=30).status_code == 404

    def test_05_verify_with_doc(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        doc = self._mkdoc(tokens, dep_setup["d1"], "Paspor", dep_setup["future"])
        item = next(i for i in self._items(tokens, pid) if i["requirement_code"] == "passport")
        r = requests.put(f"{API}/departures/{pid}/checklist/{item['id']}/verify",
                         headers=_headers(tokens["admin"]),
                         json={"document_id": doc, "note": "Fisik cocok"}, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json()["verified_by"] and r.json()["verified_at"]

    def test_06_verify_expired_rejected(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d2"]})["id"]
        doc = self._mkdoc(tokens, dep_setup["d2"], "Paspor", dep_setup["past"])
        item = next(i for i in self._items(tokens, pid) if i["requirement_code"] == "passport")
        r = requests.put(f"{API}/departures/{pid}/checklist/{item['id']}/verify",
                         headers=_headers(tokens["admin"]), json={"document_id": doc}, timeout=30)
        assert r.status_code == 400, r.text

    def test_07_verify_foreign_doc(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        other = d.documents.find_one({"student_id": dep_setup["d2"]})
        item = next(i for i in self._items(tokens, pid) if i["requirement_code"] == "coe")
        r = requests.put(f"{API}/departures/{pid}/checklist/{item['id']}/verify",
                         headers=_headers(tokens["admin"]), json={"document_id": other["id"]}, timeout=30)
        assert r.status_code == 400, r.text

    def test_08_reject_needs_reason(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        item = next(i for i in self._items(tokens, pid) if i["requirement_code"] == "visa")
        h = _headers(tokens["staff"])
        assert requests.post(f"{API}/departures/{pid}/checklist/{item['id']}/reject",
                             headers=h, json={"reason": ""}, timeout=30).status_code == 400
        r = requests.post(f"{API}/departures/{pid}/checklist/{item['id']}/reject",
                          headers=h, json={"reason": "Foto buram"}, timeout=30)
        assert r.status_code == 200 and r.json()["status"] == "rejected", r.text

    def test_09_exception_flow(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        item = next(i for i in self._items(tokens, pid) if i["requirement_code"] == "medical")
        r = requests.post(f"{API}/departures/{pid}/checklist/{item['id']}/exception",
                          headers=_headers(tokens["staff"]), json={"reason": "MCU susulan RS lain"}, timeout=30)
        assert r.status_code == 200 and r.json()["exception_status"] == "requested", r.text
        r = requests.post(f"{API}/departures/{pid}/checklist/{item['id']}/exception",
                          headers=_headers(tokens["owner"]), json={"reason": "Disetujui Owner"}, timeout=30)
        assert r.status_code == 200 and r.json()["status"] == "exception", r.text
        assert r.json()["exception_approved_by"]

    def test_10_readiness_aggregation(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        r = requests.get(f"{API}/departures/{pid}/readiness", headers=_headers(tokens["owner"]), timeout=30).json()
        assert set(r.keys()) >= {"readiness_status", "checklist_total", "checklist_verified",
                                 "checklist_pending", "checklist_rejected", "checklist_exception",
                                 "blockers", "warnings", "days_to_departure"}
        assert r["checklist_total"] == 6 and r["readiness_status"] == "BLOCKED"

    def test_12_visa_expiry_blocker(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d2"]})["id"]
        doc = self._mkdoc(tokens, dep_setup["d2"], "Visa", dep_setup["future"])
        item = next(i for i in self._items(tokens, pid) if i["requirement_code"] == "visa")
        assert requests.put(f"{API}/departures/{pid}/checklist/{item['id']}/verify",
                            headers=_headers(tokens["admin"]), json={"document_id": doc}, timeout=30).status_code == 200
        assert requests.put(f"{API}/students/{dep_setup['d2']}/documents/status",
                            headers=_headers(tokens["admin"]),
                            json={"jenis": "Visa", "kategori": "jepang", "status": "tersedia",
                                  "tanggal_kadaluarsa": dep_setup["past"]}, timeout=30).status_code == 200
        r = requests.get(f"{API}/departures/{pid}/readiness", headers=_headers(tokens["owner"]), timeout=30).json()
        assert r["readiness_status"] == "BLOCKED"
        assert any("expired" in b for b in r["blockers"]), r["blockers"]

    def test_14_arrears_warning_not_blocker(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        r = requests.get(f"{API}/departures/{pid}/readiness", headers=_headers(tokens["owner"]), timeout=30).json()
        assert any("Sisa tagihan" in w for w in r["warnings"]), r["warnings"]
        assert not any("tagihan" in b.lower() or "rp" in b.lower() for b in r["blockers"]), r["blockers"]

    def test_10b_ready_requires_conditions(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        r = requests.post(f"{API}/departures/{pid}/ready", headers=_headers(tokens["owner"]),
                          json={"reason": "coba"}, timeout=30)
        assert r.status_code == 400, r.text

    def test_15_full_verify_then_ready(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        h = _headers(tokens["owner"])
        for it in self._items(tokens, pid):
            if it["status"] in ("verified", "exception"):
                continue
            r = requests.put(f"{API}/departures/{pid}/checklist/{it['id']}/verify",
                             headers=h, json={"note": "verifikasi fisik"}, timeout=30)
            assert r.status_code == 200, r.text
        r = requests.post(f"{API}/departures/{pid}/ready", headers=h, json={"reason": "Siap berangkat"}, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json()["final_decision"] == "READY" and r.json()["final_decision_by"]
        assert r.json()["readiness"]["readiness_status"] == "READY"

    def test_16_block_flow(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d2"]})["id"]
        h = _headers(tokens["owner"])
        assert requests.post(f"{API}/departures/{pid}/block", headers=h, json={"reason": ""}, timeout=30).status_code == 400
        r = requests.post(f"{API}/departures/{pid}/block", headers=h, json={"reason": "Tunggu COE"}, timeout=30)
        assert r.status_code == 200, r.text
        assert requests.get(f"{API}/departures/{pid}/readiness", headers=h, timeout=30).json()["readiness_status"] == "BLOCKED"

    def test_17_dates_and_ticket(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        h = _headers(tokens["owner"])
        assert requests.put(f"{API}/departures/{pid}", headers=h,
                            json={"target_departure_date": "2099-06-01"}, timeout=30).status_code == 400
        r = requests.put(f"{API}/departures/{pid}", headers=h,
                         json={"target_departure_date": "2099-06-01", "actual_departure_date": "2099-06-02",
                               "ticket_airline": "GA", "flight_number": "GA881",
                               "departure_airport": "CGK", "arrival_airport": "NRT",
                               "alasan": "Jadwal fix"}, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json()["flight_number"] == "GA881" and r.json()["actual_departure_date"] == "2099-06-02"

    def test_19_20_notification(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d2"]})["id"]
        owner_h = _headers(tokens["owner"])
        rows = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        mine = [n for n in rows if n.get("tipe") == "departure" and n.get("entity_id") == pid]
        assert len(mine) == 1 and mine[0]["level"] == "danger", [n.get("judul") for n in mine]
        rows2 = requests.get(f"{API}/notifications", headers=owner_h, timeout=30).json()
        assert [n.get("dedupe_key") for n in rows if n.get("tipe") == "departure"] == \
               [n.get("dedupe_key") for n in rows2 if n.get("tipe") == "departure"]

    def test_21_rbac(self, tokens, dep_setup):
        sid = dep_setup["d1"]
        for role in ("owner", "admin", "staff"):
            assert requests.get(f"{API}/departures", headers=_headers(tokens[role]), timeout=15).status_code == 200, role
        for role in ("finance", "hr", "guru", "marketing"):
            assert requests.get(f"{API}/departures", headers=_headers(tokens[role]), timeout=15).status_code == 200, role
        for role in ("finance", "hr", "guru", "marketing"):
            r = requests.post(f"{API}/departures", headers=_headers(tokens[role]),
                              json={"student_id": sid}, timeout=15)
            assert r.status_code == 403, role
        assert requests.put(f"{API}/departures/xxx/checklist/yyy/verify", headers=_headers(tokens["guru"]),
                            json={}, timeout=15).status_code in (403, 404)
        assert requests.post(f"{API}/departures/xxx/ready", headers=_headers(tokens["admin"]),
                             json={}, timeout=15).status_code in (403, 404)

    def test_22_idor_student(self, tokens, dep_setup):
        import uuid
        tag = uuid.uuid4().hex[:6]
        owner_h = _headers(tokens["owner"])
        r = requests.post(f"{API}/users", headers=owner_h,
                          json={"name": f"TEST Dep Portal {tag}", "email": f"test.depportal.{tag}@lpk.id",
                                "password": "portal123", "role": "student",
                                "student_id": dep_setup["d1"]}, timeout=30)
        assert r.status_code == 200, r.text
        uid = r.json()["id"]
        try:
            t = requests.post(f"{API}/student/auth/login",
                              json={"email": f"test.depportal.{tag}@lpk.id", "password": "portal123"},
                              timeout=30).json()["token"]
            sh = {"Authorization": f"Bearer {t}"}
            assert requests.get(f"{API}/departures", headers=sh, timeout=30).status_code == 403
            assert requests.post(f"{API}/student/auth/change-password",
                                 headers=sh, json={"old_password": "portal123", "new_password": "Baru12345"},
                                 timeout=30).status_code == 200
            t = requests.post(f"{API}/student/auth/login",
                              json={"email": f"test.depportal.{tag}@lpk.id", "password": "Baru12345"},
                              timeout=30).json()["token"]
            sh = {"Authorization": f"Bearer {t}"}
            r = requests.get(f"{API}/student/departure", headers=sh, timeout=30)
            assert r.status_code == 200, r.text
            assert r.json()["profile"] is not None
            assert "final_decision_by" not in str(r.json())
        finally:
            requests.delete(f"{API}/users/{uid}", headers=owner_h, timeout=30)

    def test_24_financial_isolation(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        fin_h = _headers(tokens["finance"])
        owner_h = _headers(tokens["owner"])
        s0 = (d.transactions.count_documents({}), d.payments.count_documents({}),
              requests.get(f"{API}/finance/summary?period=bulan", headers=fin_h, timeout=30).json()["saldo_kas"])
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        requests.put(f"{API}/departures/{pid}", headers=owner_h,
                     json={"notes": "fin check", "alasan": "test"}, timeout=30)
        requests.get(f"{API}/notifications", headers=owner_h, timeout=30)
        s1 = (d.transactions.count_documents({}), d.payments.count_documents({}),
              requests.get(f"{API}/finance/summary?period=bulan", headers=fin_h, timeout=30).json()["saldo_kas"])
        assert s0 == s1

    def test_25_audit(self, tokens, dep_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pid = d.departure_profiles.find_one({"student_id": dep_setup["d1"]})["id"]
        rows = requests.get(f"{API}/audit-logs?entity=departure&limit=100",
                            headers=_headers(tokens["owner"]), timeout=30).json()
        mine = [x for x in rows if x.get("entity_id") == pid]
        assert {x["action"] for x in mine} >= {"create", "update", "ready"}
        assert all(x.get("user_name") and x.get("timestamp") for x in mine)

    def test_262728_spot_regression(self, tokens, dep_setup):
        assert requests.get(f"{API}/job-orders", headers=_headers(tokens["admin"]), timeout=30).status_code == 200
        assert requests.get(f"{API}/collections/overview", headers=_headers(tokens["owner"]),
                            timeout=30).status_code == 200
        assert requests.get(f"{API}/candidate-followups/overview",
                            headers=_headers(tokens["marketing"]), timeout=30).status_code == 200
        assert requests.get(f"{API}/students/{dep_setup['d1']}/documents",
                            headers=_headers(tokens["admin"]), timeout=30).status_code == 200


# ---------------- DEPARTURE PIC NOTIFICATION (QA-DEP-01) ----------------
class TestDeparturePIC:
    @pytest.fixture(scope="class")
    def pic_setup(self, tokens):
        import uuid
        tag = uuid.uuid4().hex[:6]
        owner_h, admin_h = _headers(tokens["owner"]), _headers(tokens["admin"])
        users = requests.get(f"{API}/users", headers=owner_h, timeout=30).json()
        staff_a = next(u for u in users if u["role"] == "staff")
        r = requests.post(f"{API}/users", headers=owner_h,
                          json={"name": f"TEST StaffB {tag}", "email": f"test.staffb.{tag}@lpk.id",
                                "password": "Rahasia123", "role": "staff"}, timeout=30)
        assert r.status_code == 200, r.text
        staff_b = r.json()
        r = requests.post(f"{API}/students", headers=admin_h,
                          json={"nama_lengkap": f"TEST Dep PIC {tag}", "jenis_kelamin": "L"}, timeout=30)
        assert r.status_code == 200, r.text
        sid = r.json()["id"]
        r = requests.put(f"{API}/students/{sid}/status", headers=admin_h,
                         json={"status": "pemberkasan", "catatan": "test pic"}, timeout=30)
        assert r.status_code == 200, r.text
        r = requests.post(f"{API}/departures", headers=owner_h, json={"student_id": sid}, timeout=30)
        assert r.status_code == 200, r.text
        pid = r.json()["id"]
        yield {"sid": sid, "pid": pid, "tag": tag, "staff_a": staff_a, "staff_b": staff_b}
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        d.departure_checklist.delete_many({"departure_profile_id": pid})
        d.departure_profiles.delete_many({"id": pid})
        requests.delete(f"{API}/students/{sid}", headers=owner_h, timeout=30)
        requests.delete(f"{API}/users/{staff_b['id']}", headers=owner_h, timeout=30)

    def _tok(self, email, pwd):
        return requests.post(f"{API}/auth/login", json={"email": email, "password": pwd},
                             timeout=15).json()["token"]

    def _dep(self, token, tag):
        rows = requests.get(f"{API}/notifications", headers={"Authorization": f"Bearer {token}"},
                            timeout=30).json()
        return [n for n in rows if n.get("tipe") == "departure" and tag in n.get("judul", "")]

    def _set_pic(self, tokens, pid, pic_id):
        r = requests.put(f"{API}/departures/{pid}", headers=_headers(tokens["owner"]),
                         json={"pic_user_id": pic_id, "alasan": "QA PIC test"}, timeout=30)
        assert r.status_code == 200, r.text

    def test_a_pic_receives(self, tokens, pic_setup):
        self._set_pic(tokens, pic_setup["pid"], pic_setup["staff_a"]["id"])
        tb = self._tok("staff@lpk.id", "password123")
        got = self._dep(tb, pic_setup["tag"])
        assert len(got) == 1 and got[0]["level"] == "danger", [n.get("judul") for n in got]

    def test_b_non_pic_no_receive(self, tokens, pic_setup):
        tb = self._tok(f"test.staffb.{pic_setup['tag']}@lpk.id", "Rahasia123")
        assert self._dep(tb, pic_setup["tag"]) == []

    def test_c_owner_admin_receive(self, tokens, pic_setup):
        to = self._tok("owner@lpk.id", "owner123")
        ta = self._tok("admin@lpk.id", "password123")
        assert len(self._dep(to, pic_setup["tag"])) == 1
        assert len(self._dep(ta, pic_setup["tag"])) == 1

    def test_d_pic_owner_no_duplicate(self, tokens, pic_setup):
        owner = next(u for u in requests.get(f"{API}/users", headers=_headers(tokens["owner"]),
                                             timeout=30).json() if u["role"] == "owner")
        self._set_pic(tokens, pic_setup["pid"], owner["id"])
        to = self._tok("owner@lpk.id", "owner123")
        got = self._dep(to, pic_setup["tag"])
        assert len(got) == 1, [n.get("judul") for n in got]
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        assert d.notifications.count_documents({"dedupe_key": f"dep:{pic_setup['pid']}"}) == 1
        self._set_pic(tokens, pic_setup["pid"], pic_setup["staff_a"]["id"])

    def test_e_empty_pic(self, tokens, pic_setup):
        self._set_pic(tokens, pic_setup["pid"], pic_setup["staff_a"]["id"])
        r = requests.put(f"{API}/departures/{pic_setup['pid']}", headers=_headers(tokens["owner"]),
                         json={"pic_user_id": None, "alasan": "QA PIC clear"}, timeout=30)
        assert r.status_code == 200, r.text
        tb = self._tok("staff@lpk.id", "password123")
        assert self._dep(tb, pic_setup["tag"]) == []
        to = self._tok("owner@lpk.id", "owner123")
        assert len(self._dep(to, pic_setup["tag"])) == 1
        self._set_pic(tokens, pic_setup["pid"], pic_setup["staff_a"]["id"])

    def test_f_lifecycle_intact(self, tokens, pic_setup):
        to = self._tok("owner@lpk.id", "owner123")
        assert len(self._dep(to, pic_setup["tag"])) == 1
        pid = pic_setup["pid"]
        items = requests.get(f"{API}/departures/{pid}/checklist",
                             headers=_headers(tokens["owner"]), timeout=30).json()
        for it in items:
            requests.put(f"{API}/departures/{pid}/checklist/{it['id']}/verify",
                         headers=_headers(tokens["owner"]), json={"note": "ok"}, timeout=30)
        assert requests.post(f"{API}/departures/{pid}/ready", headers=_headers(tokens["owner"]),
                             json={"reason": "QA"}, timeout=30).status_code == 200
        assert self._dep(to, pic_setup["tag"]) == []
        assert requests.post(f"{API}/departures/{pid}/block", headers=_headers(tokens["owner"]),
                             json={"reason": "QA reopen"}, timeout=30).status_code == 200
        assert len(self._dep(to, pic_setup["tag"])) == 1

    def test_g_read_isolated(self, tokens, pic_setup):
        ta_tok = self._tok("staff@lpk.id", "password123")
        to_tok = self._tok("owner@lpk.id", "owner123")
        nid = self._dep(ta_tok, pic_setup["tag"])[0]["id"]
        assert requests.post(f"{API}/notifications/{nid}/read",
                             headers={"Authorization": f"Bearer {ta_tok}"}, timeout=30).status_code == 200
        own = [n for n in self._dep(to_tok, pic_setup["tag"]) if n["id"] == nid][0]
        assert own["read"] is False
        tb = self._tok(f"test.staffb.{pic_setup['tag']}@lpk.id", "Rahasia123")
        assert requests.post(f"{API}/notifications/{nid}/read",
                             headers={"Authorization": f"Bearer {tb}"}, timeout=30).status_code == 404

    def test_h_no_dup_refresh(self, tokens, pic_setup):
        to = self._tok("owner@lpk.id", "owner123")
        a = [n.get("dedupe_key") for n in self._dep(to, pic_setup["tag"])]
        b = [n.get("dedupe_key") for n in self._dep(to, pic_setup["tag"])]
        assert a == b and len(a) == 1


# ---------------- PAYROLL PROVENANCE (P1.4) ----------------
class TestPayrollProvenance:
    @pytest.fixture(scope="class")
    def prov_setup(self, tokens):
        import uuid
        tag = uuid.uuid4().hex[:6]
        hr_h, owner_h, admin_h = _headers(tokens["hr"]), _headers(tokens["owner"]), _headers(tokens["admin"])
        mk = {"tag": tag}
        r = requests.post(f"{API}/employees", headers=hr_h,
                          json={"nama": f"TEST Prov Guru {tag}", "tipe": "guru", "jabatan": "Pengajar",
                                "gaji_pokok": 0, "tunjangan": 200000, "honor_per_pertemuan": 100000}, timeout=30)
        assert r.status_code == 200, r.text
        mk["guru"] = r.json()["id"]
        r = requests.post(f"{API}/employees", headers=hr_h,
                          json={"nama": f"TEST Prov Kar {tag}", "tipe": "karyawan", "jabatan": "Staff",
                                "gaji_pokok": 3000000, "tunjangan": 300000}, timeout=30)
        assert r.status_code == 200, r.text
        mk["kar"] = r.json()["id"]
        r = requests.post(f"{API}/students", headers=admin_h,
                          json={"nama_lengkap": f"TEST Prov Siswa {tag}", "jenis_kelamin": "L"}, timeout=30)
        mk["siswa"] = r.json()["id"]
        mk["classes"] = []
        for nm in ("A", "B"):
            r = requests.post(f"{API}/classes", headers=admin_h,
                              json={"nama": f"TEST Prov Kelas {nm} {tag}", "guru_id": mk["guru"]}, timeout=30)
            assert r.status_code == 200, r.text
            cid = r.json()["id"]
            mk["classes"].append(cid)
            assert requests.post(f"{API}/classes/{cid}/students", headers=admin_h,
                                 json={"student_ids": [mk["siswa"]]}, timeout=30).status_code == 200
        for tgl, cids in (("2021-05-03", [mk["classes"][0], mk["classes"][1]]),
                          ("2021-05-10", [mk["classes"][0]])):
            for cid in cids:
                r = requests.post(f"{API}/attendance", headers=admin_h,
                                  json={"class_id": cid, "tanggal": tgl,
                                        "records": [{"student_id": mk["siswa"], "status": "hadir"}]}, timeout=30)
                assert r.status_code == 200, r.text
        yield mk
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        pids = [p["id"] for p in d.payrolls.find({"employee_id": {"$in": [mk["guru"], mk["kar"]]}},
                                                 {"_id": 0, "id": 1})]
        d.payrolls.delete_many({"id": {"$in": pids}})
        d.audit_logs.delete_many({"entity_id": {"$in": pids + [mk["guru"], mk["kar"]]}})
        for cid in mk["classes"]:
            d.attendance.delete_many({"class_id": cid})
            requests.delete(f"{API}/classes/{cid}", headers=admin_h, timeout=30)
        requests.delete(f"{API}/students/{mk['siswa']}", headers=owner_h, timeout=30)
        for eid in (mk["guru"], mk["kar"]):
            requests.delete(f"{API}/employees/{eid}", headers=hr_h, timeout=30)

    def test_01_suggest_distinct(self, tokens, prov_setup):
        r = requests.get(f"{API}/payrolls/suggest-meetings",
                         headers=_headers(tokens["hr"]),
                         params={"employee_id": prov_setup["guru"], "periode": "2021-05"}, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["suggested_meetings"] == 3, d
        keys = {(s["tanggal"], s["guru_id"], s["class_id"]) for s in d["suggestions"]}
        assert len(keys) == 3

    def test_02_suggest_two_classes_same_date(self, tokens, prov_setup):
        r = requests.get(f"{API}/payrolls/suggest-meetings",
                         headers=_headers(tokens["hr"]),
                         params={"employee_id": prov_setup["guru"], "periode": "2021-05"}, timeout=30)
        same_day = [s for s in r.json()["suggestions"] if s["tanggal"] == "2021-05-03"]
        assert len(same_day) == 2
        assert {s["class_id"] for s in same_day} == set(prov_setup["classes"])

    def test_03_suggest_flags(self, tokens, prov_setup):
        r = requests.get(f"{API}/payrolls/suggest-meetings",
                         headers=_headers(tokens["hr"]),
                         params={"employee_id": prov_setup["guru"], "periode": "2021-05"}, timeout=30)
        for s in r.json()["suggestions"]:
            assert s["unverified_presence"] is True
            assert s["evidence_count"] >= 1 and s["class_nama"]

    def test_04_suggest_no_side_effect(self, tokens, prov_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        n0 = d.payrolls.count_documents({"employee_id": prov_setup["guru"]})
        requests.get(f"{API}/payrolls/suggest-meetings", headers=_headers(tokens["hr"]),
                     params={"employee_id": prov_setup["guru"], "periode": "2021-05"}, timeout=30)
        assert d.payrolls.count_documents({"employee_id": prov_setup["guru"]}) == n0

    def test_05_meeting_provenance(self, tokens, prov_setup):
        r = requests.post(f"{API}/payrolls/calculate", headers=_headers(tokens["hr"]),
                          json={"periode": "2021-05", "items": [
                              {"employee_id": prov_setup["guru"], "honor_pertemuan": 3,
                               "honor_note": "Konfirmasi rekap kelas"}]}, timeout=30)
        assert r.status_code == 200, r.text
        prov = r.json()["rows"][0]["komponen"]["provenance"]
        assert prov["meetings"]["note"] == "Konfirmasi rekap kelas"
        assert prov["meetings"]["actor"] and prov["meetings"]["at"]

    def test_06_overtime_needs_reason(self, tokens, prov_setup):
        h = _headers(tokens["hr"])
        base = {"periode": "2021-06", "items": [{"employee_id": prov_setup["kar"], "lembur": 100000}]}
        assert requests.post(f"{API}/payrolls/calculate", headers=h, json=base, timeout=30).status_code == 400
        base["items"][0].update({"lembur_reason": " lembur proyek ", "lembur_source_note": "2021-06-01 s.d. 2021-06-03"})
        r = requests.post(f"{API}/payrolls/calculate", headers=h, json=base, timeout=30)
        assert r.status_code == 200, r.text
        prov = r.json()["rows"][0]["komponen"]["provenance"]
        assert prov["lembur"]["reason"] == "lembur proyek"

    def test_07_bonus_needs_reason(self, tokens, prov_setup):
        h = _headers(tokens["hr"])
        base = {"periode": "2021-07", "items": [{"employee_id": prov_setup["kar"], "bonus": 50000}]}
        assert requests.post(f"{API}/payrolls/calculate", headers=h, json=base, timeout=30).status_code == 400
        base["items"][0]["bonus_reason"] = "Target tercapai"
        r = requests.post(f"{API}/payrolls/calculate", headers=h, json=base, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json()["rows"][0]["komponen"]["provenance"]["bonus"]["reason"] == "Target tercapai"

    def test_08_deduction_reference(self, tokens, prov_setup):
        h = _headers(tokens["hr"])
        bad = {"periode": "2021-08", "items": [{"employee_id": prov_setup["kar"], "potongan": [
            {"jenis": "Kasbon", "nominal": 50000, "keterangan": ""}]}]}
        assert requests.post(f"{API}/payrolls/calculate", headers=h, json=bad, timeout=30).status_code == 400
        bad["items"][0]["potongan"] = [
            {"jenis": "Kasbon", "nominal": 50000, "keterangan": "Kasbon Mei",
             "source_ref": {"source_type": "external_doc", "source_note": "Slip kasbon #12"}}]
        r = requests.post(f"{API}/payrolls/calculate", headers=h, json=bad, timeout=30)
        assert r.status_code == 200, r.text
        pot = r.json()["rows"][0]["komponen"]["potongan"]
        assert pot[0]["source_ref"]["source_note"] == "Slip kasbon #12"
        bad2 = {"periode": "2021-09", "items": [{"employee_id": prov_setup["kar"], "potongan": [
            {"jenis": "Kasbon", "nominal": 1, "keterangan": "x",
             "source_ref": {"source_type": "alien"}}]}]}
        assert requests.post(f"{API}/payrolls/calculate", headers=h, json=bad2, timeout=30).status_code == 400

    def test_09_alfa_separate(self, tokens, prov_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        p = d.payrolls.find_one({"employee_id": prov_setup["kar"], "periode": "2021-06"}, {"_id": 0})
        assert p["komponen"]["potongan_alfa"] >= 0
        assert all("alfa" not in (x.get("jenis", "").lower()) for x in p["komponen"]["potongan"])

    def test_11_input_change_audit(self, tokens, prov_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        p = d.payrolls.find_one({"employee_id": prov_setup["kar"], "periode": "2021-07"}, {"_id": 0, "id": 1})
        r = requests.put(f"{API}/payrolls/{p['id']}", headers=_headers(tokens["hr"]),
                         json={"bonus": 80000, "bonus_reason": "Revisi target", "alasan": "Koreksi QA"}, timeout=30)
        assert r.status_code == 200, r.text
        au = requests.get(f"{API}/audit-logs?entity=payroll&limit=100",
                          headers=_headers(tokens["owner"]), timeout=30).json()
        upd = [x for x in au if x.get("entity_id") == p["id"] and x["action"] == "update"]
        assert upd and upd[0]["after"].get("bonus") == 80000
        assert "provenance" in upd[0]["after"] and upd[0].get("alasan") == "Koreksi QA"

    def test_12_legacy_readable(self, tokens, prov_setup):
        from pymongo import MongoClient
        import uuid
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        legacy = {"id": f"legacy-{uuid.uuid4().hex[:8]}", "periode": "2019-01",
                  "employee_id": prov_setup["kar"], "employee_nama": "Legacy",
                  "no_slip": "SLIP-LEGACY", "snapshot": {"model": "bulanan"},
                  "komponen": {"bonus": 0, "lembur": 0, "honor_pertemuan": 0, "potongan": []},
                  "bruto": 1, "bersih": 1, "status": "draft", "created_at": "2019-01-01"}
        d.payrolls.insert_one(dict(legacy))
        try:
            r = requests.get(f"{API}/payrolls/{legacy['id']}", headers=_headers(tokens["hr"]), timeout=30)
            assert r.status_code == 200 and "provenance" not in r.json().get("komponen", {})
            r = requests.put(f"{API}/payrolls/{legacy['id']}", headers=_headers(tokens["hr"]),
                             json={"bonus": 10000, "bonus_reason": "Koreksi legacy", "alasan": "QA"}, timeout=30)
            assert r.status_code == 200, r.text
            assert r.json()["komponen"]["provenance"]["bonus"]["reason"] == "Koreksi legacy"
        finally:
            d.payrolls.delete_many({"id": legacy["id"]})

    def test_13_rbac_isolation(self, tokens, prov_setup):
        assert requests.get(f"{API}/payrolls", headers=_headers(tokens["guru"]), timeout=30).status_code == 403
        assert requests.post(f"{API}/payrolls/calculate", headers=_headers(tokens["marketing"]),
                             json={"periode": "2021-05", "items": []}, timeout=30).status_code == 403
        assert requests.get(f"{API}/payrolls/suggest-meetings", headers=_headers(tokens["guru"]),
                            params={"employee_id": prov_setup["guru"], "periode": "2021-05"},
                            timeout=30).status_code == 403

    def _race_once(self, tokens, emp_id, periode):
        import concurrent.futures
        h = _headers(tokens["hr"])

        def hit(_):
            r = requests.post(f"{API}/payrolls/calculate", headers=h,
                              json={"periode": periode, "items": [{"employee_id": emp_id}]}, timeout=60)
            return r.status_code, r.json()
        with concurrent.futures.ThreadPoolExecutor(max_workers=5) as ex:
            return list(ex.map(hit, range(5)))

    def test_14_concurrent_5x(self, tokens, prov_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        res = self._race_once(tokens, prov_setup["kar"], "2021-10")
        assert all(c == 200 for c, _ in res), res
        assert sum(1 for _, b in res if b.get("skipped")) == 4
        assert d.payrolls.count_documents({"employee_id": prov_setup["kar"], "periode": "2021-10"}) == 1

    def test_15_concurrent_3runs(self, tokens, prov_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        for i, per in enumerate(("2021-11", "2021-12", "2022-01")):
            res = self._race_once(tokens, prov_setup["kar"], per)
            assert all(c == 200 for c, _ in res), (per, res)
            assert sum(1 for _, b in res if b.get("skipped")) == 4, (per, res)
            assert d.payrolls.count_documents({"employee_id": prov_setup["kar"], "periode": per}) == 1, per

    def test_17_no_dup_audit(self, tokens, prov_setup):
        au = requests.get(f"{API}/audit-logs?entity=payroll&limit=200",
                          headers=_headers(tokens["owner"]), timeout=30).json()
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        for per in ("2021-10", "2021-11", "2021-12", "2022-01"):
            p = d.payrolls.find_one({"employee_id": prov_setup["kar"], "periode": per}, {"_id": 0, "id": 1})
            assert p is not None
            creates = [x for x in au if x.get("entity_id") == p["id"] and x["action"] == "create"]
            assert len(creates) == 1, (per, len(creates))

    def test_18_financial_unchanged(self, tokens, prov_setup):
        from pymongo import MongoClient
        d = MongoClient("mongodb://127.0.0.1:27017")["lpk"]
        fin_h = _headers(tokens["finance"])
        tx0, saldo0 = d.transactions.count_documents({}), \
            requests.get(f"{API}/finance/summary?period=bulan", headers=fin_h, timeout=30).json()["saldo_kas"]
        p = d.payrolls.find_one({"employee_id": prov_setup["guru"], "periode": "2021-05"},
                                {"_id": 0, "komponen": 1, "bruto": 1, "bersih": 1, "snapshot": 1})
        k = p["komponen"]
        assert p["bruto"] == k["base"] + k["tunjangan_hitung"] + k["lembur"] + k["bonus"]
        assert p["bersih"] == p["bruto"] - k["potongan_alfa"] - sum(x["nominal"] for x in k["potongan"])
        assert k["honor_total"] == p["snapshot"]["honor_per_pertemuan"] * k["honor_pertemuan"]
        assert (d.transactions.count_documents({}),
                requests.get(f"{API}/finance/summary?period=bulan", headers=fin_h, timeout=30).json()["saldo_kas"]) == (tx0, saldo0)
