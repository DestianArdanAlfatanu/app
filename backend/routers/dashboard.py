import asyncio
import re
from typing import Optional, Dict, Any
from datetime import date, timedelta
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from core import (db, get_current_user, require_roles, deny_student, period_range, today_str, payment_summary_map, fee_total,
                  attendance_summary, grade_average, doc_progress, STUDENT_STATUSES, compute_age, ROLES, DOC_TYPES,
                  new_id, now_iso, clean, log_audit, logger, guru_class_ids, today, local_day_start_utc)
from routers.finance import account_balances

router = APIRouter()
AUDIT_ROLES = ("owner", "admin", "finance", "hr")
# Role yang boleh melihat nominal tagihan siswa (sama dengan FIN_VISIBLE di routers/students.py).
FIN_VISIBLE = ("owner", "admin", "finance", "hr")
STUDENT_MONEY_KEYS = ("total_tagihan", "dibayar", "sisa")
DOC_DONE_STATUSES = ["tersedia", "verified"]

NOTIF_FIN = ["owner", "admin", "finance"]
NOTIF_ABSEN = ["owner", "admin", "guru"]
NOTIF_HR = ["owner", "hr", "admin"]
NOTIF_COLLECTION = ["owner", "finance"]
NOTIF_CANDIDATE = ["owner", "admin", "marketing"]
NOTIF_DEP = ["owner", "admin"]


def _notif_scope(user: dict) -> dict:
    """Satu notifikasi logis terlihat oleh role-nya ATAU user-id yang ditunjuk (PIC).
    read_by tetap per-user sehingga read-state terisolasi."""
    return {"$or": [{"roles": user["role"]}, {"recipients": user["id"]}]}


async def _compute_notif_events():
    """Bangun daftar event notifikasi dari kondisi live. Murni derive, tanpa I/O tulis."""
    from datetime import date as _date
    t = today_str()
    soon = (today() + timedelta(days=30)).isoformat()
    events = []

    def ev(tipe, level, judul, pesan, link, tanggal, roles, entity_type="", entity_id="", dedupe="", wa=None,
           recipients=None):
        events.append({"tipe": tipe, "level": level, "judul": judul, "pesan": pesan, "link": link,
                       "tanggal": tanggal, "roles": roles, "entity_type": entity_type, "entity_id": entity_id,
                       "dedupe_key": dedupe or f"{tipe}:{entity_type}:{entity_id}:{tanggal}",
                       "recipients": recipients or [],
                       "wa": wa})
        return events[-1]

    students = await db.students.find({"status": {"$nin": ["calon_siswa", "gagal", "alumni"]}}, {"_id": 0}).to_list(None)
    pay = await payment_summary_map()
    for s in students:
        sisa = fee_total(s) - pay.get(s["id"], {}).get("bayar", 0)
        jt = s.get("jatuh_tempo")
        if sisa > 0 and jt and jt <= (today() + timedelta(days=7)).isoformat():
            ev("pembayaran", "danger" if jt < t else "warning", f"Tagihan {s['nama_lengkap']}",
               f"Sisa Rp{sisa:,.0f} {'terlambat' if jt < t else 'jatuh tempo'} {jt}", f"/siswa/{s['id']}", jt,
               NOTIF_FIN, "student", s["id"], f"tagihan:{s['id']}:{jt}",
               wa={"template": "payment_due", "student_id": s["id"], "suffix": f"tagihan:{s['id']}:{jt}",
                   "recipients": ["siswa", "wali"],
                   "variables": {"nama": s["nama_lengkap"], "jumlah": f"Rp{sisa:,.0f}".replace(",", "."),
                                 "jatuh_tempo": jt, "cara_bayar": "Transfer/QRIS LPK"}})
    docs = await db.documents.find({"is_deleted": False, "tanggal_kadaluarsa": {"$ne": None, "$lte": soon}}, {"_id": 0}).to_list(None)
    names = {s["id"]: s["nama_lengkap"] for s in await db.students.find({"id": {"$in": [d["student_id"] for d in docs]}}, {"_id": 0, "id": 1, "nama_lengkap": 1}).to_list(None)}
    for d in docs:
        exp = d["tanggal_kadaluarsa"]
        ev("dokumen", "danger" if exp < t else "warning", f"{d['jenis']} - {names.get(d['student_id'], '')}",
           f"{'Sudah expired' if exp < t else 'Akan expired'} {exp}", f"/siswa/{d['student_id']}", exp,
           ROLES, "document", d["id"], f"dok:{d['id']}:{exp}",
           wa={"template": "document_expiry", "student_id": d["student_id"], "suffix": f"dok:{d['id']}:{exp}",
               "recipients": ["siswa", "wali"],
               "variables": {"nama": names.get(d["student_id"], ""), "jenis_dokumen": d["jenis"], "tanggal": exp}})
    for iv in await db.interviews.find({"tanggal": {"$gte": t, "$lte": (today() + timedelta(days=7)).isoformat()}, "hasil": "menunggu"}, {"_id": 0}).to_list(None):
        ev("interview", "info", f"Interview {iv['student_nama']}", f"{iv['perusahaan']} - {iv['posisi']} pada {iv['tanggal']}",
           "/job-order", iv["tanggal"], ROLES, "interview", iv["id"], f"iv:{iv['id']}",
           wa={"template": "interview_reminder", "student_id": iv.get("student_id"), "suffix": f"iv:{iv['id']}",
               "recipients": ["siswa"],
               "variables": {"nama": iv["student_nama"], "perusahaan": iv["perusahaan"],
                             "posisi": iv["posisi"], "tanggal": iv["tanggal"]}})
    absen = await db.attendance.find({"tanggal": t, "status": "alfa"}, {"_id": 0}).to_list(None)
    anames = {s["id"]: s["nama_lengkap"] for s in await db.students.find({"id": {"$in": [a["student_id"] for a in absen]}}, {"_id": 0, "id": 1, "nama_lengkap": 1}).to_list(None)}
    for a in absen:
        ev("absensi", "warning", f"{anames.get(a['student_id'], '')} tidak hadir", "Alfa hari ini",
           f"/siswa/{a['student_id']}", t, NOTIF_ABSEN, "student", a["student_id"], f"alfa:{a['student_id']}:{t}")
    for e in await db.employees.find({"kontrak_berakhir": {"$ne": None, "$lte": soon}, "aktif": True}, {"_id": 0}).to_list(None):
        ev("kontrak", "warning", f"Kontrak {e['nama']}", f"Berakhir {e['kontrak_berakhir']}", "/sdm",
           e["kontrak_berakhir"], NOTIF_HR, "employee", e["id"], f"kontrak:{e['id']}:{e['kontrak_berakhir']}")
    for lv in await db.leaves.find({"status": "menunggu"}, {"_id": 0}).sort("dari", 1).to_list(None):
        emp = await db.employees.find_one({"id": lv["employee_id"]}, {"_id": 0, "nama_lengkap": 1, "nama": 1})
        nama = (emp or {}).get("nama") or (emp or {}).get("nama_lengkap") or ""
        ev("cuti", "info", f"Pengajuan cuti {nama}", f"{lv['dari']} → {lv['sampai']} ({lv.get('durasi_hari', '?')} hari)",
           "/sdm", lv["dari"], NOTIF_HR, "leave", lv["id"], f"cuti:{lv['id']}")
    # P1.1 collection: reminder follow-up penagihan. QA-COL-01: tentukan activity
    # TERBARU dari SEMUA activity per siswa dulu, baru evaluasi. Activity lama
    # tidak mewariskan reminder bila sudah digantikan activity baru; outcome
    # paid_after_contact atau tanpa follow-up => tidak ada reminder.
    # Tanpa auto-WA (tanpa key "wa"). Overdue bila fu < hari ini.
    latest_col = {}
    for c in await db.collection_activities.find({}, {"_id": 0}).sort("created_at", -1).to_list(None):
        latest_col.setdefault(c["student_id"], c)
    fu_names = {s["id"]: s["nama_lengkap"] for s in await db.students.find(
        {"id": {"$in": list(latest_col.keys())}}, {"_id": 0, "id": 1, "nama_lengkap": 1}).to_list(None)} if latest_col else {}
    for sid, c in latest_col.items():
        fu = c.get("next_follow_up_at")
        if not fu or c.get("outcome") == "paid_after_contact":
            continue
        if fu <= t:
            ev("collection", "danger" if fu < t else "warning",
               f"Follow-up penagihan {fu_names.get(sid, '')}",
               f"{'Terlambat' if fu < t else 'Jatuh tempo'} {fu} — {c.get('note', '')[:80]}",
               "/pembayaran?tab=tunggakan", fu, NOTIF_COLLECTION, "collection", c["id"], f"col-fu:{c['id']}")
    # P1.2 candidate: pola latest-wins yang sama (versi benar: terbaru dari SEMUA
    # activity dulu). Closed = mendaftar/menolak/tidak_aktif. Tanpa auto-WA.
    latest_calon = {}
    for c in await db.candidate_followups.find({}, {"_id": 0}).sort("created_at", -1).to_list(None):
        latest_calon.setdefault(c["student_id"], c)
    calon_names = {s["id"]: s["nama_lengkap"] for s in await db.students.find(
        {"id": {"$in": list(latest_calon.keys())}}, {"_id": 0, "id": 1, "nama_lengkap": 1}).to_list(None)} if latest_calon else {}
    for sid, c in latest_calon.items():
        fu = c.get("next_follow_up_at")
        if not fu or c.get("outcome") in ("mendaftar", "menolak", "tidak_aktif"):
            continue
        if fu <= t:
            ev("candidate", "danger" if fu < t else "warning",
               f"Follow-up calon {calon_names.get(sid, '')}",
               f"{'Terlambat' if fu < t else 'Jatuh tempo'} {fu} — {c.get('note', '')[:80]}",
               "/followup", fu, NOTIF_CANDIDATE, "candidate", c["id"], f"candidate-fu:{c['id']}")
    # P1.3 departure: satu event per profile dari readiness terkini (latest-wins:
    # derive ulang tiap sync; selesai/berangkat => tanpa event => stale terprune).
    # QA-DEP-01: penerima = owner/admin + PIC spesifik (pic_user_id), BUKAN
    # seluruh role staff. Tanpa auto-WA.
    from routers.departures import compute_readiness
    for p in await db.departure_profiles.find({}, {"_id": 0}).to_list(None):
        r = await compute_readiness(p)
        nm = (await db.students.find_one({"id": p["student_id"]}, {"_id": 0, "nama_lengkap": 1}) or {}).get("nama_lengkap", "")
        pic = [p["pic_user_id"]] if p.get("pic_user_id") else []
        if r["readiness_status"] == "BLOCKED":
            ev("departure", "danger", f"Keberangkatan BLOCKED: {nm}",
               (r["blockers"][0] if r["blockers"] else "Blocked")[:120],
               f"/departure?student={p['student_id']}", t, NOTIF_DEP, "departure", p["id"], f"dep:{p['id']}",
               recipients=pic)
        elif r["readiness_status"] != "READY" and (r["target_departure_date"] or "") != "":
            try:
                dd = (_date.fromisoformat(r["target_departure_date"]) - today()).days
            except ValueError:
                dd = None
            if dd is not None and 0 <= dd <= 14:
                ev("departure", "warning", f"Keberangkatan dekat: {nm}",
                   f"H-{dd} {r['target_departure_date']} — {r['checklist_verified']}/{r['checklist_total']} verified",
                   f"/departure?student={p['student_id']}", r["target_departure_date"], NOTIF_DEP,
                   "departure", p["id"], f"dep:{p['id']}", recipients=pic)
    return events


_sync_task: Optional[asyncio.Task] = None


async def _sync_notifications():
    """Satu sinkronisasi pada satu waktu. Pemanggil menunggu sinkronisasi yang sedang berjalan (mungkin dimulai
    sebelum perubahannya), lalu ikut sinkronisasi berikutnya yang dipakai bersama semua pemanggil yang menunggu."""
    global _sync_task
    if _sync_task is not None and not _sync_task.done():
        try:
            await asyncio.shield(_sync_task)
        except Exception:
            pass
    if _sync_task is None or _sync_task.done():
        _sync_task = asyncio.ensure_future(_run_sync_notifications())
    return await asyncio.shield(_sync_task)


async def _run_sync_notifications():
    """Upsert event kini + hapus event basi. Idempoten: refresh berulang tidak duplikat."""
    events = await _compute_notif_events()
    keys = []
    for e in events:
        keys.append(e["dedupe_key"])
        await db.notifications.update_one(
            {"dedupe_key": e["dedupe_key"]},
            {"$set": {k: e[k] for k in ("tipe", "level", "judul", "pesan", "link", "tanggal", "roles", "recipients",
                                        "entity_type", "entity_id")},
             "$setOnInsert": {"id": new_id(), "read_by": [], "created_at": now_iso()}}, upsert=True)
    await db.notifications.delete_many({"dedupe_key": {"$nin": keys}})
    await db.notifications.delete_many({"expires_at": {"$lt": today_str()}})
    try:
        from routers.whatsapp import dispatch_event
        for e in events:
            w = e.get("wa")
            if w and w.get("student_id"):
                await dispatch_event(w["template"], w["student_id"], w["suffix"], w["variables"], w["recipients"])
    except Exception as ex:
        logger.error(f"WA dispatch gagal: {ex}")
    return keys


@router.get("/dashboard")
async def dashboard(period: str = "bulan", user: dict = Depends(get_current_user)):
    deny_student(user)
    start, end = period_range(period)
    students = await db.students.find({}, {"_id": 0, "id": 1, "status": 1, "fee_plan": 1, "created_at": 1}).to_list(None)
    by_status = {s: 0 for s in STUDENT_STATUSES}
    for s in students:
        by_status[s["status"]] = by_status.get(s["status"], 0) + 1
    start_utc = local_day_start_utc(date.fromisoformat(start))  # created_at disimpan UTC
    baru = sum(1 for s in students if s["created_at"] >= start_utc)
    aktif_statuses = ["diterima", "pelatihan", "ujian", "lulus", "matching", "pemberkasan", "visa"]
    out: Dict[str, Any] = {"period": period, "dari": start, "sampai": end, "siswa": {
        "total": len(students), "calon": by_status["calon_siswa"] + by_status["pendaftaran"] + by_status["seleksi"],
        "aktif": sum(by_status[s] for s in aktif_statuses), "pelatihan": by_status["pelatihan"], "lulus": by_status["lulus"],
        "gagal": by_status["gagal"], "berangkat": by_status["berangkat"], "alumni": by_status["alumni"], "baru_periode": baru,
        "per_status": by_status}}
    role = user["role"]
    if role in ("owner", "admin", "finance"):
        rows = await db.transactions.aggregate([{"$match": {"tanggal": {"$gte": start, "$lte": end}}},
                                                {"$group": {"_id": "$jenis", "n": {"$sum": "$nominal"}}}]).to_list(5)
        m = {r["_id"]: r["n"] for r in rows}
        pay = await payment_summary_map()
        piutang = sum(max(fee_total(s) - pay.get(s["id"], {}).get("bayar", 0), 0) for s in students if s["status"] not in ("calon_siswa", "gagal"))
        tunggakan = sum(1 for s in students if s["status"] not in ("calon_siswa", "gagal", "alumni") and fee_total(s) - pay.get(s["id"], {}).get("bayar", 0) > 0)
        today_pay = await db.payments.aggregate([{"$match": {"tanggal": today_str()}}, {"$group": {"_id": None, "n": {"$sum": "$nominal"}}}]).to_list(1)
        accounts = await account_balances()
        out["keuangan"] = {"pemasukan": m.get("pemasukan", 0), "pengeluaran": m.get("pengeluaran", 0), "saldo_kas": sum(a["saldo"] for a in accounts),
                           "piutang": piutang, "siswa_tunggakan": tunggakan, "pembayaran_hari_ini": today_pay[0]["n"] if today_pay else 0,
                           "accounts": [{"nama": a["nama"], "saldo": a["saldo"]} for a in accounts]}
    guru = await db.employees.count_documents({"tipe": "guru", "aktif": True})
    karyawan = await db.employees.count_documents({"tipe": "karyawan", "aktif": True})
    kelas_aktif = await db.classes.find({"status": "aktif"}, {"_id": 0, "id": 1, "nama": 1, "jadwal": 1, "guru_id": 1}).to_list(None)
    hari_map = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
    today_name = hari_map[today().weekday()]
    teachers = {t["id"]: t["nama"] for t in await db.employees.find({"tipe": "guru"}, {"_id": 0, "id": 1, "nama": 1}).to_list(None)}
    jadwal_hari_ini = [{"kelas": c["nama"], "guru": teachers.get(c.get("guru_id")), **j} for c in kelas_aktif for j in c.get("jadwal", []) if j.get("hari") == today_name]
    upcoming = (today() + timedelta(days=14)).isoformat()
    exams = await db.exams.find({"tanggal": {"$gte": today_str(), "$lte": upcoming}}, {"_id": 0}).to_list(None)
    ujian_peserta = 0
    for e in exams:
        if e.get("class_id"):
            c = await db.classes.find_one({"id": e["class_id"]}, {"_id": 0, "student_ids": 1})
            ujian_peserta += len(c.get("student_ids", [])) if c else 0
    interviews = await db.interviews.find({"tanggal": {"$gte": today_str()}, "hasil": "menunggu"}, {"_id": 0}).sort("tanggal", 1).to_list(20)
    out["operasional"] = {"guru": guru, "karyawan": karyawan, "kelas_aktif": len(kelas_aktif), "jadwal_hari_ini": jadwal_hari_ini,
                          "ujian_mendatang": [{"nama": e["nama"], "tanggal": e["tanggal"], "jenis": e["jenis"]} for e in exams],
                          "siswa_akan_ujian": ujian_peserta, "interview_mendatang": interviews[:5],
                          "job_order_terbuka": await db.job_orders.count_documents({"status": "terbuka"})}
    absen_today = await db.attendance.find({"tanggal": today_str(), "status": {"$ne": "hadir"}}, {"_id": 0}).to_list(None)
    out["absensi_hari_ini"] = {"tidak_hadir": len(absen_today)}
    pending: Dict[str, Any] = {}
    if role in ("owner", "admin", "finance"):
        exps = await db.expenses.find({"status": "diajukan"}, {"_id": 0, "nominal": 1}).to_list(None)
        pending["expense"] = {"count": len(exps), "total": sum(e.get("nominal", 0) for e in exps)}
    if role in ("owner", "admin", "hr"):
        pending["leave"] = {"count": await db.leaves.count_documents({"status": "menunggu"})}
        drafts = await db.payrolls.find({"status": "draft"}, {"_id": 0, "bersih": 1}).to_list(None)
        pending["payroll"] = {"count": len(drafts), "total": sum(p.get("bersih", 0) for p in drafts)}
    out["pending"] = pending
    # Audit log memuat nilai sebelum/sesudah (nominal payroll, pembayaran): hanya role yang boleh membuka Audit Log.
    if role in AUDIT_ROLES:
        out["aktivitas_terbaru"] = await db.audit_logs.find({}, {"_id": 0}).sort("timestamp", -1).to_list(8)
    else:
        out["aktivitas_terbaru"] = []
    return out


@router.get("/notifications")
async def notifications(user: dict = Depends(get_current_user)):
    deny_student(user)
    await _sync_notifications()
    rows = await db.notifications.find(_notif_scope(user), {"_id": 0}).sort("tanggal", 1).to_list(None)
    for r in rows:
        r["read"] = user["id"] in (r.get("read_by") or [])
    return rows


@router.get("/notifications/unread-count")
async def notifications_unread_count(user: dict = Depends(get_current_user)):
    # Hanya membaca; sinkronisasi dilakukan oleh GET /notifications dan job berkala di server.py.
    deny_student(user)
    return {"unread": await db.notifications.count_documents({**_notif_scope(user), "read_by": {"$ne": user["id"]}})}


@router.post("/notifications/{notif_id}/read")
async def notification_read(notif_id: str, user: dict = Depends(get_current_user)):
    deny_student(user)
    n = await db.notifications.find_one({"id": notif_id, **_notif_scope(user)}, {"_id": 0, "id": 1})
    if not n:
        raise HTTPException(status_code=404, detail="Notifikasi tidak ditemukan")
    await db.notifications.update_one({"id": notif_id}, {"$addToSet": {"read_by": user["id"]}})
    return {"ok": True}


@router.post("/notifications/read-all")
async def notifications_read_all(user: dict = Depends(get_current_user)):
    deny_student(user)
    await db.notifications.update_many(_notif_scope(user), {"$addToSet": {"read_by": user["id"]}})
    return {"ok": True}


@router.get("/search")
async def search(q: str = Query(..., min_length=1), user: dict = Depends(get_current_user)):
    deny_student(user)
    query: Dict[str, Any] = {"$or": [{"nama_lengkap": {"$regex": re.escape(q), "$options": "i"}}, {"nik": {"$regex": re.escape(q)}}]}
    if user["role"] == "guru":
        emp_classes = await db.classes.find({"guru_id": user.get("employee_id")}, {"_id": 0, "id": 1}).to_list(None)
        query["class_id"] = {"$in": [c["id"] for c in emp_classes]}
    rows = await db.students.find(query, {"_id": 0}).to_list(8)
    show_pay = user["role"] in ("owner", "admin", "finance", "hr")
    pay = await payment_summary_map([r["id"] for r in rows]) if show_pay else {}
    classes = {c["id"]: c["nama"] for c in await db.classes.find({}, {"_id": 0, "id": 1, "nama": 1}).to_list(None)}
    out = []
    for s in rows:
        iv = await db.interviews.find_one({"student_id": s["id"]}, {"_id": 0}, sort=[("tanggal", -1)])
        item = {"id": s["id"], "nama_lengkap": s["nama_lengkap"], "status": s["status"], "kelas": classes.get(s.get("class_id")),
                "absensi": await attendance_summary(s["id"]), "nilai": await grade_average(s["id"]),
                "job_order": iv["perusahaan"] if iv else None, "dokumen": await doc_progress(s["id"])}
        if show_pay:
            item["pembayaran"] = {"total": fee_total(s), "bayar": pay.get(s["id"], {}).get("bayar", 0)}
        out.append(item)
    result: Dict[str, Any] = {"students": out}
    role = user["role"]
    if role in ("owner", "admin", "finance", "hr", "guru", "marketing", "staff"):
        emps = await db.employees.find({"nama": {"$regex": re.escape(q), "$options": "i"}},
                                       {"_id": 0, "id": 1, "nama": 1, "tipe": 1, "jabatan": 1}).to_list(8)
        result["employees"] = [{"entity_type": "employee", "entity_id": e["id"], "title": e["nama"],
                                "subtitle": f"{e.get('tipe', '')} · {e.get('jabatan', '')}", "route": "/sdm"} for e in emps]
    if role in ("owner", "admin", "marketing", "staff", "finance", "hr", "guru"):
        jobs = await db.job_orders.find({"$or": [{"perusahaan": {"$regex": re.escape(q), "$options": "i"}},
                                                 {"posisi": {"$regex": re.escape(q), "$options": "i"}}]},
                                        {"_id": 0, "id": 1, "perusahaan": 1, "posisi": 1, "status": 1}).to_list(8)
        result["job_orders"] = [{"entity_type": "job_order", "entity_id": j["id"],
                                 "title": f"{j.get('perusahaan', '')} — {j.get('posisi', '')}",
                                 "subtitle": j.get("status", ""), "route": "/job-order"} for j in jobs]
    return result


@router.get("/audit-logs")
async def audit_logs(entity: Optional[str] = None, limit: int = 200, user: dict = Depends(require_roles("admin", "finance", "hr"))):
    q = {"entity": entity} if entity else {}
    return await db.audit_logs.find(q, {"_id": 0}).sort("timestamp", -1).to_list(max(1, min(limit, 1000)))


@router.get("/reports/students")
async def report_students(page: Optional[int] = None, limit: int = 200,
                           user: dict = Depends(require_roles("admin", "marketing", "hr", "finance"))):
    data = await _students_report(show_money=user["role"] in FIN_VISIBLE)
    return _paginate(data, "rows", page, limit)


def _paginate(data: dict, key: str, page: Optional[int], limit: int):
    rows = data.get(key, [])
    total = len(rows)
    if not page or page < 1:
        return {**data, "total": total, "page": 1, "limit": total}
    limit = max(1, min(limit or 200, 1000))
    start = (page - 1) * limit
    return {**data, key: rows[start:start + limit], "total": total, "page": page, "limit": limit}


async def _students_report(show_money: bool = True):
    students = await db.students.find({}, {"_id": 0}).to_list(None)
    ids = [s["id"] for s in students]
    pay = await payment_summary_map()
    classes = {c["id"]: c["nama"] for c in await db.classes.find({}, {"_id": 0, "id": 1, "nama": 1}).to_list(None)}
    att_rows = await db.attendance.aggregate([
        {"$match": {"student_id": {"$in": ids}}},
        {"$group": {"_id": {"s": "$student_id", "st": "$status"}, "n": {"$sum": 1}}}]).to_list(None)
    att_map: Dict[str, Dict[str, int]] = {}
    for r in att_rows:
        att_map.setdefault(r["_id"]["s"], {}).update({r["_id"]["st"]: r["n"]})
    grade_rows = await db.grades.aggregate([
        {"$match": {"student_id": {"$in": ids}}},
        {"$group": {"_id": "$student_id", "avg": {"$avg": "$nilai_akhir"}}}]).to_list(None)
    grade_map = {r["_id"]: round(r["avg"], 1) for r in grade_rows}
    doc_rows = await db.documents.find({"student_id": {"$in": ids}, "is_deleted": False, "status": {"$in": DOC_DONE_STATUSES}},
                                       {"_id": 0, "student_id": 1, "jenis": 1}).to_list(None)
    doc_map: Dict[str, set] = {}
    for d in doc_rows:
        doc_map.setdefault(d["student_id"], set()).add(d["jenis"])
    doc_total = len(DOC_TYPES)
    rows = []
    for s in students:
        total = fee_total(s)
        bayar = pay.get(s["id"], {}).get("bayar", 0)
        ac = att_map.get(s["id"], {})
        tot = sum(ac.values())
        hadir_pct = round(ac.get("hadir", 0) / tot * 100, 1) if tot else 0
        rows.append({"nama": s["nama_lengkap"], "nik": s.get("nik"), "jenis_kelamin": s.get("jenis_kelamin"), "usia": compute_age(s.get("tanggal_lahir")),
                     "status": s["status"], "kelas": classes.get(s.get("class_id")), "bahasa": s.get("kemampuan_bahasa_jepang"),
                     "total_tagihan": total, "dibayar": bayar, "sisa": max(total - bayar, 0), "kehadiran": hadir_pct,
                     "nilai": grade_map.get(s["id"]), "dokumen": len(doc_map.get(s["id"], set())), "no_hp": s.get("no_hp")})
        if not show_money:
            for k in STUDENT_MONEY_KEYS:
                rows[-1].pop(k)
    per_status = {}
    for r in rows:
        per_status[r["status"]] = per_status.get(r["status"], 0) + 1
    return {"rows": rows, "per_status": per_status}


@router.get("/reports/finance")
async def report_finance(dari: Optional[str] = None, sampai: Optional[str] = None, user: dict = Depends(require_roles("admin", "finance"))):
    return await _finance_report(dari, sampai)


async def _finance_report(dari: Optional[str] = None, sampai: Optional[str] = None):
    q: Dict[str, Any] = {}
    if dari or sampai:
        q["tanggal"] = {k: v for k, v in (("$gte", dari), ("$lte", sampai)) if v}
    rows = await db.transactions.find(q, {"_id": 0}).sort("tanggal", 1).to_list(None)
    accs = {a["id"]: a["nama"] for a in await db.accounts.find({}, {"_id": 0, "id": 1, "nama": 1}).to_list(None)}
    for r in rows:
        r["account_nama"] = accs.get(r["account_id"])
    pemasukan = sum(r["nominal"] for r in rows if r["jenis"] == "pemasukan")
    pengeluaran = sum(r["nominal"] for r in rows if r["jenis"] == "pengeluaran")
    per_kat: Dict[str, float] = {}
    for r in rows:
        per_kat[f"{r['jenis']}|{r['kategori']}"] = per_kat.get(f"{r['jenis']}|{r['kategori']}", 0) + r["nominal"]
    return {"rows": rows, "pemasukan": pemasukan, "pengeluaran": pengeluaran, "laba": pemasukan - pengeluaran,
            "per_kategori": [{"jenis": k.split("|")[0], "kategori": k.split("|")[1], "nominal": v} for k, v in per_kat.items()],
            "accounts": await account_balances()}


@router.get("/reports/hr")
async def report_hr(user: dict = Depends(require_roles("admin", "hr"))):
    return await _hr_report()


async def _hr_report():
    emps = await db.employees.find({}, {"_id": 0}).to_list(None)
    classes = await db.classes.find({}, {"_id": 0, "id": 1, "nama": 1, "guru_id": 1}).to_list(None)
    for e in emps:
        e["jumlah_kelas"] = sum(1 for c in classes if c.get("guru_id") == e["id"])
    return {"rows": emps, "total_gaji": sum(e.get("gaji_pokok", 0) + e.get("tunjangan", 0) for e in emps if e.get("aktif", True))}


@router.get("/reports/training")
async def report_training(class_id: Optional[str] = None, page: Optional[int] = None, limit: int = 200,
                          user: dict = Depends(require_roles("admin", "guru", "hr"))):
    data = await _training_report(class_id, await guru_class_ids(user))
    return _paginate(data, "rows", page, limit)


async def _training_report(class_id: Optional[str] = None, only_classes: Optional[set] = None):
    q = {"class_id": class_id} if class_id else {"class_id": {"$ne": None}}
    if only_classes is not None:
        if class_id and class_id not in only_classes:
            raise HTTPException(status_code=403, detail="Anda bukan pengajar kelas ini")
        if not class_id:
            q = {"class_id": {"$in": sorted(only_classes)}}
    students = await db.students.find(q, {"_id": 0, "id": 1, "nama_lengkap": 1, "class_id": 1, "status": 1}).to_list(None)
    ids = [s["id"] for s in students]
    classes = {c["id"]: c["nama"] for c in await db.classes.find({}, {"_id": 0, "id": 1, "nama": 1}).to_list(None)}
    att_rows = await db.attendance.aggregate([
        {"$match": {"student_id": {"$in": ids}}},
        {"$group": {"_id": {"s": "$student_id", "st": "$status"}, "n": {"$sum": 1}}}]).to_list(None)
    att_map: Dict[str, Dict[str, Any]] = {}
    for r in att_rows:
        d = att_map.setdefault(r["_id"]["s"], {"hadir": 0, "izin": 0, "sakit": 0, "alfa": 0})
        d[r["_id"]["st"]] = d.get(r["_id"]["st"], 0) + r["n"]
    grade_rows = await db.grades.aggregate([
        {"$match": {"student_id": {"$in": ids}}},
        {"$group": {"_id": "$student_id", "avg": {"$avg": "$nilai_akhir"}}}]).to_list(None)
    grade_map = {r["_id"]: round(r["avg"], 1) for r in grade_rows}
    class_ids = list({s["class_id"] for s in students if s.get("class_id")})
    exams = await db.exams.find({"class_id": {"$in": class_ids}}, {"_id": 0, "nama": 1, "tanggal": 1, "results": 1}).sort("tanggal", 1).to_list(None)
    exam_by_student: Dict[str, list] = {}
    for e in exams:
        for res in (e.get("results") or []):
            if res.get("student_id") in ids:
                exam_by_student.setdefault(res["student_id"], []).append(
                    {"nama": e["nama"], "tanggal": e["tanggal"], "nilai": res.get("nilai")})
    rows = []
    for s in students:
        ac = att_map.get(s["id"], {"hadir": 0, "izin": 0, "sakit": 0, "alfa": 0})
        tot = sum(v for k, v in ac.items() if k in ("hadir", "izin", "sakit", "alfa"))
        ac["total"] = tot
        ac["persentase"] = round(ac.get("hadir", 0) / tot * 100, 1) if tot else 0
        rows.append({"nama": s["nama_lengkap"], "kelas": classes.get(s["class_id"]), "status": s["status"], **ac,
                     "nilai": grade_map.get(s["id"]), "ujian": exam_by_student.get(s["id"], [])})
    return {"rows": rows}


# ---------- Export infrastructure (D2): XLSX + PDF dari builder yang sama ----------
EXPORT_ROLES = {
    "siswa": ("admin", "marketing", "hr", "finance"),
    "keuangan": ("admin", "finance"),
    "sdm": ("admin", "hr"),
    "pelatihan": ("admin", "guru", "hr"),
}
EXPORT_COLS = {
    "siswa": [("Nama", "nama", "text"), ("NIK", "nik", "text"), ("JK", "jenis_kelamin", "text"), ("Usia", "usia", "num"),
              ("Status", "status", "text"), ("Kelas", "kelas", "text"), ("Bahasa", "bahasa", "text"),
              ("Tagihan", "total_tagihan", "num"), ("Dibayar", "dibayar", "num"), ("Sisa", "sisa", "num"),
              ("Hadir %", "kehadiran", "num"), ("Nilai", "nilai", "num"), ("Dokumen", "dokumen", "num"), ("HP", "no_hp", "text")],
    "keuangan": [("Tanggal", "tanggal", "text"), ("Jenis", "jenis", "text"), ("Kategori", "kategori", "text"),
                 ("Deskripsi", "deskripsi", "text"), ("Rekening", "account_nama", "text"), ("Nominal", "nominal", "num"),
                 ("Petugas", "petugas", "text")],
    "sdm": [("Nama", "nama", "text"), ("Tipe", "tipe", "text"), ("Jabatan", "jabatan", "text"),
            ("Status", "status_kerja", "text"), ("Masuk", "tanggal_masuk", "text"),
            ("Gaji Pokok", "gaji_pokok", "num"), ("Tunjangan", "tunjangan", "num"),
            ("Honor/Pertemuan", "honor_per_pertemuan", "num"), ("Kelas", "jumlah_kelas", "num"),
            ("Kontrak", "kontrak_berakhir", "text")],
    "pelatihan": [("Nama", "nama", "text"), ("Kelas", "kelas", "text"), ("Status", "status", "text"),
                  ("Hadir", "hadir", "num"), ("Izin", "izin", "num"), ("Sakit", "sakit", "num"), ("Alfa", "alfa", "num"),
                  ("% Hadir", "persentase", "num"), ("Nilai", "nilai", "num"), ("Ujian", "ujian", "text")],
}


def _export_flat_rows(tab: str, rows: list) -> list:
    flat = []
    for r in rows:
        d = dict(r)
        if tab == "pelatihan":
            d["ujian"] = "; ".join(f"{u.get('nama')}:{u.get('nilai') if u.get('nilai') is not None else '-'}" for u in (r.get("ujian") or []))
        if tab == "sdm":
            d["sertifikat"] = "; ".join(r.get("sertifikat") or [])
        flat.append(d)
    return flat


async def _report_data_for_export(tab: str, dari: Optional[str], sampai: Optional[str], class_id: Optional[str], user: dict):
    if tab == "siswa":
        data = await _students_report(show_money=user["role"] in FIN_VISIBLE)
    elif tab == "keuangan":
        data = await _finance_report(dari, sampai)
    elif tab == "sdm":
        data = await _hr_report()
    elif tab == "pelatihan":
        data = await _training_report(class_id, await guru_class_ids(user))
    else:
        raise HTTPException(status_code=400, detail="Jenis laporan tidak valid")
    return data


def _export_cols(tab: str, rows: list) -> list:
    cols = EXPORT_COLS[tab]
    if tab == "siswa" and rows and "sisa" not in rows[0]:
        cols = [c for c in cols if c[1] not in STUDENT_MONEY_KEYS]
    return cols


def _build_xlsx(tab: str, rows: list) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = f"laporan-{tab}"[:31]
    cols = _export_cols(tab, rows)
    head_fill = PatternFill("solid", fgColor="0F172A")
    for ci, (h, _, _) in enumerate(cols, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for ri, r in enumerate(rows, 2):
        for ci, (_, k, t) in enumerate(cols, 1):
            v = r.get(k)
            if t == "num":
                try:
                    v = 0 if v is None else float(v)
                except (TypeError, ValueError):
                    v = 0
            elif v is None:
                v = "-"
            else:
                v = str(v)
            ws.cell(row=ri, column=ci, value=v)
    for ci in range(1, len(cols) + 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_pdf(tab: str, rows: list, subtitle: str) -> bytes:
    import io
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = [Paragraph(f"Laporan {tab.capitalize()} — LPK Penyaluran Kerja Jepang", styles["Title"]),
             Paragraph(f"{subtitle} · Dicetak {today_str()}", styles["Normal"]), Spacer(1, 12)]
    cols = _export_cols(tab, rows)
    data = [[h for h, _, _ in cols]]
    for r in rows:
        line = []
        for _, k, t in cols:
            v = r.get(k)
            if t == "num":
                try:
                    line.append(f"{float(0 if v is None else v):,.0f}".replace(",", "."))
                except (TypeError, ValueError):
                    line.append("-")
            else:
                line.append("-" if v is None else str(v))
        data.append(line)
    avail = landscape(A4)[0] - 48
    tbl = Table(data, repeatRows=1, colWidths=[avail / len(cols)] * len(cols))
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(tbl)
    doc.build(story)
    return buf.getvalue()


@router.get("/reports/{tab}/export")
async def report_export(tab: str, format: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
                        dari: Optional[str] = None, sampai: Optional[str] = None,
                        class_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    if tab not in EXPORT_ROLES:
        raise HTTPException(status_code=400, detail="Jenis laporan tidak valid")
    if user["role"] != "owner" and user["role"] not in EXPORT_ROLES[tab]:
        raise HTTPException(status_code=403, detail="Anda tidak memiliki akses ke laporan ini")
    data = await _report_data_for_export(tab, dari, sampai, class_id, user)
    rows = _export_flat_rows(tab, data.get("rows", []))
    stamp = today_str()
    subtitle = f"Periode {dari or '-'} s/d {sampai or '-'}" if tab == "keuangan" else f"Data per {stamp}"
    if format == "xlsx":
        payload = _build_xlsx(tab, rows)
        media, ext = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        payload = _build_pdf(tab, rows, subtitle)
        media, ext = "application/pdf", "pdf"
    await log_audit("export", tab, "export", user, None,
                    {"tab": tab, "format": format, "dari": dari, "sampai": sampai, "rows": len(rows)})
    import io as _io
    return StreamingResponse(_io.BytesIO(payload), media_type=media,
                             headers={"Content-Disposition": f'attachment; filename="laporan-{tab}-{stamp}.{ext}"'})
